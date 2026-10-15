"""내려받는 문서의 배치 — docx(A4 · 열 폭 · 행 쪽 넘김 · 표 머리 반복 · 공백)와 hwp 열 폭.

예전에는 docx 표의 열이 모두 같은 폭이라 사업비 표 '집행 계획' 칸의 글이 한 글자씩 세로로 끊기고, 행이 쪽 끝에서 갈라지고, 표 머리가 반복되지 않았고,
숫자와 한글 사이가 벌어져("5 건 수행") 보였다. hwp는 일정 표 '구분' 칸(순번용 3238)과 초기창업 사업비 금액 칸(5946)이 좁아 글이 끊겼다.
docx는 실제 Word로 열어 확인했고(PDF로 바꿔 눈으로 봄) 테스트는 그 설정이 파일에 들어 있는지 확인한다.

실행:
    pytest tests/test_document_layout.py -v
"""
import io
import os
import re
import shutil
import subprocess
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from app import hwp_export, hwpx_rows
from app.plan_document_export import TEXT_WIDTH_CM, render_plan_docx
from tests.test_hwp_table_layout import _big_data, _data

_RHWP_AVAILABLE = shutil.which(hwp_export.RHWP_BIN) is not None or os.path.isfile(hwp_export.RHWP_BIN)


def _docx(template, data=None):
    return Document(io.BytesIO(render_plan_docx(data or _big_data(template), template=template)))


@pytest.mark.parametrize('template', ['preliminary', 'early_general'])
def test_docx_page_is_a4_with_two_centimeter_margins(template):
    section = _docx(template).sections[0]

    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    assert round(section.left_margin.cm, 1) == 2.0 and round(section.right_margin.cm, 1) == 2.0


@pytest.mark.parametrize('template', ['preliminary', 'early_general'])
def test_every_docx_table_has_fixed_column_widths_that_add_up_to_the_text_width(template):
    doc = _docx(template)

    assert doc.tables
    for table in doc.tables:
        assert table._tbl.tblPr.find(qn('w:tblLayout')).get(qn('w:type')) == 'fixed'
        widths = [column.width.cm for column in table.columns]
        assert sum(widths) == pytest.approx(TEXT_WIDTH_CM, abs=0.05)
        # 칸 폭도 같은 값 — 열 폭만 있고 칸 폭이 없으면 Word가 열을 같은 폭으로 나눈다
        assert [round(cell.width.cm, 1) for cell in table.rows[-1].cells] == [round(w, 1) for w in widths]


def test_the_long_text_columns_are_wider_than_the_short_ones():
    doc = _docx('early_general')
    budget = next(t for t in doc.tables if '집행 계획' in [c.text for c in t.rows[0].cells])
    team = next(t for t in doc.tables if '보유 역량(경력 및 학력 등)' in [c.text for c in t.rows[0].cells])
    schedule = next(t for t in doc.tables if '추진 기간' in [c.text for c in t.rows[0].cells])

    def width_of(table, header):
        return next(col.width.cm for col, cell in zip(table.columns, table.rows[0].cells, strict=True) if cell.text == header)

    assert width_of(budget, '집행 계획') > 3 * width_of(budget, '비목') / 2 and width_of(budget, '집행 계획') >= 4
    assert width_of(team, '보유 역량(경력 및 학력 등)') > 2 * width_of(team, '순번')
    assert width_of(schedule, '세부 내용') > width_of(schedule, '구분')


@pytest.mark.parametrize('template', ['preliminary', 'early_general'])
def test_rows_do_not_split_across_pages_and_grid_tables_repeat_their_header(template):
    doc = _docx(template)
    grids = [t for t in doc.tables if t.rows[0]._tr.find('.//' + qn('w:tblHeader')) is not None]

    assert len(grids) >= 4  # 팀 구성 현황 · 일정 두 표 · 사업비 · 팀 구성(안) · 협력기관 등
    for table in doc.tables:   # 이 자료는 칸 글이 짧아 모든 행이 나뉘지 않는다(긴 칸은 아래 테스트)
        assert all(row._tr.find('.//' + qn('w:cantSplit')) is not None for row in table.rows)
    for table in grids:
        assert [row._tr.find('.//' + qn('w:tblHeader')) is not None for row in table.rows][:2] == [True, False]


@pytest.mark.parametrize('template', ['preliminary', 'early_general'])
def test_grid_table_header_row_stays_on_the_page_of_the_first_content_row(template):
    """쪽 맨 아래에 표 제목과 머리 줄만 남고 내용이 다음 쪽에서 시작하지 않게 — 머리 줄 문단은 다음(첫 내용 줄)과 붙는다."""
    doc = _docx(template)
    grids = [t for t in doc.tables if t.rows[0]._tr.find('.//' + qn('w:tblHeader')) is not None]

    for table in grids:
        assert all(p.paragraph_format.keep_with_next for cell in table.rows[0].cells for p in cell.paragraphs)
        assert not any(p.paragraph_format.keep_with_next for row in table.rows[1:] for cell in row.cells for p in cell.paragraphs)


def test_long_rows_may_split_across_pages_so_the_page_above_is_not_left_half_empty():
    """개요(요약) 표 '문제 인식' 같은 20줄 넘는 칸을 통째로 다음 쪽으로 넘기면 앞 쪽 아래 절반이 비었다(실제 실행 건 10-13) — 긴 행만 나뉠 수 있다."""
    long_text = '\n'.join(f'◦ 둘째 줄까지 넘어가는 긴 문장 {i} — 회원 관리와 수업 예약을 한 화면에서 처리하는 서비스를 개발할 계획임.' for i in range(8))
    doc = _docx('early_general', _data(요약_문제인식=long_text, 요약_실현가능성='짧은 요약'))
    summary = next(t for t in doc.tables if t.rows[0].cells[0].text == '명 칭')
    rows = {row.cells[0].text: row._tr.find('.//' + qn('w:cantSplit')) is not None for row in summary.rows}

    assert rows['문제 인식(Problem)'] is False   # 긴 행은 쪽 끝에서 나뉜다
    assert rows['실현 가능성(Solution)'] is True and rows['명 칭'] is True   # 짧은 행은 예전처럼 나뉘지 않는다


def test_captions_stay_with_their_table_or_picture():
    doc = _docx('preliminary')
    captions = [p for p in doc.paragraphs if p.text.startswith('< ') and p.text.endswith(' >')]

    assert captions and all(p.paragraph_format.keep_with_next for p in captions)


def test_docx_turns_off_the_automatic_gap_between_hangul_and_digits():
    styles_xml = zipfile.ZipFile(io.BytesIO(render_plan_docx(_data(), template='preliminary'))).read('word/styles.xml').decode('utf-8')

    assert re.search(r'w:autoSpaceDE w:val="0"', styles_xml) and re.search(r'w:autoSpaceDN w:val="0"', styles_xml)


# ── hwp 열 폭 ───────────────────────────────────────────────────────────────────
def _cell(col, span, width):
    return (f'<hp:tc><hp:cellAddr colAddr="{col}" rowAddr="0"/><hp:cellSpan colSpan="{span}" rowSpan="1"/>'
            f'<hp:cellSz width="{width}" height="500"/></hp:tc>')


def test_resize_columns_sets_each_cell_to_the_sum_of_the_columns_it_covers():
    xml = ('<hp:tbl id="1" rowCnt="2" colCnt="4"><hp:sz width="1000" widthRelTo="ABSOLUTE" height="1000" protect="0"/>'
           f'<hp:tr>{_cell(0, 1, 100)}{_cell(1, 1, 200)}{_cell(2, 2, 700)}</hp:tr>'
           f'<hp:tr>{_cell(0, 2, 300)}{_cell(2, 1, 300)}{_cell(3, 1, 400)}</hp:tr></hp:tbl>')

    out = hwpx_rows.resize_columns(xml, 0, [400, 100, 250, 250])

    assert re.findall(r'<hp:cellSz width="(\d+)"', out) == ['400', '100', '500', '500', '250', '250']
    with pytest.raises(ValueError):
        hwpx_rows.resize_columns(xml, 0, [500, 500])  # 열 수가 맞지 않으면 거부한다


def test_every_adjusted_hwp_table_keeps_its_total_width():
    for template, tables in hwp_export._COLUMN_WIDTHS_BY_TEMPLATE.items():
        for table, widths in tables.items():
            assert len(widths) in (4, 6), (template, table)
    # 일정 표는 원래 합(47962)을, 초기창업 사업비 표는 47924를 그대로 둔다
    assert sum(hwp_export._SCHEDULE_COLUMNS) == 47962
    assert sum(hwp_export._COLUMN_WIDTHS_BY_TEMPLATE['early_general'][13]) == 47924


@pytest.mark.skipif(not _RHWP_AVAILABLE, reason='rhwp 실행 파일이 없다(RHWP_BIN)')
@pytest.mark.parametrize('template', ['preliminary', 'early_general'])
def test_real_hwp_schedule_division_column_is_wide_enough_and_amounts_fit_on_one_line(template, tmp_path):
    path = tmp_path / 'out.hwp'
    path.write_bytes(hwp_export.render_plan_hwp(_big_data(template), template))
    dump = subprocess.run([hwp_export.RHWP_BIN, 'dump', str(path)], capture_output=True, text=True, encoding='utf-8', errors='replace').stdout
    widths: dict[str, list[int]] = {}
    for width, text in re.findall(r' w=(\d+) .*?text="([^"]*)"', dump):
        widths.setdefault(text, []).append(int(width))

    assert widths['구분'].count(7000) == 2  # 일정 표 두 개의 '구분' 머리 칸(순번용 3238에서 넓힘) — 팀 구성 표의 '구분'은 그대로
    if template == 'early_general':
        assert 7300 in widths['현금'] and 7300 in widths['현물']  # 사업비 금액 칸(5946에서 넓힘) — "28,000,000원"이 한 줄에 들어간다
