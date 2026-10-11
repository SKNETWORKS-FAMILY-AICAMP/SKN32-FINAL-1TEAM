"""계획서 그림(서비스 흐름도 · 서비스 구조도 SVG)을 내려받기 문서에 넣기 (SB-340).

그림은 오케스트레이터 파일 저장소의 SVG다. 웹이 PNG로 바꿔 docx(개요 표 아래) · hwp(맨 뒤 새 쪽)에 넣고, 원본 SVG는 결과물 압축에 넣는다.
그림 하나를 못 받거나 못 바꾸면 그 그림만 빼고 문서는 만든다. hwp 실제 확인은 rhwp 실행 파일이 있을 때만 한다(없으면 건너뜀).

실행:
    pytest tests/test_diagram_images.py -v
"""
import io
import json
import os
import shutil
import subprocess
import zipfile
from types import SimpleNamespace as NS

import pytest
from docx import Document
from orch_fakes import make_diagram, make_file_ref, make_outputs, make_plan_doc, make_section
from sbrain.agents.partner_sw import diagram_svg

from app import diagram_images, hwp_export
from app.diagram_images import MAX_SVG_BYTES, PNG_WIDTH_PX, diagram_pictures, svg_to_png
from app.models import Project
from app.orch import OrchError
from app.plan_texts import PlanTexts
from app.routers.projects import _build_plan_document_data

NODES = ['사용자 가입', '공고 선택', '계획서 생성']
USER_FLOW_SVG = diagram_svg.render({'flowType': 'USER_FLOW', 'nodes': NODES})
ARCHITECTURE_SVG = diagram_svg.render({'flowType': 'SERVICE_ARCHITECTURE', 'nodes': ['웹', '서버', '저장소']})
SVGS = {'userflow.svg': USER_FLOW_SVG, 'architecture.svg': ARCHITECTURE_SVG}


def _diagram(flow_type, name, **file_overrides):
    return make_diagram(flow_type=flow_type, image_file=make_file_ref(name, key=f'r1/e1/uniq/{name}', **file_overrides))


def _plan_doc(*diagrams):
    return make_plan_doc([make_section('3.4.1', '문제인식', '본문')], diagrams=list(diagrams))


def _gateway(files=None):
    files = SVGS if files is None else files

    def read(project_id, key):
        name = key.rsplit('/', 1)[-1]
        if name not in files:
            raise OrchError('FILE_NOT_FOUND')
        return NS(name=name, data=files[name], media_type='image/svg+xml')

    return NS(read_artifact_file=read)


BOTH = (_diagram('SERVICE_ARCHITECTURE', 'architecture.svg'), _diagram('USER_FLOW', 'userflow.svg'))


# ── SVG → PNG ────────────────────────────────────────────────────────────────────────────
def test_the_real_generator_output_becomes_a_png_of_the_set_width():
    png = svg_to_png(USER_FLOW_SVG)
    assert png[:8] == b'\x89PNG\r\n\x1a\n'
    assert int.from_bytes(png[16:20], 'big') == PNG_WIDTH_PX


@pytest.mark.parametrize('svg', [
    b'<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///etc/passwd" width="10" height="10"/></svg>',
    b'<svg xmlns="http://www.w3.org/2000/svg" xmlns:x="http://www.w3.org/1999/xlink"><use x:href="other.svg#a"/></svg>',
    b'<svg xmlns="http://www.w3.org/2000/svg"><rect width="5" height="5" fill="url(http://example.com/a.svg#g)"/></svg>',
])
def test_svg_with_external_references_is_not_drawn(svg):
    with pytest.raises(ValueError):
        svg_to_png(svg)


def test_the_generator_output_has_no_external_reference():
    for svg in SVGS.values():
        assert not diagram_images._EXTERNAL_REFERENCE.search(svg)


# ── 계획서 → 그림 목록 ───────────────────────────────────────────────────────────────────
def test_pictures_come_user_flow_first_with_titles_and_the_original_svg():
    pictures = diagram_pictures(_gateway(), 1, _plan_doc(*BOTH))
    assert [p.title for p in pictures] == ['서비스 흐름도', '서비스 구조도']
    assert [p.svg_name for p in pictures] == ['userflow.svg', 'architecture.svg']
    assert pictures[0].svg == USER_FLOW_SVG
    assert all(p.png[:4] == b'\x89PNG' for p in pictures)
    assert pictures[0].size_px[0] == PNG_WIDTH_PX


def test_no_plan_or_no_diagrams_gives_no_pictures():
    assert diagram_pictures(_gateway(), 1, None) == []
    assert diagram_pictures(_gateway(), 1, _plan_doc()) == []


def test_a_picture_that_cannot_be_read_is_left_out(caplog):
    pictures = diagram_pictures(_gateway({'userflow.svg': USER_FLOW_SVG}), 1, _plan_doc(*BOTH))
    assert [p.title for p in pictures] == ['서비스 흐름도']
    assert 'FILE_NOT_FOUND' in caplog.text


def test_a_broken_svg_is_left_out():
    pictures = diagram_pictures(_gateway({'userflow.svg': b'<svg', 'architecture.svg': ARCHITECTURE_SVG}), 1, _plan_doc(*BOTH))
    assert [p.title for p in pictures] == ['서비스 구조도']


def test_an_oversized_svg_is_not_even_read():
    calls = []
    gateway = NS(read_artifact_file=lambda *a: calls.append(a))
    big = _diagram('USER_FLOW', 'userflow.svg', size=MAX_SVG_BYTES + 1)
    assert diagram_pictures(gateway, 1, _plan_doc(big)) == [] and calls == []


# ── docx · zip ───────────────────────────────────────────────────────────────────────────
def _ready(orch, *diagrams, files=None):
    orch.responses['outputs'] = lambda project_id: make_outputs(plan_doc=_plan_doc(*diagrams), deliverable=NS(
        plan_doc_file=make_file_ref('plan.docx'), prototype_files=[], infographic_file=None))
    orch.responses['read_artifact_file'] = _gateway(files).read_artifact_file


def _create(client, applicant_type='corp'):
    payload = {'applicant_type': applicant_type, 'biz_type': '법인', 'ceo_name': '박테스트', 'founded_at': '2024-01-01',
               'description': 'AI 서비스', 'team_members': [], 'pricing_items': []}
    r = client.post('/projects', data={'payload': json.dumps(payload)})
    assert r.status_code == 201, r.text
    return r.json()['project_id']


def _pictures_in(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    return doc, len(doc.inline_shapes)


def test_docx_has_both_pictures_after_the_summary_table_in_order(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, *BOTH)

    r = authed_client.get(f'/projects/{pid}/plan-document.docx')

    assert r.status_code == 200, r.text
    doc, count = _pictures_in(r.content)
    assert count == 2
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs.index('□ 창업 아이템 개요(요약)') < paragraphs.index('< 서비스 흐름도 >')
    assert paragraphs.index('< 서비스 흐름도 >') < paragraphs.index('< 서비스 구조도 >')
    assert paragraphs.index('< 서비스 구조도 >') < paragraphs.index('1. 문제 인식(Problem)_창업 아이템의 필요성')


def test_docx_without_diagrams_has_no_picture_and_no_caption(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch)

    doc, count = _pictures_in(authed_client.get(f'/projects/{pid}/plan-document.docx').content)

    assert count == 0 and not [p for p in doc.paragraphs if p.text in ('< 서비스 흐름도 >', '< 서비스 구조도 >')]


def test_docx_is_still_made_when_one_diagram_file_is_missing(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, *BOTH, files={'userflow.svg': USER_FLOW_SVG})

    r = authed_client.get(f'/projects/{pid}/plan-document.docx')

    assert r.status_code == 200 and _pictures_in(r.content)[1] == 1


def test_zip_has_the_pictures_in_the_docx_and_the_original_svgs(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, *BOTH)

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 200, r.text
    archive = zipfile.ZipFile(io.BytesIO(r.content))
    assert archive.namelist() == ['plan/business-plan.docx', 'plan/diagrams/userflow.svg', 'plan/diagrams/architecture.svg']
    assert archive.read('plan/diagrams/userflow.svg') == USER_FLOW_SVG
    assert _pictures_in(archive.read('plan/business-plan.docx'))[1] == 2


# ── hwp: 맨 뒤 새 쪽에 그림 ──────────────────────────────────────────────────────────────
def test_hwp_diagram_page_is_a_new_last_page_with_pictures_stacked_downward(monkeypatch):
    commands = []
    pictures = diagram_pictures(_gateway(), 1, _plan_doc(*BOTH))
    answers = iter([json.dumps({'paragraphCount': 86, 'pageCount': 8}), '', json.dumps({'pageCount': 9}), '', ''])

    def fake_run(args):
        commands.append(args)
        return next(answers)

    monkeypatch.setattr(hwp_export, '_run_rhwp', fake_run)
    tmp = iter(f'/tmp/{i}.hwp' for i in range(10))

    result = hwp_export._append_diagram_page('/tmp/in.hwp', pictures, lambda: next(tmp))

    assert commands[0][:2] == ['explain', '/tmp/in.hwp']
    page_break = commands[1]
    assert page_break[:2] == ['edit', 'insert-page-break'] and page_break[page_break.index('--para') + 1] == '85'
    first, second = commands[3], commands[4]
    assert first[:2] == ['edit', 'insert-image'] and first[first.index('--page') + 1] == '8' == second[second.index('--page') + 1]
    first_y, second_y = int(first[first.index('--y') + 1]), int(second[second.index('--y') + 1])
    first_height = int(first[first.index('--height') + 1])
    assert first_y == hwp_export._PICTURE_TOP and second_y == first_y + first_height + hwp_export._PICTURE_GAP
    assert first[first.index('--x') + 1] == str(hwp_export._PICTURE_X)
    assert first[first.index('--width') + 1] == str(hwp_export._PICTURE_WIDTH)
    assert result == '/tmp/2.hwp'  # 쪽 나눔 파일(0번) 뒤로 그림마다 새 파일(1 · 2번) — 마지막 파일을 돌려준다


_RHWP_AVAILABLE = shutil.which(hwp_export.RHWP_BIN) is not None or os.path.isfile(hwp_export.RHWP_BIN)


def _page_count(hwp_bytes, tmp_path):
    path = tmp_path / 'out.hwp'
    path.write_bytes(hwp_bytes)
    out = subprocess.run([hwp_export.RHWP_BIN, 'explain', str(path), '--json'], capture_output=True, text=True, encoding='utf-8', check=True)
    return json.loads(out.stdout)['pageCount']


@pytest.mark.skipif(not _RHWP_AVAILABLE, reason='rhwp 실행 파일이 없다(RHWP_BIN)')
@pytest.mark.parametrize('applicant_type', ['preliminary', 'corp'])
def test_real_hwp_gets_one_more_page_with_the_pictures(authed_client, db_session, tmp_path, applicant_type):
    pid = _create(authed_client, applicant_type)
    data, template = _build_plan_document_data(db_session, db_session.get(Project, pid), PlanTexts())
    plain = hwp_export.render_plan_hwp(data, template)

    data.그림 = diagram_pictures(_gateway(), pid, _plan_doc(*BOTH))
    with_pictures = hwp_export.render_plan_hwp(data, template)

    assert _page_count(with_pictures, tmp_path) == _page_count(plain, tmp_path) + 1
