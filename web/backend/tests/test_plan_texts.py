"""계획서(planDoc) 항목 문장 → 공식 양식 글 칸 (SB-339 — 본문 채우기).

새 계획서는 양식의 모든 항목을 항목 번호(예비창업자 2.1.1 ~ 2.7.4, 개인사업자 · 법인 3.1.1 ~ 3.7.4)로 담는다. 뒤 두 자리(a.b)가 별첨1 구조와
맞는다 — 담당자 section_mapping.json과 항목 제목이 근거(app/plan_texts.py 설명). 표 항목의 서술은 표를 글자로 옮긴 줄 · 단계 미지정 원본 JSON을 빼고 넣는다."""
import io
import json
from types import SimpleNamespace as NS

from docx import Document
from orch_fakes import make_outputs, make_plan_doc, make_section, make_table_fallback

from app.hwp_export import _body_text
from app.plan_texts import PlanTexts, plan_texts, table_narrative

TABLE_TEXT = ['구분 | 추진내용 | 추진기간 | 세부내용', '개발 | 모델 고도화 | 2026.05~2026.07 | 정확도 80%']


def _items(prefix: str = '2', **extra: list[str]):
    """prefix.a.b 항목들 — 값은 그 항목의 문단들(줄). extra로 항목을 더하거나 바꾼다(예: {'5.2': [...]} )."""
    contents = {
        '1.1': ['초안-일반현황'], '1.2': ['초안-개요'], '1.3': ['초안-문제'], '1.4': ['초안-실현'], '1.5': ['초안-성장'], '1.6': ['초안-팀'],
        '2.1': ['창업아이템명 한 줄'], '2.2': ['산출물 목록 문장'],
        '3.1': ['아이템 개요 문장'], '3.2': ['요약 문제인식'], '3.3': ['요약 실현가능성'], '3.4': ['요약 성장전략'], '3.5': ['요약 팀구성'],
        '4.1': ['문제인식 본문 1', '문제인식 본문 2'],
        '5.1': ['실현가능성 본문'], '5.2': [*TABLE_TEXT, '', '개발계획 서술:', '1. 모델 고도화를 추진한다.'],
        '6.1': ['성장전략 본문'], '6.2': [*TABLE_TEXT, '', '개발계획 서술:', '1. 사업화를 추진한다.'],
        '7.1': ['대표자 역량 문장'], '7.2': ['장비 문장'], '7.3': ['채용 역량 문장'], '7.4': ['협력 역량 문장'],
    }
    contents.update(extra)
    return [make_section(f'{prefix}.{number}', f'{prefix}.{number} 제목', *lines) for number, lines in contents.items()]


def _plan(prefix: str = '2', **extra):
    return make_plan_doc(_items(prefix, **extra))


# ── 항목 번호 → 양식 칸 ──────────────────────────────────────────────────────────────────
def test_item_numbers_map_to_the_form_fields_for_both_forms():
    for prefix in ('2', '3'):
        texts = plan_texts(_plan(prefix))
        assert texts.by_item_number is True
        assert texts.산출물 == '산출물 목록 문장'
        assert texts.아이템_개요 == '아이템 개요 문장'
        assert (texts.요약_문제인식, texts.요약_실현가능성, texts.요약_성장전략, texts.요약_팀구성) == (
            '요약 문제인식', '요약 실현가능성', '요약 성장전략', '요약 팀구성')
        assert texts.문제인식_본문 == '문제인식 본문 1\n문제인식 본문 2'
        assert texts.실현가능성_본문 == '실현가능성 본문'
        assert texts.성장전략_본문 == '성장전략 본문'


def test_first_group_drafts_are_not_used():
    texts = plan_texts(_plan())
    everything = ' '.join(str(v) for v in vars(texts).values())
    assert '초안-' not in everything


def test_company_name_and_team_roster_are_not_taken_from_the_plan():
    """창업아이템명(x.2.1)은 웹이 받은 값이라 여기서 다루지 않는다."""
    texts = plan_texts(_plan())
    assert '창업아이템명 한 줄' not in ' '.join(str(v) for v in vars(texts).values())


def test_team_body_joins_the_four_items_with_labels_in_order():
    texts = plan_texts(_plan())
    assert texts.팀구성_본문 == (
        '[대표자 보유 역량]\n대표자 역량 문장\n[구매 후 보유 예정 장비 · 시설]\n장비 문장\n'
        '[추가 채용 시 보유 예정 역량]\n채용 역량 문장\n[협력기관 협업 시 보유 예정 기술 · 역량]\n협력 역량 문장')


def test_team_body_skips_missing_items_and_is_none_when_all_are_missing():
    only_two = make_plan_doc([s for s in _items() if s.section_code in ('2.7.1', '2.7.4')])
    assert plan_texts(only_two).팀구성_본문 == '[대표자 보유 역량]\n대표자 역량 문장\n[협력기관 협업 시 보유 예정 기술 · 역량]\n협력 역량 문장'
    none = make_plan_doc([s for s in _items() if not s.section_code.startswith('2.7.')])
    assert plan_texts(none).팀구성_본문 is None


def test_missing_items_are_none_so_the_builder_can_fall_back():
    texts = plan_texts(make_plan_doc([s for s in _items() if s.section_code in ('2.4.1',)]))
    assert texts.문제인식_본문 is not None
    assert texts.산출물 is None and texts.요약_문제인식 is None and texts.성장전략_본문 is None


# ── 표 항목의 서술 ───────────────────────────────────────────────────────────────────────
def test_table_narrative_drops_the_table_text_and_keeps_the_numbered_sentences():
    text = '\n'.join([*TABLE_TEXT, '', '개발계획 서술:', '1. 첫째.', '2. 둘째.'])
    assert table_narrative(text) == '1. 첫째.\n2. 둘째.'
    budget = '\n'.join(['비목 | 산출근거 | 정부지원사업비', '재료비 | 라벨링 | 5,000,000원', '', '사업비 집행계획 서술:', '1. 재료비: 라벨링.'])
    assert table_narrative(budget) == '1. 재료비: 라벨링.'


def test_table_narrative_drops_the_unassigned_original_rows_appendix():
    text = '\n'.join([*TABLE_TEXT, '', '사업비 집행계획 서술:', '1. 재료비.', '단계 미지정 원본 (단계별 합계에 미포함):', '[', '  {"비목": "x"}', ']'])
    assert table_narrative(text) == '1. 재료비.'


def test_table_narrative_keeps_a_fallback_text_and_drops_a_table_only_text():
    fallback = '표 구조 검증에 실패하여 본문 설명으로 대체함.\n· 구분: 개발; 추진내용: 모델'
    assert table_narrative(fallback) == fallback
    assert table_narrative('\n'.join(TABLE_TEXT)) == ''
    assert table_narrative('') == ''


def test_budget_narrative_is_split_by_phase_for_preliminary():
    budget1 = ['비목 | 산출근거 | 정부지원사업비', '재료비 | 라벨링 | 5,000,000원', '', '사업비 집행계획 서술:', '1. 1단계 문장.']
    budget2 = ['비목 | 산출근거 | 정부지원사업비', '광고 | SNS | 1,000원', '', '사업비 집행계획 서술:', '1. 2단계 문장.']
    texts = plan_texts(_plan('2', **{'5.3': budget1, '5.4': budget2}))
    assert texts.사업비_서술 == '1. 1단계 문장.' and texts.사업비_2단계_서술 == '1. 2단계 문장.'
    early = plan_texts(_plan('3', **{'5.3': budget1}))
    assert early.사업비_서술 == '1. 1단계 문장.' and early.사업비_2단계_서술 == ''


def test_schedule_narratives_come_from_the_two_schedule_tables():
    texts = plan_texts(_plan())
    assert texts.실현가능성_일정_서술 == '1. 모델 고도화를 추진한다.'
    assert texts.성장전략_일정_서술 == '1. 사업화를 추진한다.'


# ── 옛 계획서(4개 태그) · 계획서 없음 ───────────────────────────────────────────────────────
def test_legacy_four_tag_plan_is_read_by_tag():
    legacy = make_plan_doc([make_section('1-1', '문제', '문제 본문'), make_section('2-1', '실현', '실현 본문'), make_section('3-1', '성장', '성장 본문')])
    texts = plan_texts(legacy)
    assert texts.by_item_number is False
    assert (texts.문제인식_본문, texts.실현가능성_본문, texts.성장전략_본문) == ('문제 본문', '실현 본문', '성장 본문')
    assert (texts.요약_문제인식, texts.요약_실현가능성, texts.요약_성장전략) == ('문제 본문', '실현 본문', '성장 본문')
    assert texts.산출물 is None and texts.팀구성_본문 is None


def test_no_plan_gives_empty_texts():
    assert plan_texts(None) == PlanTexts()


# ── hwp: 서술은 그 절의 본문 뒤에 ───────────────────────────────────────────────────────────
def test_hwp_body_appends_the_table_narratives_to_their_section():
    data = NS(문제인식_본문='문제', 실현가능성_본문='실현\n둘째 줄', 성장전략_본문='성장', 팀구성_본문='팀',
              실현가능성_일정_서술='일정 서술', 사업비_서술='사업비 서술', 성장전략_일정_서술='전체 일정 서술')
    assert _body_text(data, '실현가능성_본문') == '실현 / 둘째 줄 / 일정 서술 / 사업비 서술'
    assert _body_text(data, '성장전략_본문') == '성장 / 전체 일정 서술'
    assert _body_text(data, '문제인식_본문') == '문제'
    assert _body_text(data, '팀구성_본문') == '팀'


def test_hwp_body_skips_empty_narratives():
    data = NS(실현가능성_본문='실현', 실현가능성_일정_서술='', 사업비_서술='')
    assert _body_text(data, '실현가능성_본문') == '실현'


# ── .docx 내려받기 ────────────────────────────────────────────────────────────────────────
def _payload(**overrides):
    base = {'biz_type': '개인', 'ceo_name': '박테스트', 'description': 'AI 기반 동네 헬스장 통합 예약 서비스',
            'output_summary': 'WEB산출물 입력값', 'team_members': [{'name': '박테스트', 'role': '대표', 'experience': '4년'}],
            'pricing_items': []}
    base.update(overrides)
    return {'payload': json.dumps(base)}


def _docx(client, project_id):
    r = client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200, r.text
    doc = Document(io.BytesIO(r.content))
    return '\n'.join(p.text for p in doc.paragraphs), '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)


def test_docx_is_filled_from_item_numbers_for_the_early_startup_form(authed_client, orch):
    project_id = authed_client.post('/projects', data=_payload()).json()['project_id']
    orch.responses['outputs'] = lambda pid: make_outputs(plan_doc=_plan('3'))
    paragraphs, tables = _docx(authed_client, project_id)

    assert '문제인식 본문 1' in paragraphs and '문제인식 본문 2' in paragraphs
    assert '실현가능성 본문' in paragraphs and '성장전략 본문' in paragraphs
    assert '[대표자 보유 역량]' in paragraphs and '협력 역량 문장' in paragraphs
    assert '산출물 목록 문장' in tables and 'WEB산출물 입력값' not in tables  # 산출물은 담당자가 짓는 값
    assert '아이템 개요 문장' in tables and '요약 문제인식' in tables and '요약 팀구성' in tables
    assert '초안-' not in paragraphs and '초안-' not in tables  # 첫 묶음(x.1.*)은 쓰지 않는다


def test_docx_puts_the_table_narratives_after_the_tables_without_the_table_text_lines(authed_client, orch):
    project_id = authed_client.post('/projects', data=_payload()).json()['project_id']
    orch.responses['outputs'] = lambda pid: make_outputs(plan_doc=_plan('3', **{
        '5.3': ['비목 | 산출근거 | 정부지원사업비', '재료비 | 라벨링 | 5,000,000원', '', '사업비 집행계획 서술:', '1. 재료비는 라벨링에 쓴다.']}))
    paragraphs, _tables = _docx(authed_client, project_id)

    assert '1. 모델 고도화를 추진한다.' in paragraphs and '1. 사업화를 추진한다.' in paragraphs
    assert '1. 재료비는 라벨링에 쓴다.' in paragraphs
    assert ' | ' not in paragraphs  # 표를 글자로 옮긴 줄은 문서에 나오지 않는다
    assert '개발계획 서술:' not in paragraphs and '사업비 집행계획 서술:' not in paragraphs


def test_docx_falls_back_to_the_web_output_summary_when_the_plan_has_no_item(authed_client, orch):
    """항목 x.2.2가 없는 계획서(옛 4개 태그 모양)는 웹 산출물 입력으로 돌아간다."""
    project_id = authed_client.post('/projects', data=_payload()).json()['project_id']
    orch.responses['outputs'] = lambda pid: make_outputs(plan_doc=make_plan_doc([make_section('1-1', '문제', '옛 문제 본문')]))
    paragraphs, tables = _docx(authed_client, project_id)
    assert 'WEB산출물 입력값' in tables
    assert '옛 문제 본문' in paragraphs


# ── 표를 본문 서술로 바꾼 항목 (SB-354) ───────────────────────────────
FALLBACK_LINE = '· 구분: 개발; 추진내용: 모델 고도화; 추진기간: 2026.05~2026.07; 세부내용: 정확도 80%'


def test_a_replaced_table_item_gives_no_narrative_because_the_table_is_drawn_from_the_input_rows():
    plan = _plan('2', **{'5.2': [FALLBACK_LINE], '5.3': ['· 비목: 외주용역비'], '5.4': ['1단계 서술'], '6.2': ['성장 서술']})
    texts = plan_texts(plan, [make_table_fallback('2.5.2'), make_table_fallback('2.5.3')])

    assert texts.실현가능성_일정_서술 == '' and texts.사업비_서술 == ''
    assert texts.사업비_2단계_서술 == '1단계 서술' and texts.성장전략_일정_서술 == '성장 서술'  # 바꾸지 않은 표의 서술은 그대로
    assert texts.실현가능성_본문 == '실현가능성 본문'  # 본문 항목은 건드리지 않는다


def test_without_replaced_tables_the_narrative_is_kept():
    plan = _plan('2', **{'5.2': [FALLBACK_LINE]})
    assert plan_texts(plan, []).실현가능성_일정_서술 == FALLBACK_LINE
    assert plan_texts(plan).실현가능성_일정_서술 == FALLBACK_LINE


def test_replacing_every_table_item_does_not_turn_the_plan_into_a_legacy_plan():
    only_tables = make_plan_doc([make_section('2.5.2', '표', FALLBACK_LINE), make_section('2.4.1', '문제', '문제 본문')])
    texts = plan_texts(only_tables, [make_table_fallback('2.5.2')])

    assert texts.by_item_number is True and texts.문제인식_본문 == '문제 본문'


def test_a_fallback_with_an_unreadable_item_number_is_ignored():
    plan = _plan('2', **{'5.2': [FALLBACK_LINE]})
    texts = plan_texts(plan, [make_table_fallback('없는번호'), make_table_fallback('')])

    assert texts.실현가능성_일정_서술 == FALLBACK_LINE


def test_the_downloaded_docx_keeps_the_input_table_and_drops_the_repeated_lines(authed_client, orch):
    payload = {'applicant_type': 'preliminary', 'biz_type': '개인', 'ceo_name': '박', 'description': 'AI', 'team_members': [],
               'pricing_items': [], 'budget_items': [], 'schedule_items': [
                   {'section': 'feasibility', 'category': '개발', 'content': '모델 고도화', 'period': '2026.05~2026.07', 'detail': '정확도 80%'}]}
    created = authed_client.post('/projects', data={'payload': json.dumps(payload)})
    assert created.status_code == 201, created.text
    pid = created.json()['project_id']
    plan = _plan('2', **{'5.2': [FALLBACK_LINE], '6.2': [FALLBACK_LINE]})

    orch.responses['outputs'] = lambda p: make_outputs(plan_doc=plan, table_fallbacks=[make_table_fallback('2.5.2')])
    doc = Document(io.BytesIO(authed_client.get(f'/projects/{pid}/plan-document.docx').content))
    lines = [p.text for p in doc.paragraphs]

    assert lines.count(FALLBACK_LINE) == 1  # 바꾼 5.2의 줄은 빠지고, 바꾸지 않은 6.2의 줄만 남는다
    assert any(c.text == '2026.05~2026.07' for table in doc.tables for row in table.rows for c in row.cells)  # 표는 입력 행 그대로
