"""GET /projects/{id}/deliverable.zip — 최종 결과물 묶음 내려받기 (SB-316).

압축은 웹이 만든다: 프로토타입 · 인포그래픽은 결과물(deliverable)의 파일을 참조마다 read_artifact_file로 받아 묶고, 계획서는 파일 참조(자리채움 파일이라
쓰지 않는다)가 아니라 계획서(planDoc)로 웹이 만든 .docx를 넣는다(SB-341).

실행:
    pytest tests/test_deliverable_zip.py -v
"""
import io
import json
import zipfile
from types import SimpleNamespace as NS

import pytest
from docx import Document
from orch_fakes import make_outputs, make_plan_doc, make_section

from app.models import User
from app.orch import OrchError
from app.routers import artifact_files


def _ref(name: str, size: int = 10, execution: str = 'e' * 12):
    return NS(key=f'{"r" * 12}/{execution}/{"u" * 8}/{name}', name=name, size=size)


FILES = {  # key → (파일 이름, 내용) — 자리채움 계획서 파일(plan.docx)은 일부러 뺐다: 읽으면 KeyError라 쓰지 않는지 확인된다
    _ref('index.html').key: ('index.html', '<html>실행 파일</html>'.encode()),
    _ref('README.md').key: ('README.md', b'# readme'),
    _ref('infographic.png').key: ('infographic.png', bytes(range(200))),
}


def _deliverable(**overrides):
    base = dict(
        plan_doc_file=_ref('plan.docx'),
        prototype_files=[_ref('index.html'), _ref('README.md')],
        infographic_file=_ref('infographic.png'))
    base.update(overrides)
    return NS(**base)


def _plan_doc():
    """새 계획서 모양(항목 번호) — 개인사업자 · 법인 양식(3.x.x)."""
    return make_plan_doc([
        make_section('3.2.2', '산출물', '압축 안 계획서의 산출물 문장'),
        make_section('3.4.1', '문제인식', '압축 안 계획서의 문제인식 본문'),
    ])


def _read(project_id, key):
    name, data = FILES[key]
    return NS(name=name, media_type='application/octet-stream', size=len(data), sha256='0' * 64, data=data)


def _create(client) -> int:
    r = client.post('/projects', data={'payload': json.dumps({'description': '압축 테스트'})})
    assert r.status_code == 201, r.text
    return r.json()['project_id']


def _ready(orch, deliverable, plan_doc='default'):
    plan_doc = _plan_doc() if plan_doc == 'default' else plan_doc
    orch.responses['outputs'] = lambda project_id: make_outputs(deliverable=deliverable, plan_doc=plan_doc)
    orch.responses['read_artifact_file'] = _read


def _zip(response):
    return zipfile.ZipFile(io.BytesIO(response.content))


def test_owner_gets_a_zip_with_the_plan_prototype_and_infographic(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 200, r.text
    assert r.headers['content-type'] == 'application/zip'
    assert r.headers['content-disposition'] == f'attachment; filename="deliverable-{pid}.zip"'
    assert r.headers['cache-control'] == 'private, no-store' and r.headers['x-content-type-options'] == 'nosniff'
    archive = _zip(r)
    assert archive.namelist() == ['plan/business-plan.docx', 'prototype/index.html', 'prototype/README.md', 'infographic/infographic.png']
    assert archive.read('prototype/index.html').decode('utf-8') == '<html>실행 파일</html>'
    assert archive.read('infographic/infographic.png') == bytes(range(200))
    assert archive.testzip() is None


def test_plan_in_the_zip_is_a_docx_built_from_the_plan_doc(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    archive = _zip(authed_client.get(f'/projects/{pid}/deliverable.zip'))

    data = archive.read('plan/business-plan.docx')
    assert data[:2] == b'PK'
    doc = Document(io.BytesIO(data))
    paragraphs = '\n'.join(p.text for p in doc.paragraphs)
    tables = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert '압축 안 계획서의 문제인식 본문' in paragraphs  # x.4.1 → 문제인식 본문
    assert '압축 안 계획서의 산출물 문장' in tables         # x.2.2 → 일반현황의 산출물
    assert '초기창업패키지' in paragraphs


def test_the_placeholder_plan_file_in_the_deliverable_is_never_read(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable(plan_doc_file=_ref('plan.docx')))

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 200, r.text
    assert _ref('plan.docx').key not in [c[1][1] for c in orch.calls if c[0] == 'read_artifact_file']


def test_each_file_is_read_by_its_reference_key(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    authed_client.get(f'/projects/{pid}/deliverable.zip')

    reads = [c for c in orch.calls if c[0] == 'read_artifact_file']
    assert [c[1] for c in reads] == [
        (pid, _ref('index.html').key), (pid, _ref('README.md').key), (pid, _ref('infographic.png').key)]


def test_missing_optional_files_are_left_out(authed_client, orch):
    """진입 파일 · 안내 문서가 없으면(프로토타입 목록이 비어 있으면) 그 폴더는 만들지 않는다."""
    pid = _create(authed_client)
    _ready(orch, _deliverable(prototype_files=[]))

    archive = _zip(authed_client.get(f'/projects/{pid}/deliverable.zip'))

    assert archive.namelist() == ['plan/business-plan.docx', 'infographic/infographic.png']


def test_duplicate_names_in_one_folder_get_a_number(authed_client, orch):
    pid = _create(authed_client)
    other = _ref('README.md', execution='f' * 12)
    FILES[other.key] = ('README.md', b'# second')
    try:
        _ready(orch, _deliverable(prototype_files=[_ref('README.md'), other]))

        archive = _zip(authed_client.get(f'/projects/{pid}/deliverable.zip'))

        assert archive.namelist()[1:3] == ['prototype/README.md', 'prototype/README_2.md']
        assert archive.read('prototype/README_2.md') == b'# second'
    finally:
        del FILES[other.key]


def test_no_deliverable_yet_is_404(authed_client, orch):
    pid = _create(authed_client)
    orch.responses['outputs'] = lambda project_id: make_outputs(deliverable=None)

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 404 and r.json()['code'] == 'DELIVERABLE_NOT_READY'
    assert not [c for c in orch.calls if c[0] == 'read_artifact_file']


def test_deliverable_without_a_plan_doc_is_404(authed_client, orch):
    """계획서(planDoc)가 없으면 계획서 없는 압축을 주지 않는다."""
    pid = _create(authed_client)
    _ready(orch, _deliverable(), plan_doc=None)

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 404 and r.json()['code'] == 'DELIVERABLE_NOT_READY'


def test_too_large_to_bundle_is_413_and_nothing_is_read(authed_client, orch, monkeypatch):
    pid = _create(authed_client)
    monkeypatch.setattr(artifact_files, 'MAX_ZIP_SOURCE_BYTES', 20)
    _ready(orch, _deliverable(prototype_files=[_ref('index.html', size=10)], infographic_file=_ref('infographic.png', size=15)))

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 413 and r.json()['code'] == 'DELIVERABLE_TOO_LARGE'
    assert not [c for c in orch.calls if c[0] == 'read_artifact_file']


def test_a_file_that_cannot_be_read_fails_the_whole_zip(authed_client, orch):
    """일부만 든 압축을 주지 않는다."""
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    def flaky(project_id, key):
        if key.endswith('README.md'):
            raise OrchError('FILE_NOT_FOUND')
        return _read(project_id, key)

    orch.responses['read_artifact_file'] = flaky

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 404 and r.json()['code'] == 'FILE_NOT_FOUND'


@pytest.mark.parametrize('code, status', [('RUN_NOT_FOUND', 404), ('RUN_NOT_VIEWABLE', 409)])
def test_orchestrator_errors_on_outputs_become_http_errors(authed_client, orch, code, status):
    pid = _create(authed_client)
    orch.responses['outputs'] = OrchError(code)

    assert authed_client.get(f'/projects/{pid}/deliverable.zip').status_code == status


def test_other_users_get_the_same_404_and_the_orchestrator_is_not_asked(login_as, orch):
    owner = login_as('owner@example.com')
    pid = _create(owner)
    _ready(orch, _deliverable())
    calls_before = len(orch.calls)

    other = login_as('other@example.com')
    r = other.get(f'/projects/{pid}/deliverable.zip')

    assert r.status_code == 404 and r.json()['code'] == 'FILE_NOT_FOUND'
    assert len([c for c in orch.calls if c[0] in ('outputs', 'read_artifact_file')]) == 0
    assert len(orch.calls) >= calls_before


def test_admin_cannot_download_someone_elses_zip(login_as, db_session, orch):
    owner = login_as('owner@example.com')
    pid = _create(owner)
    _ready(orch, _deliverable())
    admin = login_as('admin@example.com')
    db_session.query(User).filter(User.email == 'admin@example.com').update({'role': 'admin'})
    db_session.commit()

    assert admin.get(f'/projects/{pid}/deliverable.zip').status_code == 404


def test_requires_login(client, orch):
    assert client.get('/projects/1/deliverable.zip').status_code == 401
