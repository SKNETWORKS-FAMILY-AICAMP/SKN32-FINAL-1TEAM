"""계획서 파일에 인포그래픽 · 첫 장 고지 넣기 (기준 문서 T-C4, SB-354).

기준 문서는 계획서 파일에 결과물의 인포그래픽(`deliverable.infographicFile`, SVG)과 첫 장 고지(`deliverable.disclaimer`)를 넣도록 정했다.
웹이 만드는 .docx · .hwp(와 결과물 .zip 안의 .docx)가 그렇게 한다. 인포그래픽 SVG는 계획서 그림과 달리 데이터 주소(`data:`) 이미지 · 같은
파일 안 참조를 쓰므로 그 둘은 허용하고 바깥 파일 · 주소를 부르는 것은 막는다. 인포그래픽을 못 받거나 못 바꾸면 그것만 빼고 문서는 만든다.
hwp 실제 확인은 rhwp 실행 파일이 있을 때만 한다.

실행:
    pytest tests/test_infographic_in_document.py -v
"""
import io
import json
import os
import shutil
import subprocess
import zipfile
from types import SimpleNamespace as NS

import pytest
from docx import Document
from docx.shared import Cm
from orch_fakes import make_file_ref, make_outputs, make_plan_doc, make_section

from app import diagram_images, hwp_export
from app.diagram_images import (
    MAX_INFOGRAPHIC_SVG_BYTES,
    document_pictures,
    infographic_picture,
    svg_to_png,
)
from app.models import Project
from app.orch import OrchError
from app.plan_document_export import DEFAULT_DISCLAIMER
from app.plan_texts import PlanTexts
from app.routers.projects import _build_plan_document_data

_PIXEL = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=='


def _infographic_svg(height: int = 1400, extra: str = '') -> bytes:
    """실제 인포그래픽이 쓰는 요소를 닮은 SVG — 글꼴 · 이미지는 데이터 주소, 이미지는 defs에 두고 <use>로 부른다."""
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" viewBox="0 0 1000 {height}" width="1000" height="{height}">'
        f'<style>@font-face{{font-family:"Pretendard";src:url(data:font/woff2;base64,{_PIXEL}) format("woff2");}}</style>'
        f'<defs><image id="art" width="40" height="40" href="data:image/png;base64,{_PIXEL}"/></defs>'
        f'<rect width="1000" height="{height}" fill="#f4f6fb"/><use href="#art" x="40" y="40"/>'
        f'<text x="100" y="120" font-size="48" font-family="Pretendard, sans-serif">동네 헬스장 통합 예약</text>{extra}</svg>'
    ).encode()


INFOGRAPHIC_SVG = _infographic_svg()
DISCLAIMER = '본 문서는 S-Brain이 생성한 초안입니다. 제출 전 작성자 본인의 확인과 수정이 필요합니다.(결과물 값)'


def _ref(name='infographic.svg', **overrides):
    return make_file_ref(name, key=f'r1/e1/uniq/{name}', **overrides)


def _gateway(files=None):
    files = {'infographic.svg': INFOGRAPHIC_SVG} if files is None else files

    def read(project_id, key):
        name = key.rsplit('/', 1)[-1]
        if name not in files:
            raise OrchError('FILE_NOT_FOUND')
        return NS(name=name, data=files[name], media_type='image/svg+xml')

    return NS(read_artifact_file=read)


def _deliverable(**overrides):
    base = dict(plan_doc_file=make_file_ref('plan.docx'), prototype_files=[], infographic_file=_ref(), disclaimer=DISCLAIMER)
    base.update(overrides)
    return NS(**base)


# ── SVG → PNG: 인포그래픽은 데이터 주소 · 안쪽 참조를 허용하고 그림은 그대로 엄격하다 ─────────────────────
def test_an_infographic_with_embedded_data_and_inner_references_becomes_a_png():
    png = svg_to_png(INFOGRAPHIC_SVG, allow_embedded=True)
    assert png[:8] == b'\x89PNG\r\n\x1a\n'


def test_diagrams_stay_strict_and_still_reject_embedded_images():
    with pytest.raises(ValueError):
        svg_to_png(INFOGRAPHIC_SVG)  # allow_embedded=False(계획서 그림)


@pytest.mark.parametrize('extra', [
    '<image href="http://example.com/a.png" width="5" height="5"/>',
    '<image xlink:href="https://example.com/a.png" width="5" height="5"/>',
    '<image href="file:///etc/passwd" width="5" height="5"/>',
    '<image href="other.svg#a" width="5" height="5"/>',
    '<rect width="5" height="5" fill="url(http://example.com/a.svg#g)"/>',
    '<style>@import url(http://example.com/a.css);</style>',
    '<foreignObject width="5" height="5"/>',
    '<script>alert(1)</script>',
])
def test_an_infographic_with_an_outside_reference_is_not_drawn(extra):
    with pytest.raises(ValueError):
        svg_to_png(_infographic_svg(extra=extra), allow_embedded=True)


# ── 결과물 → 그림 ───────────────────────────────────────────────────────────────
def test_the_infographic_picture_is_marked_and_keeps_the_original_svg():
    picture = infographic_picture(_gateway(), 1, _deliverable())
    assert picture.is_infographic is True and picture.title == '인포그래픽'
    assert picture.svg == INFOGRAPHIC_SVG and picture.svg_name == 'infographic.svg'
    assert picture.png[:4] == b'\x89PNG'


def test_no_deliverable_or_no_infographic_gives_nothing():
    assert infographic_picture(_gateway(), 1, None) is None
    assert infographic_picture(_gateway(), 1, _deliverable(infographic_file=None)) is None


def test_an_infographic_that_cannot_be_used_is_left_out(caplog):
    assert infographic_picture(_gateway({}), 1, _deliverable()) is None  # 받지 못함
    assert infographic_picture(_gateway({'infographic.svg': b'<svg'}), 1, _deliverable()) is None  # 깨진 SVG
    assert infographic_picture(_gateway({'infographic.svg': _infographic_svg(extra='<image href="http://x/y.png"/>')}), 1, _deliverable()) is None
    assert 'FILE_NOT_FOUND' in caplog.text


def test_an_oversized_infographic_is_not_even_read():
    calls = []
    gateway = NS(read_artifact_file=lambda *a: calls.append(a))
    big = _deliverable(infographic_file=_ref(size=MAX_INFOGRAPHIC_SVG_BYTES + 1))
    assert infographic_picture(gateway, 1, big) is None and calls == []


def test_the_infographic_comes_after_the_plan_diagrams():
    plan = make_plan_doc([make_section('3.4.1', '문제인식', '본문')])
    pictures = document_pictures(_gateway(), 1, plan, _deliverable())
    assert [p.is_infographic for p in pictures] == [True]  # 계획서 그림이 없으면 인포그래픽만
    assert document_pictures(_gateway(), 1, plan, None) == []


# ── docx ─────────────────────────────────────────────────────────────────────
def _create(client, applicant_type='corp'):
    payload = {'applicant_type': applicant_type, 'biz_type': '개인', 'ceo_name': '박테스트', 'description': 'AI 서비스', 'team_members': [],
               'pricing_items': []}
    if applicant_type == 'corp':
        payload.update(biz_type='법인', founded_at='2024-01-01')
    r = client.post('/projects', data={'payload': json.dumps(payload)})
    assert r.status_code == 201, r.text
    return r.json()['project_id']


def _ready(orch, deliverable, files=None):
    plan = make_plan_doc([make_section('3.4.1', '문제인식', '본문'), make_section('3.2.2', '산출물', '산출물 문장')])
    orch.responses['outputs'] = lambda project_id: make_outputs(plan_doc=plan, deliverable=deliverable)
    orch.responses['read_artifact_file'] = _gateway(files).read_artifact_file


def _docx(client, pid):
    r = client.get(f'/projects/{pid}/plan-document.docx')
    assert r.status_code == 200, r.text
    return Document(io.BytesIO(r.content))


def test_docx_has_the_infographic_with_its_caption_after_the_summary_table(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    doc = _docx(authed_client, pid)

    assert '< 인포그래픽 >' in [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 1


def test_a_tall_infographic_is_scaled_down_to_fit_a_page(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable(), files={'infographic.svg': _infographic_svg(height=6000)})

    shape = _docx(authed_client, pid).inline_shapes[0]

    assert shape.height <= Cm(22) + 100 and shape.width < Cm(16)  # 세로 상한 22cm, 비율을 지켜 가로도 줄어든다


def test_docx_first_page_uses_the_deliverable_disclaimer(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    lines = [p.text for p in _docx(authed_client, pid).paragraphs]

    assert DISCLAIMER in lines
    assert not any('초안입니다' in line and line != DISCLAIMER for line in lines)  # 안내 문장 끝에 붙던 옛 한 줄은 없다


def test_docx_uses_the_default_disclaimer_before_the_deliverable_exists(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, None)

    doc = _docx(authed_client, pid)

    assert DEFAULT_DISCLAIMER in [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 0 and '< 인포그래픽 >' not in [p.text for p in doc.paragraphs]


def test_docx_is_still_made_when_the_infographic_cannot_be_used(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable(), files={})

    doc = _docx(authed_client, pid)

    assert len(doc.inline_shapes) == 0 and DISCLAIMER in [p.text for p in doc.paragraphs]


# ── zip ──────────────────────────────────────────────────────────────────────
def test_zip_plan_docx_has_the_infographic_and_the_svg_is_kept_only_in_the_infographic_folder(authed_client, orch):
    pid = _create(authed_client)
    _ready(orch, _deliverable())

    r = authed_client.get(f'/projects/{pid}/deliverable.zip')
    assert r.status_code == 200, r.text
    archive = zipfile.ZipFile(io.BytesIO(r.content))

    assert sorted(archive.namelist()) == ['infographic/infographic.svg', 'plan/business-plan.docx']
    assert archive.read('infographic/infographic.svg') == INFOGRAPHIC_SVG
    doc = Document(io.BytesIO(archive.read('plan/business-plan.docx')))
    assert '< 인포그래픽 >' in [p.text for p in doc.paragraphs] and DISCLAIMER in [p.text for p in doc.paragraphs]


# ── hwp ──────────────────────────────────────────────────────────────────────
_RHWP_AVAILABLE = shutil.which(hwp_export.RHWP_BIN) is not None or os.path.isfile(hwp_export.RHWP_BIN)


def _rhwp_json(args, path):
    out = subprocess.run([hwp_export.RHWP_BIN, *args, str(path), '--json'], capture_output=True, text=True, encoding='utf-8', check=True)
    return json.loads(out.stdout)


def _hwp(data, template, tmp_path, name):
    path = tmp_path / name
    path.write_bytes(hwp_export.render_plan_hwp(data, template))
    return path


@pytest.mark.skipif(not _RHWP_AVAILABLE, reason='rhwp 실행 파일이 없다(RHWP_BIN)')
@pytest.mark.parametrize('applicant_type', ['preliminary', 'corp'])
def test_real_hwp_has_the_disclaimer_on_the_first_page_and_the_infographic_on_a_new_page(authed_client, db_session, tmp_path, applicant_type):
    pid = _create(authed_client, applicant_type)
    data, template = _build_plan_document_data(db_session, db_session.get(Project, pid), PlanTexts())
    plain = _hwp(data, template, tmp_path, 'plain.hwp')

    data.고지 = DISCLAIMER
    data.인포그래픽 = infographic_picture(_gateway(), pid, _deliverable())
    full = _hwp(data, template, tmp_path, 'full.hwp')

    assert _rhwp_json(['explain'], full)['pageCount'] == _rhwp_json(['explain'], plain)['pageCount'] + 1
    first_page = _rhwp_json(['export-text'], full)['pages'][0]['text']
    assert DISCLAIMER in first_page and DISCLAIMER not in _rhwp_json(['export-text'], plain)['pages'][0]['text']


@pytest.mark.skipif(not _RHWP_AVAILABLE, reason='rhwp 실행 파일이 없다(RHWP_BIN)')
def test_real_hwp_puts_the_diagrams_and_the_infographic_on_two_separate_pages(authed_client, db_session, tmp_path):
    from test_diagram_images import BOTH, _plan_doc
    from test_diagram_images import _gateway as diagram_gateway

    pid = _create(authed_client, 'corp')
    data, template = _build_plan_document_data(db_session, db_session.get(Project, pid), PlanTexts())
    plain = _hwp(data, template, tmp_path, 'plain.hwp')
    data.그림 = diagram_images.diagram_pictures(diagram_gateway(), pid, _plan_doc(*BOTH))
    data.인포그래픽 = infographic_picture(_gateway(), pid, _deliverable())
    both = _hwp(data, template, tmp_path, 'both.hwp')
    data.인포그래픽 = None
    diagrams_only = _hwp(data, template, tmp_path, 'diagrams.hwp')

    plain_pages = _rhwp_json(['explain'], plain)['pageCount']
    assert _rhwp_json(['explain'], diagrams_only)['pageCount'] == plain_pages + 1
    assert _rhwp_json(['explain'], both)['pageCount'] == plain_pages + 2  # 그림 쪽 하나, 인포그래픽 쪽 하나


@pytest.mark.skipif(not _RHWP_AVAILABLE, reason='rhwp 실행 파일이 없다(RHWP_BIN)')
def test_real_hwp_without_a_deliverable_disclaimer_uses_the_default(authed_client, db_session, tmp_path):
    pid = _create(authed_client, 'corp')
    data, template = _build_plan_document_data(db_session, db_session.get(Project, pid), PlanTexts())

    full = _hwp(data, template, tmp_path, 'default.hwp')

    assert DEFAULT_DISCLAIMER in _rhwp_json(['export-text'], full)['pages'][0]['text']


def test_the_unused_helper_names_exist():
    assert callable(diagram_images.document_pictures)
