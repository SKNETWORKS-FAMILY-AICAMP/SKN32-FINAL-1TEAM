"""계획서 본문 줄의 글머리 — docx · hwp가 줄에 붙어 온 글머리 위에 " ◦ "를 또 붙이지 않는다.

예전에는 담당자 글의 "◦ 글" · "- 글" · "① 글" 줄에 문서가 " ◦ "를 더 붙여 " ◦ ◦ 글" · " ◦ - 글"이 됐고, docx는 줄이 넘어가면 둘째 줄이
글머리 앞(여백)부터 시작했다. 지금은 글머리를 떼고 단계(본문 · 하위)를 가려 한 번만 달고, docx는 내어쓰기로 넘어간 줄을 글 첫 글자에 맞춘다.

실행:
    pytest tests/test_plan_body_lines.py -v
"""
import io

import pytest
from docx import Document

from app import hwp_export
from app.plan_document_export import body_line, render_plan_docx
from tests.test_hwp_table_layout import _data


@pytest.mark.parametrize('line,expected', [
    ('본문 줄', ('item', '◦', '본문 줄')),
    ('◦ 본문 줄', ('item', '◦', '본문 줄')),
    ('◦ ◦ 겹친 글머리', ('item', '◦', '겹친 글머리')),
    ('• 다른 글머리', ('item', '◦', '다른 글머리')),
    ('* 마크다운 글머리', ('item', '◦', '마크다운 글머리')),
    ('- 하위 줄', ('sub', '-', '하위 줄')),
    ('· 가운뎃점 하위 줄', ('sub', '·', '가운뎃점 하위 줄')),
    ('– 긴 줄표 하위 줄', ('sub', '-', '긴 줄표 하위 줄')),
    ('① 번호 줄', ('item', '①', '번호 줄')),
    ('1. 번호 줄', ('item', '1.', '번호 줄')),
    ('가) 번호 줄', ('item', '가)', '번호 줄')),
    ('[대표자 보유 역량]', ('heading', '', '[대표자 보유 역량]')),
    ('**굵게** 쓴 글', ('item', '◦', '굵게 쓴 글')),
    ('2026.11 ~ 2027.02 개발', ('item', '◦', '2026.11 ~ 2027.02 개발')),   # 날짜는 번호가 아니다
    ('3.5배 성장', ('item', '◦', '3.5배 성장')),
    ('-5% 절감', ('item', '◦', '-5% 절감')),                               # 글머리 뒤에는 빈칸이 있어야 한다
    ('등록·조회', ('item', '◦', '등록·조회')),
])
def test_body_line_takes_off_the_mark_and_tells_the_level(line, expected):
    assert body_line(line) == expected


def test_body_line_without_plain_mark_leaves_unmarked_lines_plain():
    assert body_line('표 아래 서술', '') == ('plain', '', '표 아래 서술')
    assert body_line('1. 번호 서술', '') == ('item', '1.', '번호 서술')


def _paragraphs(data, template='early_general'):
    return Document(io.BytesIO(render_plan_docx(data, template=template))).paragraphs


def test_docx_body_has_one_mark_per_line_and_keeps_the_levels():
    body = '[소제목]\n◦ 본문 첫째\n◦ ◦ 겹친 글머리\n- 하위 줄\n① 번호 줄\n글머리 없는 줄\n* 마크다운 **굵게**'
    paragraphs = _paragraphs(_data(문제인식_본문=body))
    start = next(i for i, p in enumerate(paragraphs) if p.text.startswith('1. 문제 인식'))
    texts = [p.text for p in paragraphs[start + 1:start + 8]]

    assert texts == ['[소제목]', '◦\t본문 첫째', '◦\t겹친 글머리', '-\t하위 줄', '①\t번호 줄', '◦\t글머리 없는 줄', '◦\t마크다운 굵게']
    assert not any(t.startswith(' ◦') for t in texts)


def test_docx_body_lines_hang_so_wrapped_lines_start_under_the_text():
    paragraphs = _paragraphs(_data(문제인식_본문='[소제목]\n본문 줄\n- 하위 줄'))
    start = next(i for i, p in enumerate(paragraphs) if p.text.startswith('1. 문제 인식'))
    heading, item, sub = paragraphs[start + 1:start + 4]

    assert heading.runs[0].bold and heading.paragraph_format.left_indent is None   # 소제목은 굵게, 들여쓰지 않는다
    for paragraph, indent in ((item, 0.7), (sub, 1.3)):
        fmt = paragraph.paragraph_format
        assert round(fmt.left_indent.cm, 1) == indent
        assert round(fmt.first_line_indent.cm, 1) == -0.5            # 내어쓰기 — 첫 줄만 글머리 자리만큼 왼쪽으로
        assert [round(t.position.cm, 1) for t in fmt.tab_stops] == [indent]   # 글머리 뒤 탭이 글 첫 글자를 왼쪽 들여쓰기에 맞춘다
    assert sub.paragraph_format.left_indent > item.paragraph_format.left_indent


def test_docx_narrative_lines_keep_unnumbered_lines_plain():
    paragraphs = [p.text for p in _paragraphs(_data(성장전략_일정_서술='1. 첫째 서술.\n번호 없는 서술.'))]
    assert '1.\t첫째 서술.' in paragraphs and '번호 없는 서술.' in paragraphs


def _cell(document, label):
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text == label:
                return row.cells[1]
    raise AssertionError(label)


def test_docx_multi_line_cell_gets_one_kind_of_mark_per_line_and_header_cells_keep_their_line_break():
    document = Document(io.BytesIO(render_plan_docx(_data(산출물='◦ 개발 목표 산출물\n- 회원 관리 기능\n• 예약 기능', 아이템_개요='한 줄 개요'))))
    output = _cell(document, '산출물\n(협약기간 내 목표)')   # 제목 칸은 칸 안 줄바꿈 그대로

    assert [p.text for p in output.paragraphs] == ['◦\t개발 목표 산출물', '-\t회원 관리 기능', '◦\t예약 기능']
    assert [p.text for p in _cell(document, '아이템 개요').paragraphs] == ['한 줄 개요']   # 글머리 없는 한 줄 값은 그대로


def test_docx_multi_line_cell_does_not_add_a_mark_to_unmarked_lines():
    document = Document(io.BytesIO(render_plan_docx(_data(아이템_개요='개요 첫째\n개요 둘째'))))
    assert [p.text for p in _cell(document, '아이템 개요').paragraphs] == ['개요 첫째', '개요 둘째']


# ── hwp 본문 ─────────────────────────────────────────────────────
@pytest.mark.parametrize('line,expected', [
    ('본문 줄', ' ◦ 본문 줄'),
    ('◦ 본문 줄', ' ◦ 본문 줄'),
    ('• 다른 글머리', ' ◦ 다른 글머리'),
    ('- 하위 줄', '   - 하위 줄'),
    ('① 번호 줄', ' ① 번호 줄'),
    ('[소제목]', '[소제목]'),
])
def test_hwp_body_paragraph_has_one_mark(line, expected):
    assert hwp_export._body_paragraph(line) == expected


def test_hwp_body_replaces_the_first_bullet_when_the_first_line_has_its_own_mark(monkeypatch):
    commands = []
    monkeypatch.setattr(hwp_export, '_run_rhwp', lambda args: commands.append(args) or '')
    tmp = iter(f'/tmp/{i}.hwp' for i in range(500))
    data = _data(실현가능성_본문='- 하위 줄로 시작\n◦ 본문 줄')
    data.문제인식_본문 = data.팀구성_본문 = data.성장전략_본문 = ''

    hwp_export._write_body_sections('/tmp/in.hwp', 'preliminary', data, lambda: next(tmp))

    texts = [c[c.index('--text') + 1] for c in commands if c[1] == 'insert-text' and 26 <= int(c[c.index('--para') + 1]) <= 27]
    assert texts == [' ◦ 본문 줄', '   - 하위 줄로 시작']   # 마지막 줄부터 — 첫 줄은 양식의 " ◦ "를 지우고 통째로
    deletes = [c for c in commands if c[1] == 'delete-text']
    assert len(deletes) == 1 and deletes[0][deletes[0].index('--count') + 1] == '3'


def test_hwp_body_writes_after_the_template_bullet_when_the_first_line_is_a_plain_item(monkeypatch):
    commands = []
    monkeypatch.setattr(hwp_export, '_run_rhwp', lambda args: commands.append(args) or '')
    tmp = iter(f'/tmp/{i}.hwp' for i in range(500))
    data = _data(실현가능성_본문='◦ 첫 줄')
    data.문제인식_본문 = data.팀구성_본문 = data.성장전략_본문 = ''

    hwp_export._write_body_sections('/tmp/in.hwp', 'preliminary', data, lambda: next(tmp))

    first = [c for c in commands if c[1] == 'insert-text' and c[c.index('--para') + 1] == '26']
    assert [c[c.index('--text') + 1] for c in first] == ['첫 줄']   # 양식의 " ◦ " 뒤(offset 3)에 글만 — " ◦ ◦ 첫 줄"이 되지 않는다
    assert not [c for c in commands if c[1] == 'delete-text']
