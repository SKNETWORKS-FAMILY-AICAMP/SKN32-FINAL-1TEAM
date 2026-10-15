"""예비창업 사업비 집행계획의 1단계 · 2단계 표 나누기 (SB-339 후반 — 사업비 줄의 단계 phase는 SB-327).

원본 양식(별첨1)은 예비창업 사업비를 1단계 · 2단계 두 표로 나눈다. 사업비 줄의 단계로 docx는 표 둘, hwp는 표12(1단계) · 표14(2단계)를 채운다.
단계가 없는 줄은 1단계 쪽이다. 초기 · 개인사업자 · 법인은 단계가 없어 표 하나 그대로다.

실행:
    pytest tests/test_budget_phase_export.py -v
"""
import io
import json
from types import SimpleNamespace as NS

from docx import Document

from app.hwp_export import _body_text, _preliminary_steps
from app.models import Project
from app.plan_document_export import BudgetLineItem, split_budget_by_phase
from app.plan_texts import PlanTexts
from app.routers.projects import _build_plan_document_data


def _budget(**overrides):
    row = {'category': '외주용역비', 'execution_plan': '앱 개발 외주', 'total_amount': 15000000,
           'government_amount': 15000000, 'self_cash_amount': 0, 'self_in_kind_amount': 0}
    row.update(overrides)
    return row


def _create(client, applicant_type='preliminary', budget_items=None):
    payload = {'applicant_type': applicant_type, 'biz_type': '개인', 'ceo_name': '박테스트', 'description': 'AI 서비스',
               'team_members': [], 'pricing_items': [], 'budget_items': budget_items or []}
    if applicant_type == 'corp':
        payload.update(biz_type='법인', founded_at='2024-01-01')
    r = client.post('/projects', data={'payload': json.dumps(payload)})
    assert r.status_code == 201, r.text
    return r.json()['project_id']


PHASED = [
    _budget(phase='1단계', category='외주용역비', execution_plan='앱 개발 외주', total_amount=15000000, government_amount=15000000),
    _budget(phase='1단계', category='재료비', execution_plan='라벨링', total_amount=5000000, government_amount=5000000),
    _budget(phase='2단계', category='광고선전비', execution_plan='SNS 마케팅', total_amount=10000000, government_amount=10000000),
]


def _docx_tables(client, project_id):
    r = client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200, r.text
    doc = Document(io.BytesIO(r.content))
    paragraphs = [p.text for p in doc.paragraphs]
    return paragraphs, doc.tables


def _rows(table):
    return [[c.text for c in row.cells] for row in table.rows]


def test_split_puts_unphased_rows_in_the_first_table():
    rows = [BudgetLineItem('a', 'x', '1', '1', 단계='1단계'), BudgetLineItem('b', 'x', '1', '1'), BudgetLineItem('c', 'x', '1', '1', 단계='2단계')]
    first, second = split_budget_by_phase(rows)
    assert [r.비목 for r in first] == ['a', 'b'] and [r.비목 for r in second] == ['c']


def test_preliminary_docx_has_a_table_per_phase(authed_client):
    pid = _create(authed_client, budget_items=PHASED)

    paragraphs, tables = _docx_tables(authed_client, pid)

    assert '< 정부지원사업비 집행계획(1단계) >' in paragraphs and '< 정부지원사업비 집행계획(2단계) >' in paragraphs
    budget_tables = [t for t in tables if _rows(t)[0] == ['비목', '산출 근거', '정부지원사업비(원)']]
    assert len(budget_tables) == 2
    first, second = (_rows(t)[1:] for t in budget_tables)
    assert first == [['외주용역비', '앱 개발 외주', '15,000,000원'], ['재료비', '라벨링', '5,000,000원']]
    assert second == [['광고선전비', 'SNS 마케팅', '10,000,000원']]
    assert '1단계 합계: 20,000,000원' in paragraphs and '2단계 합계: 10,000,000원' in paragraphs


def test_preliminary_docx_keeps_an_empty_phase_table_as_placeholder(authed_client):
    pid = _create(authed_client, budget_items=PHASED[:2])

    _, tables = _docx_tables(authed_client, pid)

    budget_tables = [t for t in tables if _rows(t)[0] == ['비목', '산출 근거', '정부지원사업비(원)']]
    assert _rows(budget_tables[1])[1:] == [['○○', '○○', '○○']]


def test_non_preliminary_docx_keeps_the_single_budget_table(authed_client):
    pid = _create(authed_client, 'corp', [_budget()])

    paragraphs, tables = _docx_tables(authed_client, pid)

    assert '< 사업비 집행 계획 >' in paragraphs
    assert not [p for p in paragraphs if '집행계획(1단계)' in p]
    assert len([t for t in tables if '총사업비(ⓐ+ⓑ)' in _rows(t)[0]]) == 1


def _cells(steps, table):
    return {(s['row'], s['col']): s['text'] for s in steps if s['action'] == 'set_cell' and s['table'] == table}


def test_preliminary_hwp_fills_table_12_with_phase_1_and_table_14_with_phase_2(authed_client, db_session):
    pid = _create(authed_client, budget_items=PHASED)
    data, template = _build_plan_document_data(db_session, db_session.get(Project, pid), PlanTexts())
    assert template == 'preliminary'

    steps = _preliminary_steps(data)

    first, second = _cells(steps, 12), _cells(steps, 14)
    assert (first[(1, 0)], first[(1, 1)], first[(1, 2)]) == ('외주용역비', '앱 개발 외주', '15,000,000원')
    assert (first[(2, 0)], first[(2, 1)]) == ('재료비', '라벨링') and first[(6, 2)] == '20,000,000원'  # 둘째 항목도 자기 비목 칸에 쓴다(비목 칸을 먼저 나눈다)
    assert (12, 1, 0) in steps.split_cells and (14, 1, 0) in steps.split_cells
    assert (second[(1, 0)], second[(1, 2)]) == ('광고선전비', '10,000,000원') and second[(6, 2)] == '10,000,000원'
    assert second[(2, 1)] == '' and second[(2, 0)] == ''  # 2단계 표의 남는 칸은 공란


def test_hwp_body_labels_the_two_phase_narratives():
    data = NS(실현가능성_본문='실현', 실현가능성_일정_서술='', 사업비_서술='1단계 문장', 사업비_2단계_서술='2단계 문장')
    assert _body_text(data, '실현가능성_본문') == '실현 / [1단계] 1단계 문장 / [2단계] 2단계 문장'


def test_hwp_body_has_no_phase_label_without_a_second_narrative():
    data = NS(실현가능성_본문='실현', 실현가능성_일정_서술='', 사업비_서술='사업비 문장', 사업비_2단계_서술='')
    assert _body_text(data, '실현가능성_본문') == '실현 / 사업비 문장'
