# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUT = HERE / "데이터_전처리_결과서_Codex.docx"

FONT = "Noto Sans CJK KR"
MONO = "D2Coding"
NAVY = "1F4E78"
NAVY_DARK = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "6B7280"
BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_title(page_header_row):
    repeat_table_header(page_header_row)


def set_run_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(p, before=0, after=6, line=1.35, keep_with_next=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = True


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=MID_GRAY)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=7, line=1.45)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.8, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.8)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.8)
    return p


def add_bullets(doc, items):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=6, line=1.35)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    for idx, item in enumerate(items):
        run = p.add_run("• " + item)
        set_run_font(run, size=10.5)
        if idx < len(items) - 1:
            run.add_break()


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=4, line=1.35)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_table(doc, headers, rows, widths=None, aligns=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header = table.rows[0]
    set_repeat_title(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 110, 120, 110, 120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.1)
        r = p.add_run(str(text))
        set_run_font(r, size=9.3, bold=True, color=WHITE)

    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_margins(cell, 95, 120, 95, 120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[i]
            if ridx % 2:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.alignment = (aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT)
            set_paragraph_spacing(p, after=0, line=1.2)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_paragraph_spacing(p, before=12 if level == 1 else 8, after=6, line=1.1, keep_with_next=True)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_section_break(doc):
    doc.add_page_break()


def pct(n, d):
    return f"{n / d * 100:.1f}%"


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.8)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.4

    title_style = styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title_style.font.size = Pt(29)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        st = styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(46)
    r = p.add_run("S BRAIN 데이터 파이프라인")
    set_run_font(r, size=11, bold=True, color=NAVY_DARK)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, after=14, line=1.1)
    p.add_run("데이터 전처리 결과서")

    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=30, line=1.35)
    r = p.add_run("창업 지원 공고 수집 정규화 품질 검증 및 학습 데이터 구성")
    set_run_font(r, size=15, color=MID_GRAY)

    meta = [
        ("작성 범위", "data-collection"),
        ("기준 스냅샷", "2026년 9월 16일 09시 00분 KST"),
        ("문서 버전", "1.0"),
        ("작성 구분", "Codex 작성본"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in meta:
        row = table.add_row()
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(4.8)
        for cell in row.cells:
            set_cell_margins(cell, 75, 0, 75, 100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = row.cells[0].paragraphs[0]
        p2 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(label)
        r2 = p2.add_run(value)
        set_run_font(r1, size=10, bold=True, color=NAVY_DARK)
        set_run_font(r2, size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("본 결과서는 수집 원본을 공통 스키마와 검색 학습 데이터로 변환한 과정과 품질을 기록한다. 핵심 결과는 최신 일일 스냅샷 1,772건을 손실 없이 정규화했으며, 결측값과 해석 불가 값을 임의로 보정하지 않고 별도 상태로 보존했다는 점이다.")
    set_run_font(r, size=11.2)
    p.paragraph_format.line_spacing = 1.5

    doc.add_page_break()

    # Footer after cover also applies to all pages, intentionally minimal.
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("데이터 전처리 결과서  ")
    set_run_font(r, size=8.5, color=MID_GRAY)
    add_page_field(fp)

    add_heading(doc, "결과 요약", 1)
    add_body(doc, "최신 정규화 스냅샷에서는 K Startup 239건과 기업마당 1,533건, 총 1,772건을 공통 스키마로 변환했다. 필수 식별자와 제목이 누락된 거부 건은 없었고 동일 출처 중복도 없었다. 다만 접수기간을 구조화하지 못한 공고 54건과 출처 간 제목 일치 후보 3개 그룹은 후속 검토 대상으로 남겼다.")

    add_table(
        doc,
        ["구분", "결과", "해석"],
        [
            ["정규화 공고", "1,772건", "최신 일일 스냅샷 기준"],
            ["거부 행", "0건", "필수 ID와 제목을 충족"],
            ["동일 출처 중복", "0건", "notice_id 기준 중복 없음"],
            ["출처 간 검토 후보", "3개 그룹", "제목 완전 일치 후보이며 자동 병합하지 않음"],
            ["기간 해석 불가", "54건", "unknown으로 보존하고 issues에 기록"],
            ["첨부 원본", "1,721개", "약 650.8 MiB의 로컬 파일"],
            ["임베딩", "2,153건", "1024차원 BGE M3 벡터"],
            ["검색 관련도 라벨", "1,317쌍", "0 1 2의 3단 관련도"],
        ],
        widths=[Inches(1.55), Inches(1.25), Inches(4.0)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_body(doc, "판단 요약  전처리 결과는 검색과 자격요건 검토에 사용할 수 있는 수준이지만, 두 출처의 필드 제공 범위가 크게 다르다. 특히 기업마당은 업력과 지역 정형 필드를 제공하지 않으므로 결측을 제한 없음으로 해석해서는 안 된다. 또한 최신 일일 스냅샷과 누적 임베딩 파일의 모집단이 다르므로 모델 학습 시 공고 집합과 생성 시점을 고정해야 한다.", bold_lead="판단 요약")

    add_heading(doc, "문서 범위", 2)
    add_bullets(doc, [
        "대상 코드와 산출물은 data-collection 폴더로 한정한다.",
        "데이터 수집부터 정규화 저장 첨부 추출 임베딩 평가 라벨 구성까지를 전처리로 본다.",
        "머신러닝 모델의 성능 결과는 별도 학습결과서에서 다룬다.",
        "통계는 최신 로컬 산출물을 직접 집계했으며 과거 문서의 수치보다 최신 스냅샷을 우선한다.",
    ])

    add_section_break(doc)
    add_heading(doc, "1 데이터 출처", 1)
    add_body(doc, "공고 데이터는 K Startup과 기업마당의 공개 API에서 수집한다. 두 출처는 동일한 창업 및 중소기업 지원 공고를 다루지만 필드 이름과 의미가 다르므로 원본을 바로 합치지 않고 출처별 매핑 규칙을 거쳐 공통 스키마로 변환한다.")

    add_table(
        doc,
        ["출처", "입력 건수", "비율", "주요 강점", "주요 제약"],
        [
            ["K Startup", "239", "13.5%", "지원대상 업력 지역 접수기간", "첨부 URL을 API에서 제공하지 않음"],
            ["기업마당", "1,533", "86.5%", "소관기관 수행기관 세부분류 첨부", "업력 지역 정형 필드가 없음"],
            ["합계", "1,772", "100.0%", "두 출처를 공통 스키마로 통합", "출처별 결측 의미를 보존해야 함"],
        ],
        widths=[Inches(1.05), Inches(0.8), Inches(0.7), Inches(2.0), Inches(2.25)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.9,
    )

    add_heading(doc, "1 1 원본 보존 원칙", 2)
    add_body(doc, "정규화 과정에서 사용하지 않는 필드도 삭제하지 않는다. 각 공고의 raw 필드에 API 응답 전체를 저장하여 정규화 오류를 역추적하거나 이후 새로운 특성을 만들 수 있게 한다. 원본 API 파일과 정규화 파일은 날짜별 스냅샷으로 남기므로 특정 실행 시점의 결과를 재현할 수 있다.")

    add_heading(doc, "1 2 통합 식별자", 2)
    add_body(doc, "출처에서 제공하는 ID만 사용하면 서로 다른 시스템의 값이 충돌할 수 있다. 따라서 notice_id를 source와 source_id의 조합으로 생성한다. 예를 들어 K Startup 공고 179262는 kstartup:179262로 저장한다. 이 키는 MySQL 저장 임베딩 첨부 관계 평가 라벨을 연결하는 기준이다.")

    add_heading(doc, "2 전처리 파이프라인", 1)
    add_body(doc, "전처리는 수집 결과를 바로 덮어쓰지 않고 검증 가능한 중간 산출물을 단계별로 남기는 방식으로 설계했다. 한 단계가 실패해도 이전의 정상 스냅샷과 이미 처리한 첨부를 재사용한다.")

    add_table(
        doc,
        ["단계", "처리", "주요 산출물", "품질 통제"],
        [
            ["1", "API 수집", "원본 JSON", "재시도와 출처별 성공 상태 기록"],
            ["2", "공통 스키마 정규화", "normalized JSON", "필수값 검증 및 issues 보존"],
            ["3", "MySQL 저장", "notices와 import_runs", "트랜잭션 및 과거 스냅샷 보호"],
            ["4", "첨부 처리", "원본 파일과 추출 본문", "해시 중복 제거 및 형식별 추출"],
            ["5", "정형 조건 추출", "notice_conditions", "근거 문장과 검산 결과 동시 저장"],
            ["6", "문서 임베딩", "embeddings_v1.npz", "모델 설정 지문과 입력 해시 기록"],
            ["7", "검색 평가 데이터", "pool과 qrels", "질의 단위 라벨과 판정 출처 보존"],
        ],
        widths=[Inches(0.55), Inches(1.35), Inches(2.05), Inches(2.85)],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
    )

    add_heading(doc, "2 1 실패 격리와 재처리", 2)
    add_body(doc, "수집과 정규화가 실패하면 운영 DB를 교체하지 않는다. 첨부 다운로드 임베딩 생성 벡터 업로드는 변경된 항목만 다시 처리한다. 파일 해시와 입력 해시를 기준으로 이미 완료된 작업을 건너뛰므로 배치가 중단돼도 전체를 처음부터 반복하지 않는다.")

    add_heading(doc, "3 정규화 규칙", 1)

    add_heading(doc, "3 1 텍스트 정제", 2)
    add_body(doc, "HTML 본문은 태그를 제거하되 문단 목록 표의 행 경계는 줄바꿈으로 유지한다. script와 style 내용은 제외하고 XML 엔티티를 복원한다. 연속 공백은 한 칸으로 정리하며 빈 줄은 제거한다. 정제 결과가 비어 있으면 빈 문자열이 아니라 NULL로 저장한다.")

    add_heading(doc, "3 2 날짜와 접수기간", 2)
    add_body(doc, "K Startup의 8자리 날짜와 기업마당의 점, 하이픈, 슬래시 혼합 날짜를 ISO 8601 형식으로 통일한다. 날짜가 없는 표현은 버리지 않고 기간의 의미를 fixed, budget_exhaustion, rolling, until_filled, unknown 다섯 종류로 구분한다. 읽지 못한 기간은 issues에 unparsed_period로 남긴다.")

    add_table(
        doc,
        ["기간 유형", "건수", "비율", "의미"],
        [
            ["fixed", "852", "48.1%", "시작일과 종료일을 모두 구조화"],
            ["budget_exhaustion", "632", "35.7%", "예산 소진 시까지"],
            ["rolling", "143", "8.1%", "상시 또는 수시 접수"],
            ["until_filled", "91", "5.1%", "선착순 또는 모집 완료 시까지"],
            ["unknown", "54", "3.0%", "기간 문구를 안전하게 해석하지 못함"],
        ],
        widths=[Inches(1.45), Inches(0.75), Inches(0.75), Inches(3.85)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "3 3 URL과 첨부 주소", 2)
    add_body(doc, "상대 주소는 출처 도메인을 기준으로 절대 주소로 변환하고 amp 엔티티를 복원한다. http와 https 이외의 스킴 사용자 정보가 포함된 주소 공백이 섞인 주소 잘못된 포트는 invalid_url로 기록한 뒤 사용하지 않는다.")

    add_heading(doc, "3 4 결측값과 판단 불가", 2)
    add_body(doc, "결측값을 제한 없음이나 조건 충족으로 바꾸지 않는다. 데이터가 없거나 해석할 수 없는 경우 NULL과 unknown을 유지한다. 이는 기업마당 공고의 업력과 지역 필드가 대부분 비어 있는 상황에서 잘못된 자격 판정을 방지하기 위한 핵심 규칙이다.")

    add_heading(doc, "3 5 중복과 이상치", 2)
    add_body(doc, "동일 출처에서 notice_id가 중복되면 마지막 정상 행을 사용하고 교체 이력을 별도로 남긴다. 서로 다른 출처의 제목이 완전히 같은 경우에는 삭제하거나 합치지 않고 검토 후보로만 표시한다. 최신 스냅샷에는 동일 출처 중복이 없고 출처 간 검토 후보가 3개 그룹 존재한다.")

    add_heading(doc, "4 정규화 결과", 1)
    add_body(doc, "최신 스냅샷 notices_20260916T000016111196Z.json을 기준으로 주요 필드의 확보율을 집계했다. 확보율은 데이터의 존재 여부를 뜻하며 값의 의미 정확도를 자동으로 보증하지 않는다.")

    add_table(
        doc,
        ["필드", "확보 건수", "확보율", "해석"],
        [
            ["body", "1,772", "100.0%", "API 사업개요 정제 본문"],
            ["target_text", "239", "13.5%", "K Startup만 지원대상 원문 제공"],
            ["age_condition_raw", "239", "13.5%", "K Startup 정형 업력 조건"],
            ["region", "239", "13.5%", "기업마당에는 대응 필드 없음"],
            ["apply_start와 apply_end", "852", "48.1%", "고정 날짜형 공고만 두 날짜 확보"],
            ["subcategory", "1,533", "86.5%", "기업마당만 세부분류 제공"],
            ["issues 보유", "54", "3.0%", "모두 접수기간 해석 불가"],
        ],
        widths=[Inches(1.65), Inches(0.9), Inches(0.8), Inches(3.55)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9,
    )

    add_body(doc, "출처 편향  전체 공고의 86.5%를 차지하는 기업마당은 업력과 지역 정형 필드를 제공하지 않는다. 따라서 결측률이 높다는 이유로 해당 행을 제거하면 학습 데이터가 K Startup 중심으로 편향된다. 검색 관련도 모델은 두 출처를 유지하고 자격요건 모델은 판정 가능 여부를 별도 목표로 다뤄야 한다.", bold_lead="출처 편향")

    add_heading(doc, "5 첨부 데이터 전처리", 1)
    add_body(doc, "기업마당이 제공하는 공고문과 신청서식 링크는 정규화 시 notice_attachments 구조로 분리한다. 최신 정규화 스냅샷에는 첨부 링크 2,614건이 연결되어 있다. 로컬 첨부 저장소에는 내용 해시를 기준으로 1,721개 파일이 보존되어 있으며 전체 크기는 약 650.8 MiB다.")

    add_table(
        doc,
        ["형식", "파일 수", "비율", "처리 방식"],
        [
            ["PDF", "936", "54.4%", "pdfplumber 기반 텍스트 추출"],
            ["HWP", "527", "30.6%", "OLE 구조 분석과 HWP 5 파서"],
            ["HWPX", "199", "11.6%", "ZIP 내부 XML 파싱"],
            ["IMG", "58", "3.4%", "이미지 또는 형식 미식별 파일로 OCR 후보"],
            ["DOCX", "1", "0.1%", "ZIP 내부 Word XML 파싱"],
            ["합계", "1,721", "100.0%", "내용 해시 기반 중복 저장 방지"],
        ],
        widths=[Inches(1.0), Inches(0.8), Inches(0.75), Inches(4.35)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "5 1 첨부 품질 상태", 2)
    add_body(doc, "첨부 처리 결과는 성공과 실패를 한 값으로 합치지 않는다. ok, image_only, unsupported, download_fail, parse_error, empty_text 등 최근 시도 상태를 구분하고 마지막 성공 본문은 이후 재시도가 실패해도 보존한다. 파일 바이트의 SHA 256과 추출기 버전을 기록하여 동일 문서를 중복 처리하지 않는다.")

    add_heading(doc, "5 2 모델 입력 시 주의", 2)
    add_bullets(doc, [
        "첨부 파일 수와 텍스트 추출 성공 건수는 같은 지표가 아니다.",
        "image_only 파일은 OCR 전에는 학습 본문으로 직접 사용할 수 없다.",
        "한 공고에 여러 첨부가 있을 수 있으므로 파일 단위 분할은 데이터 누수를 만든다.",
        "학습과 시험 분할은 notice_id 또는 공고 그룹 기준으로 수행해야 한다.",
    ])

    add_heading(doc, "6 임베딩 전처리", 1)
    add_body(doc, "검색용 문서 텍스트는 제목 사업개요 지원대상 지원대상 분류 대분류 세부분류의 여섯 필드를 결합해 만든다. BAAI bge m3 모델로 최대 512토큰까지 인코딩하며 1024차원 float32 벡터를 생성하고 코사인 유사도 계산을 위해 정규화한다.")

    add_table(
        doc,
        ["항목", "값"],
        [
            ["모델", "BAAI bge m3"],
            ["모델 리비전", "5617a9f61b028005a4858fdac845db406aefb181"],
            ["입력 버전", "v1 title body target_text target_category category subcategory"],
            ["벡터 수", "2,153건"],
            ["벡터 형상", "2153 x 1024"],
            ["자료형", "float32 little endian"],
            ["정규화", "L2 normalized true"],
            ["행렬 원시 크기", "8,818,688 bytes 약 8.41 MiB"],
        ],
        widths=[Inches(1.7), Inches(5.2)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.1,
    )

    add_body(doc, "모집단 불일치  최신 일일 정규화 스냅샷은 1,772건이지만 임베딩 파일에는 2,153건이 있다. 두 파일을 행 순서로 결합해서는 안 되며 notice_id 교집합과 입력 해시로 연결해야 한다. 모델 학습 전에는 사용할 공고 스냅샷과 임베딩 지문을 고정하고 누락 및 잔존 벡터를 검사해야 한다.", bold_lead="모집단 불일치")

    add_heading(doc, "7 머신러닝 학습 데이터 구성", 1)
    add_body(doc, "현재 검색 평가 데이터는 사용자 입력 형태의 질의와 공고 후보 쌍으로 구성되어 있다. 정상 질의 52개와 무관 질의 8개를 합쳐 60개 질의를 사용하며 Dense 상위 20건과 BM25 상위 10건의 합집합으로 후보 풀을 생성한다.")

    add_table(
        doc,
        ["항목", "건수", "비고"],
        [
            ["질의", "60", "정상 52개 무관 8개"],
            ["후보 쌍", "1,657", "Dense와 BM25 후보 합집합"],
            ["최종 관련도 라벨", "1,317", "미판정 후보는 0으로 바꾸지 않음"],
            ["관련도 0", "441", "33.5% 무관"],
            ["관련도 1", "240", "18.2% 부분 관련"],
            ["관련도 2", "636", "48.3% 매우 관련"],
            ["사람 판정", "143", "10.9%"],
            ["LLM 기존 판정", "1,174", "89.1% 약한 라벨"],
        ],
        widths=[Inches(1.8), Inches(0.9), Inches(4.2)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "7 1 라벨 정의", 2)
    add_table(
        doc,
        ["라벨", "정의", "처리"],
        [
            ["2", "사용자의 사업 내용과 지원 목적이 직접 일치", "상위 노출 목표"],
            ["1", "분야는 관련되지만 핵심 목적이나 대상이 일부 다름", "부분 관련"],
            ["0", "지원 내용과 분야가 무관", "비관련"],
            ["미판정", "사람 또는 합의된 절차로 확인하지 않음", "학습 정답에서 제외"],
        ],
        widths=[Inches(0.7), Inches(4.25), Inches(1.95)],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "7 2 권장 학습 분할", 2)
    add_body(doc, "후보 쌍을 무작위 행 단위로 나누면 같은 질의의 다른 공고가 학습과 시험에 동시에 들어가 성능이 부풀려질 수 있다. 따라서 qid 또는 scenario_group을 기준으로 학습 검증 시험을 분리한다. 권장 비율은 질의 기준 70대 15대 15이며 정상 질의와 무관 질의가 각 분할에 포함되도록 고정 seed를 사용한다.")

    add_heading(doc, "7 3 학습 특성 후보", 2)
    add_bullets(doc, [
        "Dense 코사인 유사도와 Dense 순위",
        "BM25 점수와 BM25 순위",
        "각 검색 방식의 후보 포함 여부와 reciprocal rank",
        "질의와 제목의 토큰 또는 문자 n gram 중복률",
        "공고 출처와 기간 유형 등 누수 위험이 낮은 정형 특성",
    ])

    add_body(doc, "LLM 판정 1,174건은 사람이 확정한 정답과 동일하게 취급하지 않는다. 약한 라벨로 학습할 수는 있지만 최종 성능은 사람이 검수한 질의 그룹에서 별도로 보고해야 한다. 사람 판정이 부족하면 모델의 운영 채택보다 실험 결과와 한계를 기록하는 것이 우선이다.")

    add_heading(doc, "8 품질 검증 결과", 1)
    add_body(doc, "정규화 첨부 저장 자격 판정 등 data-collection의 단위 테스트를 2026년 9월 16일에 실행했다. 총 81개 테스트 중 외부 환경이 필요한 13개가 건너뛰어졌고 실행된 68개는 모두 통과했다.")

    add_table(
        doc,
        ["검증 항목", "결과", "판정"],
        [
            ["단위 테스트 발견", "81개", "정상"],
            ["실행 및 통과", "68개", "실패 0개"],
            ["건너뜀", "13개", "외부 환경 또는 선택 의존성 조건"],
            ["최신 정규화 거부", "0건", "필수값 검증 통과"],
            ["동일 출처 중복", "0건", "중복 정책 작동 결과"],
            ["해석 경고", "54건", "삭제하지 않고 issues로 보존"],
        ],
        widths=[Inches(2.0), Inches(1.3), Inches(3.6)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "8 1 재현성 정보", 2)
    add_bullets(doc, [
        "원본 파일과 정규화 파일의 생성 시각 및 SHA 256을 import_runs와 산출물 메타데이터에 기록한다.",
        "임베딩은 모델 리비전 입력 버전 최대 토큰 차원 자료형 정규화 여부를 fingerprint로 관리한다.",
        "평가 라벨은 label_version과 판정 주체를 보존한다.",
        "실험 결과는 질의 집합 공고 스냅샷 코드 버전 random seed를 함께 기록해야 한다.",
    ])

    add_heading(doc, "9 한계와 개선 권고", 1)
    add_table(
        doc,
        ["우선순위", "한계", "영향", "권고 조치"],
        [
            ["높음", "사람 판정 라벨이 143쌍", "LLM 라벨 편향을 독립적으로 검증하기 어려움", "시험 질의 그룹의 후보 전체를 우선 사람 검수"],
            ["높음", "스냅샷 1,772건과 벡터 2,153건 불일치", "잘못된 조인과 평가 누수 가능", "notice_id 교집합과 스냅샷 해시로 고정"],
            ["높음", "기업마당의 업력 지역 결측", "자격 판정 학습 시 출처 편향", "결측을 별도 상태로 유지하고 첨부 근거 활용"],
            ["중간", "기간 unknown 54건", "마감 필터의 확인 필요 증가", "기간 패턴 추가 후 원문 기반 회귀 시험"],
            ["중간", "이미지형 첨부 58개", "본문과 조건 추출 누락", "OCR 파일럿 후 실제 복구 공고 수 측정"],
            ["중간", "출처 간 제목 일치 3개 그룹", "추천 결과 중복 가능", "기관 기간 지역 차수를 포함한 수동 확인"],
        ],
        widths=[Inches(0.7), Inches(1.65), Inches(2.25), Inches(2.3)],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.5,
    )

    add_heading(doc, "10 결론", 1)
    add_body(doc, "data-collection 파이프라인은 두 공고 API의 구조 차이를 공통 스키마로 통합하면서 원본과 결측 의미를 보존한다. 최신 스냅샷 1,772건은 거부와 동일 출처 중복 없이 정규화되었고 첨부 1,721개 임베딩 2,153건 검색 관련도 라벨 1,317쌍이 후속 분석 자산으로 준비되어 있다.")
    add_body(doc, "머신러닝 학습 전에는 최신 공고 스냅샷 임베딩 평가 후보의 모집단을 notice_id와 해시로 고정해야 한다. 검색 재정렬 모델은 현재 라벨로 실험할 수 있지만 LLM 기반 약한 라벨 비중이 89.1%이므로 사람 검수 시험셋을 확대한 뒤 성능을 최종 판단하는 것이 타당하다.")

    add_heading(doc, "부록 A 주요 산출물", 1)
    add_table(
        doc,
        ["산출물", "역할"],
        [
            ["data normalized notices timestamp json", "출처 통합 공고 스냅샷과 정규화 요약"],
            ["data attachments", "내용 해시 기반 첨부 원본 저장소"],
            ["data embeddings_v1.npz", "공고 ID 입력 해시 벡터 메타데이터"],
            ["eval queries.jsonl", "사용자 입력 형태의 평가 질의"],
            ["eval pool.jsonl", "Dense와 BM25로 생성한 판정 후보"],
            ["eval qrels.jsonl", "질의 공고 관련도와 판정 출처"],
            ["collect_log.jsonl", "일일 파이프라인 단계별 실행 기록"],
        ],
        widths=[Inches(2.65), Inches(4.25)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "부록 B 근거 코드", 1)
    add_table(
        doc,
        ["파일", "확인 내용"],
        [
            ["normalize.py", "텍스트 날짜 URL 결측 중복 정규화 규칙"],
            ["store_mysql.py", "트랜잭션 기반 DB 저장과 과거 스냅샷 보호"],
            ["attachment_pipeline.py", "첨부 다운로드 형식 판별 본문 추출"],
            ["doctext.py와 hwp5.py", "PDF HWP HWPX DOCX 텍스트 처리"],
            ["embed.py", "임베딩 입력과 모델 fingerprint 생성"],
            ["eval build_pool.py", "Dense와 BM25 후보 풀 구성"],
            ["eval merge_qrels.py", "사람 및 LLM 판정 통합"],
        ],
        widths=[Inches(2.25), Inches(4.65)],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    # Core properties and final paragraph hygiene.
    props = doc.core_properties
    props.title = "데이터 전처리 결과서"
    props.subject = "창업 지원 공고 수집 정규화 품질 검증 및 학습 데이터 구성"
    props.author = "S BRAIN 프로젝트"
    props.keywords = "데이터 전처리 공고 정규화 첨부 임베딩 머신러닝"

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
