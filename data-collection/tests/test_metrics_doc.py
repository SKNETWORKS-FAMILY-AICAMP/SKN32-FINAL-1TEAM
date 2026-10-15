# -*- coding: utf-8 -*-
"""지표 보고서(eval.metrics_doc) — 가짜 metrics.json 으로 HTML·Word 를 만들어 본다(2026-10-14).

python-docx 가 없으면 Word 부분을, Edge 가 없으면 그래프 그림 부분을 건너뛴다(환경 때문에 전체 시험이 깨지지 않게).
"""
import copy
import hashlib
import html as html_lib
import io
import json
import os
import re
import shutil
import struct
import sys
import tempfile
import unittest
import zlib
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from eval import metrics_doc as md  # noqa: E402

try:
    import docx  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
HAS_EDGE = md.find_edge() is not None

NULL_REASON = '실제 서버에 닿지 못했다'


def metric(area, name, value, unit, **kw):
    m = {'area': area, 'name': name, 'value': value, 'unit': unit, 'target': '', 'detail': '',
         'plain': name + ' 설명', 'source': 'tests', 'kind': 'live', 'measured_at': '2026-10-15'}
    m.update(kw)
    return m


def sample():
    return {
        'run_at': '2026-10-15T01:00:00Z',
        'conditions': ['공용 DB 조회만', 'OpenAI 호출 0'],
        'metrics': [
            metric('1', '누적 공고 수', 2712, 'count'),
            metric('1', '배치 성공(최근 16회)', 16, 'count', target='16', detail='16회 중 16회 통과'),
            metric('4', '신청 불가@10', 0.0, 'ratio', target='0', detail='처음(벡터만) 0.300 → 지금 0.000',
                   compare={'before': 0.3, 'after': 0.0, 'label_before': '처음(벡터만)', 'label_after': '지금 서비스',
                            'reference': {'label': '9/28 평가', 'before': 0.243, 'after': 0.0}}),
            metric('4', 'nDCG@10', 0.6034, 'score', kind='record', measured_at='2026-09-28'),
            metric('8', '응답 시간(이 PC)', 182.4, 'ms', target='<=500',
                   source='http://10.1.2.3:8000/api/match 로 58회'),
            metric('8', '응답 시간(실제 서버)', None, 'ms', detail=NULL_REASON + '(시간 초과)'),
            metric('9', '하루 LLM 비용', 0.13, 'usd'),
            metric('11', '조율 연결 상태', '연결됨', 'text'),
            metric('ML', '리랭커 nDCG@10', 0.512, 'score', kind='record', measured_at='2026-09-27'),
        ],
        'headline': [
            {'title': '추천 응답 시간', 'values': [{'label': '이 PC 기준', 'metric': '응답 시간(이 PC)'},
                                             {'label': '실제 서버 기준', 'metric': '응답 시간(실제 서버)'}],
             'chart': None, 'note': ''},
            {'title': '신청 불가 추천', 'values': [{'label': '지금', 'metric': '신청 불가@10'}],
             'chart': '신청 불가@10', 'note': '처음 방식은 30%였다.'},
        ],
        'missing': [{'name': '응답 시간(실제 서버)', 'reason': '연결 실패', 'expected': True}],
    }


AREA_TABLES = 6        # 1 · 4 · 8 · 9 · 11 · ML


def write_json(folder, data):
    path = os.path.join(folder, 'metrics.json')
    with io.open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    return path


def digest(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


class FormatTests(unittest.TestCase):
    def test_units(self):
        cases = [('ratio', 0.243, '24.3%'), ('ratio', 0.0, '0.0%'), ('score', 0.6, '0.600'),
                 ('count', 2712, '2,712'), ('count', 12.0, '12'), ('count', 1.5, '1.50'),
                 ('ms', 182.4, '182ms'), ('ms', 4.25, '4.2ms'), ('sec', 1.234, '1.23초'), ('min', 41.62, '41.6분'),
                 ('mb', 812.3, '812MB'), ('usd', 0.13, '$0.13'), ('usd', 0.0042, '$0.0042'),
                 # 1달러 미만은 셋째 자리까지 — 표의 시간당 요금 × 730 이 그래프의 한 달 요금과 맞게
                 ('usd', 0.052, '$0.052'), ('usd', 0.647, '$0.647'), ('usd', 0.05, '$0.05'), ('usd', 37.96, '$37.96'),
                 ('usd', 472.31, '$472.31'), ('usd', 0.0, '$0.00'),
                 ('text', '연결됨', '연결됨')]
        for unit, value, want in cases:
            with self.subTest(unit=unit, value=value):
                self.assertEqual(md.fmt_value({'unit': unit, 'value': value}), want)

    def test_null_value(self):
        self.assertEqual(md.fmt_value({'unit': 'ms', 'value': None}), '측정 못 함')

    def test_pass_count_uses_denominator(self):
        self.assertEqual(md.fmt_value({'unit': 'count', 'value': 16, 'target': '16', 'detail': '16회 중 16회 통과'}),
                         '16/16')
        # 통과·성공 같은 말이 없으면 분모로 쓰지 않는다
        self.assertEqual(md.fmt_value({'unit': 'count', 'value': 16, 'target': '16', 'name': '공고 수'}), '16')

    def test_target(self):
        self.assertEqual(md.fmt_target({'unit': 'ratio', 'target': '0'}), '0%')
        self.assertEqual(md.fmt_target({'unit': 'ratio', 'target': '≥0.95'}), '≥95%')
        self.assertEqual(md.fmt_target({'unit': 'ms', 'target': '<=500'}), '≤500ms')
        self.assertEqual(md.fmt_target({'unit': 'score', 'target': '≥0.70'}), '≥0.70')
        self.assertEqual(md.fmt_target({'unit': 'count', 'target': ''}), '—')

    def test_run_at_in_kst(self):
        self.assertEqual(md.fmt_run_at('2026-10-15T01:00:00Z'), '2026-10-15 10:00 KST')
        self.assertEqual(md.fmt_run_at('2026-10-15T10:00:00+09:00'), '2026-10-15 10:00 KST')

    def test_scrub_hides_addresses_and_keys(self):
        s = md.scrub('http://10.1.2.3:8000/x 와 43.201.90.238:8000, sk-abcdefghijklmnop1234, password=hunter2, '
                     'sbrain.duckdns.org')
        for leak in ('10.1.2.3', '43.201', 'sk-abc', 'hunter2', 'duckdns'):
            self.assertNotIn(leak, s)
        # 모듈 이름·파일 이름은 그대로 둔다
        self.assertEqual(md.scrub('search.app.match · eval/common.py · metrics.json'),
                         'search.app.match · eval/common.py · metrics.json')


class RenderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.json_path = write_json(self.tmp, sample())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def render(self, edge=False):
        if edge:
            return md.render(self.tmp)
        with mock.patch.object(md, 'find_edge', return_value=None):
            return md.render(self.tmp)

    def read_html(self, path):
        with io.open(path, encoding='utf-8') as f:
            return f.read()

    def test_html_structure_and_order(self):
        html_path, _ = self.render()
        self.assertTrue(os.path.isfile(html_path))
        page = self.read_html(html_path)
        order = ['공고 파트 작업 지표', '2026-10-15 10:00 KST', '공용 DB 조회만', '한눈에 보기', '1. 수집·운영',
                 '4. 추천 정확도', '8. 서버 성능', '9. 비용', '11. 약점과 한계', '서비스 밖 실험(ML)',
                 '서비스에 연결되어 있지 않다', '다시 뽑는 방법', '이번에 재지 못한 것']
        pos = [page.find(x) for x in order]
        self.assertNotIn(-1, pos, dict(zip(order, pos)))
        self.assertEqual(pos, sorted(pos))
        self.assertNotIn('2. 전처리 품질', page)       # 지표 없는 영역은 건너뛴다
        self.assertIn('eval.metrics_report --all', page)
        self.assertIn('<svg', page)
        self.assertIn('9/28 평가', page)
        self.assertIn('16/16', page)

    def test_html_is_self_contained_and_hides_addresses(self):
        html_path, _ = self.render()
        page = self.read_html(html_path)
        for bad in ('<script', '<link', 'http://', 'https://', '10.1.2.3', '@import', 'url('):
            self.assertNotIn(bad, page)

    def test_null_metric_shows_reason(self):
        html_path, _ = self.render()
        page = self.read_html(html_path)
        i = page.find('응답 시간(실제 서버)</td>')
        self.assertGreater(i, 0)
        row = page[i:page.find('</tr>', i)]
        self.assertIn('측정 못 함', row)
        self.assertIn(NULL_REASON, row)

    def test_null_reason_falls_back_to_missing(self):
        data = sample()
        data['metrics'][5]['detail'] = ''
        write_json(self.tmp, data)
        html_path, _ = self.render()
        page = self.read_html(html_path)
        i = page.find('응답 시간(실제 서버)</td>')
        self.assertIn('연결 실패', page[i:page.find('</tr>', i)])

    def test_metrics_json_untouched(self):
        before = digest(self.json_path)
        self.render()
        self.assertEqual(digest(self.json_path), before)
        names = set(os.listdir(self.tmp)) - {'metrics.json'}
        self.assertTrue(names <= {md.HTML_NAME, md.DOCX_NAME}, names)

    def test_without_python_docx_only_html(self):
        with mock.patch.object(md, '_docx', return_value=False):
            html_path, docx_path = self.render()
        self.assertTrue(os.path.isfile(html_path))
        self.assertIsNone(docx_path)

    def test_cli(self):
        with mock.patch.object(md, 'find_edge', return_value=None), mock.patch('sys.stdout', new=io.StringIO()):
            self.assertEqual(md.main([self.tmp]), 0)
        empty = tempfile.mkdtemp()
        try:
            with mock.patch('sys.stderr', new=io.StringIO()):
                self.assertEqual(md.main([empty]), 2)
        finally:
            shutil.rmtree(empty, ignore_errors=True)

    @unittest.skipUnless(HAS_DOCX, 'python-docx 없음')
    def test_docx_structure_without_edge(self):
        from docx import Document
        from docx.oxml.ns import qn
        _, docx_path = self.render(edge=False)
        self.assertTrue(os.path.isfile(docx_path))
        doc = Document(docx_path)
        heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name in ('Title', 'Heading 1', 'Heading 2')]
        texts = [t for _, t in heads]
        self.assertEqual(heads[0], ('Title', '공고 파트 작업 지표'))
        for want in ('한눈에 보기', '추천 응답 시간', '신청 불가 추천', '1. 수집·운영', '4. 추천 정확도', '8. 서버 성능',
                     '9. 비용', '11. 약점과 한계', '서비스 밖 실험(ML)', '다시 뽑는 방법', '이번에 재지 못한 것'):
            self.assertIn(want, texts)
        self.assertNotIn('2. 전처리 품질', texts)
        self.assertLess(texts.index('11. 약점과 한계'), texts.index('서비스 밖 실험(ML)'))
        self.assertLess(texts.index('서비스 밖 실험(ML)'), texts.index('다시 뽑는 방법'))
        # Edge 가 없으면 그래프는 숫자 표 1개로 대신한다
        self.assertEqual(len(doc.tables), AREA_TABLES + 1)
        self.assertEqual(len(doc.inline_shapes), 0)
        area = doc.tables[1]
        self.assertEqual([c.text for c in area.rows[0].cells], md.COLUMNS)
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        null_cells = [c for c in cells if '측정 못 함' in c]
        self.assertTrue(any(NULL_REASON in c for c in null_cells), null_cells)
        self.assertIn('16/16', cells)
        body = '\n'.join(p.text for p in doc.paragraphs)
        self.assertNotIn('10.1.2.3', body + '\n'.join(cells))
        fonts = doc.styles['Normal'].element.rPr.find(qn('w:rFonts'))
        self.assertEqual(fonts.get(qn('w:eastAsia')), '맑은 고딕')
        self.assertEqual(fonts.get(qn('w:ascii')), '맑은 고딕')

    @unittest.skipUnless(HAS_DOCX and HAS_EDGE, 'python-docx 또는 Edge 없음')
    def test_docx_chart_picture_with_edge(self):
        from docx import Document
        _, docx_path = self.render(edge=True)
        doc = Document(docx_path)
        self.assertGreaterEqual(len(doc.inline_shapes), 1)
        self.assertEqual(len(doc.tables), AREA_TABLES)
        self.assertEqual(set(os.listdir(self.tmp)), {'metrics.json', md.HTML_NAME, md.DOCX_NAME})

    def test_does_not_mutate_input_dict(self):
        data = sample()
        frozen = copy.deepcopy(data)
        md.build_html(md.prepare(data))
        self.assertEqual(data, frozen)


# ── 영역 그래프(charts) ─────────────────────────────────

DAYS = ['2026-09-%02d' % d for d in range(16, 31)] + ['2026-10-%02d' % d for d in range(1, 16)]


def chart(cid, area, form, cats, series, **kw):
    c = {'id': cid, 'area': area, 'title': cid + ' 제목', 'subtitle': cid + ' 부제', 'form': form, 'unit': 'ratio',
         'units': None, 'categories': cats, 'series': series, 'highlight': None, 'max': None, 'reference': [],
         'errors': None, 'note': cid + ' 쉬운 설명', 'source': 'tests', 'kind': 'live', 'measured_at': '2026-10-15'}
    c.update(kw)
    return c


def chart_sample():
    """다섯 모양 모두 + 빈값·구간·기준선·강조·max. 보고서 순서와 일부러 다르게 넣는다(ML 이 맨 앞)."""
    data = sample()
    data['charts'] = [
        chart('c06_pairs', 'ML', 'pairs', ['상위 3건 적합', '질의당 시간', 'GPU 최대 메모리'],
              [{'name': '학습 전', 'values': [0.62, 950, None]}, {'name': '학습 후', 'values': [0.686, 1080, 4710]}],
              unit=None, units=['ratio', 'ms', 'mb'], errors=[[0.55, 0.70], None, None]),
        chart('c01_line', '1', 'line', DAYS,
              [{'name': '누적 공고', 'values': [None if i == 5 else 2700 + 10 * i for i in range(30)]}], unit='count',
              reference=[{'value': 2900, 'label': '목표'}]),
        chart('c02_column', '1', 'column', DAYS,
              [{'name': '신규 공고', 'values': [None if i == 7 else (i * 37) % 120 for i in range(30)]}], unit='count',
              reference=[{'value': 62, 'label': '30일 평균'}]),
        chart('c03_stacked', '3', 'stacked', ['자격요건', '신청자 유형'],
              [{'name': '판정함', 'values': [0.86, 0.9]}, {'name': '판정 불가', 'values': [0.14, None]}], max=1),
        chart('c04_bar', '4', 'bar', ['처음(벡터만)', '9/28 평가', '지금 서비스'],
              [{'name': 'P@3', 'values': [0.494, 0.621, 0.615]}], max=1, highlight=[2],
              errors=[[0.43, 0.56], None, [0.54, 0.69]], reference=[{'value': 0.7, 'label': '목표 70%'}],
              source='http://10.1.2.3:8000/api/match 로 잼'),
        chart('c05_bar3', '8', 'bar', ['이 PC', '실제 서버'],
              [{'name': '중앙값', 'values': [249.7, 412.6]}, {'name': '95번째 백분위', 'values': [380.2, None]},
               {'name': '최대', 'values': [900, 1200]}], unit='ms', reference=[{'value': 500, 'label': '목표'}]),
    ]
    return data


REPORT_ORDER = ['c01_line', 'c02_column', 'c03_stacked', 'c04_bar', 'c05_bar3', 'c06_pairs']
N_CHARTS = len(REPORT_ORDER)
PNG_NAMES = ['00a_headline.png'] + ['%02d_%s.png' % (i + 1, cid) for i, cid in enumerate(REPORT_ORDER)]
_TEXT = re.compile(r'<text ([^>]*)>([^<]*)</text>')
_ATTR = re.compile(r'([\w-]+)="([^"]*)"')


def png_size(path):
    with open(path, 'rb') as f:
        head = f.read(24)
    return struct.unpack('>II', head[16:24])


def tiny_png(w, h):
    """흰색 PNG(진짜 형식). 가짜 Edge 가 쓴다."""
    def chunk(kind, body):
        return struct.pack('>I', len(body)) + kind + body + struct.pack('>I', zlib.crc32(kind + body) & 0xffffffff)
    raw = b''.join(b'\x00' + b'\xff' * (w * 3) for _ in range(h))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def fake_svg_to_png(svg, width, height, png_path, edge, work, timeout=40):
    with open(png_path, 'wb') as f:
        f.write(tiny_png(width * 2, height * 2))
    return True


class ChartRuleTests(unittest.TestCase):
    def test_normalize_skips_bad_items_with_reason(self):
        good = chart('c01_x', '1', 'bar', ['a', 'b'], [{'name': 's', 'values': [1, 2]}])
        cases = [
            (dict(good, form='pie'), 'form'),
            (dict(good, id='Bad-Id'), 'id'),
            (dict(good, categories=[]), 'categories'),
            (dict(good, series=[{'name': 's', 'values': [1]}]), 'values'),
            (dict(good, series=[{'name': str(i), 'values': [1, 2]} for i in range(4)]), 'series'),
            (dict(good, series=[{'name': 's', 'values': [None, None]}]), '숫자'),
            (dict(good, source=''), '출처'),
            (dict(good, form='stacked', series=[{'name': 's', 'values': [1, -2]}]), '음수'),
        ]
        for raw, word in cases:
            with self.subTest(word=word):
                ch, why, _ = md.normalize_chart(raw)
                self.assertIsNone(ch)
                self.assertIn(word, why)
        ch, why, notes = md.normalize_chart(good)
        self.assertIsNone(why)
        self.assertEqual(notes, [])

    def test_normalize_keeps_nulls_and_drops_bad_extras(self):
        raw = chart('c01_x', 'ml', 'bar', ['a', 'b'], [{'name': 's', 'values': [1, '2']}, {'name': 't', 'values': [3, 4]}],
                    highlight=[0], errors=[[0, 2]])
        frozen = copy.deepcopy(raw)
        ch, why, notes = md.normalize_chart(raw)
        self.assertEqual(raw, frozen)                          # 원본은 그대로
        self.assertEqual(ch['series'][0]['values'], [1, None])  # 숫자가 아닌 값은 빈칸(지어 채우지 않음)
        self.assertIsNone(ch['highlight'])                      # 계열 2개면 강조 안 함
        self.assertIsNone(ch['errors'])                         # 길이가 다른 구간은 그리지 않음
        self.assertEqual(ch['area'], 'ML')
        self.assertEqual(len(notes), 2)
        pairs, _, notes = md.normalize_chart(chart('c02_x', '1', 'pairs', ['a'], [{'name': 's', 'values': [1]}],
                                                   reference=[{'value': 1, 'label': 'x'}], max=1, units=['ms']))
        self.assertEqual((pairs['reference'], pairs['max'], pairs['units']), ([], None, ['ms']))
        self.assertTrue(notes)

    def test_reference_label_does_not_repeat_value(self):
        self.assertEqual(md._ref_text({'value': 0.7, 'label': '목표 70%'}, 'ratio'), '목표 70%')
        self.assertEqual(md._ref_text({'value': 62, 'label': '30일 평균'}, 'count'), '30일 평균 62')
        self.assertEqual(md._ref_text({'value': 3, 'label': '최근 30일'}, 'count'), '최근 30일 3')
        self.assertEqual(md._ref_text({'value': 500, 'label': ''}, 'ms'), '500ms')

    def test_axis_ticks(self):
        self.assertEqual(md._domain([0.3, 0.0], 1), (0.0, 1, 0.25, [0.0, 0.25, 0.5, 0.75, 1.0]))
        lo, hi, _, ticks = md._domain([1147, 1767], 4096)
        self.assertEqual((lo, hi, ticks[-1]), (0.0, 4096, 4000.0))
        lo, hi, _, _ = md._domain([2700, 3016], zero=False)
        self.assertEqual((lo, hi), (2700.0, 3100.0))
        for v, unit, step, want in ((0.25, 'ratio', 0.25, '25%'), (0.3, 'score', 0.1, '0.3'), (2000, 'count', 1000, '2,000'),
                                    (200, 'ms', 100, '200ms'), (0.04, 'usd', 0.02, '$0.04'), (0, 'usd', 0.02, '$0')):
            self.assertEqual(md.fmt_tick(v, unit, step), want)

    def test_every_form_draws_inside_without_text_overlap(self):
        """글자 상자(글자 폭 추정)가 그림 밖으로 나가거나 서로 겹치지 않는다. 막대·점마다 툴팁이 있다."""
        for ch in md.prepare(chart_sample())['charts']:
            with self.subTest(chart=ch['id']):
                svg, w, h = md.figure_svg(ch)
                self.assertEqual(w, md.FIG_W)
                self.assertIn('viewBox="0 0 %d %d"' % (w, h), svg)
                boxes = []
                for attrs, s in _TEXT.findall(svg):
                    a = dict(_ATTR.findall(attrs))
                    b = md._tbox(float(a['x']), float(a['y']), html_lib.unescape(s), float(a['font-size']),
                                 a.get('text-anchor', 'start'))
                    self.assertTrue(b[0] >= -0.5 and b[2] <= w + 0.5 and b[1] >= -0.5 and b[3] <= h + 0.5, (s, b))
                    boxes.append((b, s))
                for i, (a, sa) in enumerate(boxes):
                    for b, sb in boxes[i + 1:]:
                        overlap = a[0] < b[2] - 0.5 and b[0] < a[2] - 0.5 and a[1] < b[3] - 0.5 and b[1] < a[3] - 0.5
                        self.assertFalse(overlap, (sa, sb))
                tips = re.findall(r'<title>([^<]*)</title>', svg)
                self.assertTrue(tips and all(' · ' in t for t in tips), tips[:3])
                self.assertNotIn('stroke-dasharray', svg)        # 점선 금지

    def test_series_colors_fixed_order_and_highlight(self):
        model = md.prepare(chart_sample())
        by_id = {c['id']: c for c in model['charts']}
        svg = md.figure_svg(by_id['c05_bar3'])[0]
        for color in md.SERIES_COLORS:
            self.assertIn('fill="%s"' % color, svg)
        svg = md.figure_svg(by_id['c04_bar'])[0]                   # 강조 1개: 진한 색 1개 + 회색 2개
        self.assertEqual(svg.count('fill="%s"/>' % md.ONE), 1)
        self.assertEqual(svg.count('fill="%s"/>' % md.OFF), 2)

    def test_headline_chart_follows_chart_rules(self):
        """한눈에 보기 작은 그래프도 같은 규칙: 지금=청록 ONE, 앞=회색 OFF, 둥근 끝, 0이면 막대 없음."""
        m = {'name': '상위 3건 적합', 'unit': 'ratio',
             'compare': {'before': 0.49, 'after': 0.61, 'label_before': '처음', 'label_after': '지금 서비스',
                         'reference': {'label': '9/28 평가', 'before': 0.24, 'after': 0.0}}}
        svg, w, _ = md.chart_svg(m)
        self.assertEqual(w, md.CHART_W)
        self.assertNotIn(md.ACCENT, svg)
        self.assertEqual(svg.count('fill="%s"/>' % md.ONE), 1)      # 지금 서비스 0.61 만 — 9/28 서비스 0.0 은 막대 없음
        self.assertEqual(svg.count('fill="%s"/>' % md.OFF), 2)
        self.assertEqual(svg.count('<rect'), 1)                       # 바탕 하나뿐 — 막대는 둥근 끝 path
        self.assertIn('>0.0%<', svg)                                  # 막대가 없어도 값 글자는 남는다
        self.assertEqual(md.chart_svg(m, W=466)[1], 466)

    def test_html_big_numbers_use_ink(self):
        self.assertIn('.big{font-size:32px;line-height:1.15;font-weight:700;color:var(--ink)', md.CSS)

    def test_short_dates_get_spaced_axis_labels(self):
        """데이터 쪽이 보내는 'M/D' 날짜도 날짜로 알아보고 몇 칸마다 하나씩(첫날·마지막 날 포함) 쓴다 — 잘리지 않게."""
        self.assertEqual(md._axis_cat('9/8'), '9/8')
        self.assertEqual(md._axis_cat('2026-09-16'), '9/16')
        self.assertEqual(md._axis_cat('지금(10/14)'), '지금(10/14)')
        cats = ['9/%d' % d for d in range(16, 31)] + ['10/%d' % d for d in range(1, 16)]
        ch, _, _ = md.normalize_chart(chart('c02_x', '1', 'column', cats, [{'name': '새 공고', 'values': list(range(30))}],
                                            unit='count'))
        for W in (md.FIG_W, md.PNG_FIG_W):
            svg = md.figure_svg(ch, W=W)[0]
            texts = [html_lib.unescape(s) for _, s in _TEXT.findall(svg)]
            shown = [t for t in texts if t in cats]
            self.assertIn('9/16', shown)
            self.assertIn('10/15', shown)
            self.assertLess(len(shown), len(cats))
            self.assertFalse(any('…' in t for t in texts), texts)

    def test_bar_value_label_sits_at_bar_end_even_past_reference(self):
        """기준선이 있어도 값 글자는 막대 끝에 — 기준선 너머에 두면 기준을 넘은 값처럼 읽힌다. 기준선은 글자 아래에 그린다."""
        ch, _, _ = md.normalize_chart(chart('c21_x', '11', 'bar', ['LLM', 'Jev(외부 채점 모델)', 'Codex'],
                                            [{'name': '일치율', 'values': [0.6434, 0.5245, 0.5]}], max=1,
                                            reference=[{'value': 0.7, 'label': '기준 70%'}]))
        svg = md.figure_svg(ch)[0]
        self.assertIn('>기준 70%<', svg)
        self.assertNotIn('70.0%', svg)                                # 같은 값을 두 번 쓰지 않는다
        line_at = svg.find('stroke="%s" stroke-width="1"/>' % md.REF)
        label_at = svg.find('>64.3%<')
        self.assertTrue(0 < line_at < label_at)                       # 선 다음에 글자(글자가 위)
        attrs = dict(_ATTR.findall(_TEXT.findall(svg[svg.rfind('<text', 0, label_at):label_at + 20])[0][0]))
        x_line = float(re.search(r'<line x1="([\d.]+)"[^>]*stroke="%s"' % md.REF, svg).group(1))
        self.assertLess(float(attrs['x']), x_line)                    # 64.3% 글자는 70% 선 앞(막대 끝)에 있다
        self.assertEqual(md._ref_text({'value': 0.7, 'label': '기준 0.70'}, 'ratio'), '기준 0.70')

    def test_reference_line_never_cuts_a_value_label(self):
        """기준선이 값 글자를 가르지 않는다(HTML·PNG 폭 모두). 가를 자리면 막대 안쪽 끝에 넣는다(막대 밖에는 없음)."""
        ch, _, _ = md.normalize_chart(chart('c21_x', '11', 'bar', ['LLM(gpt-4.1-mini)', 'Jev(외부 채점 모델)', 'Codex(다른 표본)'],
                                            [{'name': '일치율', 'values': [0.643357, 0.524476, 0.5]}], max=1, highlight=[0, 1],
                                            reference=[{'value': 0.7, 'label': '기준 70%'}]))
        for W in (md.FIG_W, 390, md.PNG_FIG_W):
            with self.subTest(W=W):
                svg = md.figure_svg(ch, W=W)[0]
                x_line = float(re.search(r'<line x1="([\d.]+)"[^>]*stroke="%s"' % md.REF, svg).group(1))
                for value in ('64.3%', '52.4%', '50.0%'):
                    found = [(dict(_ATTR.findall(a)), s) for a, s in _TEXT.findall(svg) if s == value]
                    self.assertEqual(len(found), 1, value)
                    a = found[0][0]
                    b = md._tbox(float(a['x']), float(a['y']), value, float(a['font-size']), a.get('text-anchor', 'start'))
                    self.assertFalse(b[0] - 2 <= x_line <= b[2] + 2, (value, b, x_line))
                    self.assertLess(b[2], x_line)                             # 모두 기준선 앞(막대 쪽)
                    if a.get('text-anchor') == 'end':                         # 막대 안 글자: 막대 색 위에서 읽히는 색
                        self.assertIn(a['fill'], ('#ffffff', md.INK))
        svg = md.figure_svg(ch)[0]
        self.assertIn('fill="%s"/>' % md.OFF, svg)                     # Codex(다른 표본)는 회색
        self.assertEqual(svg.count('fill="%s"/>' % md.ONE), 2)

    def test_short_bar_label_under_reference_gets_white_backing(self):
        """막대가 짧아 안에 못 넣으면 글자 뒤에 흰 바탕을 깔아 그 높이만큼 기준선을 끊는다(선 위에 그린다)."""
        ch, _, _ = md.normalize_chart(chart('c21_x', '11', 'bar', ['짧은 막대'], [{'name': '값', 'values': [0.04]}], max=1,
                                            reference=[{'value': 0.09, 'label': '기준 9%'}]))
        svg = md.figure_svg(ch)[0]
        line_at = svg.find('stroke="%s" stroke-width="1"/>' % md.REF)
        back_at = svg.find('height="14" fill="#ffffff"/>')
        self.assertTrue(0 < line_at < back_at < svg.find('>4.0%<'))
        x = float(re.search(r'<rect x="([\d.]+)"[^>]*height="14" fill="#ffffff"', svg).group(1))
        a = dict(_ATTR.findall([a for a, s in _TEXT.findall(svg) if s == '4.0%'][0]))
        self.assertAlmostEqual(x, float(a['x']) - 2, places=1)         # 흰 바탕이 그 글자 바로 뒤

    def test_grid_line_near_reference_is_skipped_but_tick_text_stays(self):
        """기준선 바로 옆 격자선은 빼고(겹선), 아래 눈금 글자는 모두 남긴다(축 눈금이 비어 보이지 않게)."""
        ch, _, _ = md.normalize_chart(chart('c17_x', '8', 'bar', ['공고 서버', '배치 중 최대'],
                                            [{'name': '메모리', 'values': [1147, 1767]}], unit='mb', max=4096,
                                            reference=[{'value': 4096, 'label': '서버 메모리 4GB'}]))
        svg = md.figure_svg(ch)[0]
        texts = [html_lib.unescape(s) for _, s in _TEXT.findall(svg)]
        self.assertIn('4,000MB', texts)
        self.assertIn('2,000MB', texts)
        x_ref = float(re.search(r'<line x1="([\d.]+)"[^>]*stroke="%s"' % md.REF, svg).group(1))
        grid_xs = [float(x) for x in re.findall(r'<line x1="([\d.]+)"[^>]*stroke="%s"' % md.GRID, svg)]
        self.assertTrue(grid_xs)
        self.assertTrue(all(abs(x - x_ref) >= md.REF_GAP for x in grid_xs), (grid_xs, x_ref))
        ch, _, _ = md.normalize_chart(chart('c21_x', '11', 'bar', ['a', 'b'], [{'name': '값', 'values': [0.64, 0.5]}], max=1,
                                            reference=[{'value': 0.7, 'label': '기준 70%'}]))
        texts = [html_lib.unescape(s) for _, s in _TEXT.findall(md.figure_svg(ch)[0])]
        self.assertTrue({'0%', '25%', '50%', '75%', '100%'} <= set(texts), texts)

    def test_pairs_share_one_axis_per_unit(self):
        """같은 단위 칸끼리는 축 하나 — 같은 ms 인데 작은 값의 막대가 더 길어지지 않는다. 칸 제목에 단위."""
        ch, _, _ = md.normalize_chart(chart('c16_x', '8', 'pairs', ['공고 1건', '검색어 1건', '전체 다시 만들기'],
                                            [{'name': 'CPU', 'values': [408.4, 178.5, 20.53]},
                                             {'name': 'GPU', 'values': [21.2, 21.7, 1.07]}],
                                            unit=None, units=['ms', 'ms', 'min']))
        bar = re.compile(r'<g><title>([^<]*)</title><rect[^>]*/><path d="M([\d.]+),[\d.]+H[\d.]+A[\d.]+,[\d.]+ 0 0 1 ([\d.]+),')
        for W in (md.FIG_W, 390, md.PNG_FIG_W):
            with self.subTest(W=W):
                svg = md.figure_svg(ch, W=W)[0]
                length = {t.split(' · ')[0] + '/' + t.split(' · ')[1]: float(x1) - float(x0) for t, x0, x1 in bar.findall(svg)}
                self.assertEqual(len(length), 6, length)
                self.assertGreater(length['공고 1건/CPU'], length['검색어 1건/CPU'])
                self.assertAlmostEqual(length['검색어 1건/CPU'] / length['공고 1건/CPU'], 178.5 / 408.4, places=2)
                self.assertAlmostEqual(length['검색어 1건/GPU'] / length['공고 1건/GPU'], 21.7 / 21.2, places=1)
                texts = [html_lib.unescape(s) for _, s in _TEXT.findall(svg)]
                self.assertIn('공고 1건 (ms)', texts)
                self.assertIn('전체 다시 만들기 (분)', texts)
        self.assertEqual(md._pair_axes(['a', 'b'], [{'name': 's', 'values': [None, None]}], ['ms', 'ms'], None)['ms'][:2],
                         (0.0, 1.0))                                   # 값이 없어도 축이 0 폭이 되지 않는다

    def test_narrow_stacked_piece_gets_outside_label(self):
        ch, _, _ = md.normalize_chart(chart('c14_x', '7', 'stacked', ['모든 공고 2,971건'],
                                            [{'name': '통과', 'values': [2690]}, {'name': '통과 못 함', 'values': [281]}],
                                            unit='count'))
        svg = md.figure_svg(ch)[0]
        texts = [html_lib.unescape(s) for _, s in _TEXT.findall(svg)]
        self.assertIn('2,690', texts)
        self.assertIn('통과 못 함 281', texts)                         # 좁은 조각의 숫자도 그림에 남는다

    def test_png_export_scale_and_ml_subtitle(self):
        model = md.prepare(chart_sample())
        jobs = [j for j in md.png_jobs(model, word=False) if j[0][0] == 'out']
        scale = (md.CARD_W - 2 * md.CARD_PAD) / float(md.PNG_FIG_W)
        self.assertLess(scale, 1.2)                                   # 막대 18px → 24px 아래(HTML 과 같은 규격)
        self.assertLessEqual(18 * scale, 24)
        for key, svg, w, h, name in jobs:
            self.assertIn('scale(%.4f)' % scale, svg)
        by_id = {c['id']: c for c in model['charts']}
        self.assertTrue(md.png_subtitle(by_id['c06_pairs']).startswith(md.ML_TAG))      # ML 그래프는 '서비스 밖 실험'
        self.assertEqual(md.png_subtitle(by_id['c01_line']), by_id['c01_line']['subtitle'])
        head = next(svg for key, svg, w, h, name in jobs if name == '00a_headline.png')
        self.assertIn(html_lib.escape(md.plain_of(model['cards'][1]['chart'])), head)   # 머리 그림에도 부제(쉬운 설명)

    def test_crop_png_rows(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'a.png')
            with open(path, 'wb') as f:
                f.write(tiny_png(4, 10))
            self.assertTrue(md.crop_png_rows(path, 6))
            self.assertEqual(png_size(path), (4, 6))
            self.assertTrue(md._png_ready(path))
            self.assertFalse(md.crop_png_rows(path, 9))                # 원래보다 크게는 못 자른다(그대로 둔다)
            self.assertEqual(png_size(path), (4, 6))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_short_svg_is_shot_taller_then_cropped(self):
        """키가 작은 그래프(c19 등)는 Edge 창을 MIN_SHOT_H 로 키워 찍고 원래 키로 자른다 — Word 에 표로 대신 들어가지 않게."""
        tmp = tempfile.mkdtemp()
        try:
            png, seen = os.path.join(tmp, 'w_c19.png'), {}

            def fake_run(cmd, **kw):
                w, h = (int(x) for x in next(a for a in cmd if a.startswith('--window-size=')).split('=')[1].split(','))
                seen['h'] = h
                with open(png, 'wb') as f:
                    f.write(tiny_png(w * 2, h * 2))
            with mock.patch.object(md.subprocess, 'run', fake_run):
                self.assertTrue(md.svg_to_png('<svg/>', 50, 40, png, 'edge', tmp, timeout=5))
            self.assertEqual(seen['h'], md.MIN_SHOT_H)
            self.assertEqual(png_size(png), (100, 80))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class ChartRenderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        write_json(self.tmp, chart_sample())
        self.log = []

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def render(self, edge=None):
        with mock.patch.object(md, 'find_edge', return_value=edge):
            return md.render(self.tmp, log=self.log.append)

    def page(self):
        html_path, _ = self.render()
        with io.open(html_path, encoding='utf-8') as f:
            return f.read()

    def test_html_one_svg_and_number_table_per_chart(self):
        page = self.page()
        self.assertEqual(page.count('<svg class="fig-svg"'), N_CHARTS)
        self.assertEqual(page.count('<summary>숫자 보기</summary>'), N_CHARTS)
        model = md.prepare(chart_sample())
        for ch in model['charts']:
            with self.subTest(chart=ch['id']):
                start = page.find('<figure class="fig" id="chart-%s">' % ch['id'])
                self.assertGreater(start, 0)
                fig = page[start:page.find('</figure>', start)]
                for want in (ch['title'], ch['subtitle'], ch['note'], '출처: ', '측정: 이번에 잼 2026-10-15',
                             '<table class="nums">'):
                    self.assertIn(want, fig)
                table = fig[fig.find('<table class="nums">'):]
                self.assertEqual(table.count('<tr><th scope="row">'), len(ch['categories']))
        pairs = page[page.find('id="chart-c06_pairs"'):]
        self.assertIn('<td>측정 못 함</td>', pairs[:pairs.find('</figure>')])     # 빈값은 표에 '측정 못 함'
        self.assertIn('기준선: 목표 70%', page)
        self.assertIn('<title>처음(벡터만) · P@3 · 49.4% (구간 43.0% ~ 56.0%)</title>', page)

    def test_charts_sit_before_their_area_table(self):
        page = self.page()
        h1 = page.find('<h2>1. 수집·운영</h2>')
        self.assertTrue(h1 < page.find('id="chart-c01_line"') < page.find('id="chart-c02_column"')
                        < page.find('누적 공고 수</td>'))
        sec3 = page[page.find('<h2>3. AI 판정</h2>'):page.find('<h2>4. 추천 정확도</h2>')]   # 지표 없는 영역: 그래프만
        self.assertIn('id="chart-c03_stacked"', sec3)
        self.assertNotIn('<div class="tbl">', sec3)
        ml = page.find('서비스에 연결되어 있지 않다')
        self.assertTrue(ml < page.find('id="chart-c06_pairs"') < page.find('리랭커 nDCG@10</td>'))

    def test_no_addresses_or_scripts(self):
        page = self.page()
        for bad in ('<script', '<link', 'http://', 'https://', '10.1.2.3', '@import', 'url(', 'xmlns'):
            self.assertNotIn(bad, page)
        self.assertIn('(주소 생략)', page)

    def test_bad_charts_logged_not_drawn(self):
        data = chart_sample()
        data['charts'] += [{'id': 'bad_form', 'form': 'pie', 'categories': ['a'], 'series': [], 'source': 's'},
                           copy.deepcopy(data['charts'][1]),
                           chart('all_null', '1', 'bar', ['a'], [{'name': 's', 'values': [None]}])]
        write_json(self.tmp, data)
        page = self.page()
        self.assertEqual(page.count('<svg class="fig-svg"'), N_CHARTS)
        joined = '\n'.join(self.log)
        for want in ('bad_form', 'c01_line: 같은 id', 'all_null: 그릴 숫자가 하나도 없다'):
            self.assertIn(want, joined)
        self.assertIn('Edge 를 찾지 못해', joined)

    def test_without_charts_same_as_before(self):
        old = md.build_html(md.prepare(sample()))
        for charts in ([], None):
            data = sample()
            data['charts'] = charts
            self.assertEqual(md.build_html(md.prepare(data)), old)
        self.assertNotIn('fig-svg', old)
        write_json(self.tmp, sample())
        with mock.patch.object(md, 'svg_to_png', fake_svg_to_png):
            self.render(edge='edge')
        self.assertNotIn(md.CHART_DIR, os.listdir(self.tmp))
        self.assertEqual(self.log, [])

    @unittest.skipUnless(HAS_DOCX, 'python-docx 없음')
    def test_docx_without_edge_uses_number_tables(self):
        from docx import Document
        _, docx_path = self.render()
        doc = Document(docx_path)
        self.assertEqual(len(doc.inline_shapes), 0)
        self.assertEqual(len(doc.tables), AREA_TABLES + 1 + N_CHARTS)
        body = '\n'.join(p.text for p in doc.paragraphs)
        cells = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for cid in REPORT_ORDER:
            self.assertIn(cid + ' 제목', body)
            self.assertIn(cid + ' 쉬운 설명', body)
        self.assertIn('기준선: 목표 70%', body)
        self.assertIn('측정 못 함', cells)
        self.assertNotIn('10.1.2.3', body + cells)
        self.assertNotIn(md.CHART_DIR, os.listdir(self.tmp))

    @unittest.skipUnless(HAS_DOCX, 'python-docx 없음')
    def test_pictures_and_png_folder_with_fake_edge(self):
        from docx import Document
        folder = os.path.join(self.tmp, md.CHART_DIR)
        os.makedirs(folder)
        for name in ('99_old_chart.png', '메모.txt'):          # 예전 그림은 지우고, 사람이 둔 파일은 둔다
            with open(os.path.join(folder, name), 'wb') as f:
                f.write(b'x')
        with mock.patch.object(md, 'svg_to_png', fake_svg_to_png):
            _, docx_path = self.render(edge='edge')
        self.assertEqual(sorted(os.listdir(folder)), sorted(PNG_NAMES + ['메모.txt']))
        for name in PNG_NAMES:
            self.assertEqual(png_size(os.path.join(folder, name))[0], md.CARD_W * 2)
        doc = Document(docx_path)
        self.assertEqual(len(doc.inline_shapes), 1 + N_CHARTS)
        self.assertEqual(len(doc.tables), AREA_TABLES)
        self.assertEqual(set(os.listdir(self.tmp)), {'metrics.json', md.HTML_NAME, md.DOCX_NAME, md.CHART_DIR})
        self.assertTrue(any('PNG %d장' % len(PNG_NAMES) in x for x in self.log), self.log)

    def test_png_folder_without_python_docx(self):
        with mock.patch.object(md, 'svg_to_png', fake_svg_to_png), mock.patch.object(md, '_docx', return_value=False):
            _, docx_path = self.render(edge='edge')
        self.assertIsNone(docx_path)
        self.assertEqual(sorted(os.listdir(os.path.join(self.tmp, md.CHART_DIR))), sorted(PNG_NAMES))

    def test_failed_png_falls_back_to_table(self):
        def half(svg, width, height, png_path, edge, work, timeout=40):
            if os.path.basename(png_path).startswith('w_c04'):
                return False
            return fake_svg_to_png(svg, width, height, png_path, edge, work)
        with mock.patch.object(md, 'svg_to_png', half), mock.patch.object(md, '_docx', return_value=HAS_DOCX):
            _, docx_path = self.render(edge='edge')
        if HAS_DOCX:
            from docx import Document
            doc = Document(docx_path)
            self.assertEqual(len(doc.inline_shapes), N_CHARTS)          # 그래프 1개만 표로
            self.assertEqual(len(doc.tables), AREA_TABLES + 1)
            # 표로 대신 넣은 것은 실행 화면에 한 줄 남긴다
            self.assertEqual([x for x in self.log if '숫자 표로 넣었다' in x], ['Word: 그래프 그림을 만들지 못해 숫자 표로 넣었다: c04_bar'])

    @unittest.skipUnless(HAS_DOCX and HAS_EDGE, 'python-docx 또는 Edge 없음')
    def test_real_edge_pictures_and_pngs(self):
        from docx import Document
        _, docx_path = md.render(self.tmp, log=self.log.append)
        folder = os.path.join(self.tmp, md.CHART_DIR)
        self.assertEqual(sorted(os.listdir(folder)), sorted(PNG_NAMES))
        for name in PNG_NAMES:
            w, h = png_size(os.path.join(folder, name))
            self.assertEqual(w, md.CARD_W * 2)
            self.assertGreater(h, 400)
        self.assertEqual(len(Document(docx_path).inline_shapes), 1 + N_CHARTS)
        self.assertEqual(set(os.listdir(self.tmp)), {'metrics.json', md.HTML_NAME, md.DOCX_NAME, md.CHART_DIR})


if __name__ == '__main__':
    unittest.main()
