# -*- coding: utf-8 -*-
"""지표 결과 한 개(metrics.json)를 지표_보고서.html · 지표_보고서.docx 로 만든다(2026-10-14).

  python -X utf8 -m eval.metrics_doc reports/metrics_<UTC 시각>Z     폴더를 빼면 가장 최근 metrics_ 폴더

- metrics.json 은 읽기만 한다(고치지 않는다). 두 결과 파일은 같은 폴더에 쓴다.
- HTML 은 파일 하나로 끝난다. 글꼴·스크립트·그림을 밖에서 받지 않고, 그래프는 SVG 를 그대로 넣는다.
- Word 는 python-docx 로 만든다(requirements-report.txt). 설치돼 있지 않으면 HTML 만 만든다.
  그래프는 같은 SVG 를 이 PC 의 Edge(headless) 화면 저장으로 PNG 로 바꿔 넣는다(그리기 패키지를 새로 쓰지 않는다).
  Edge 가 없거나 화면 저장에 실패하면 그래프 숫자를 작은 표로 대신 넣는다.
- 문서에는 IP 주소·호스트가 든 주소(URL)·키·비밀번호를 쓰지 않는다. 입력 글에 섞여 있어도 지우고 쓴다.
- metrics.json 에 "charts"(영역 그래프 목록, 2026-10-15)가 있으면 영역마다 표 바로 앞에 그래프를 넣는다.
  모양은 line · column · bar · stacked · pairs 다섯 가지. 그래프마다 '숫자 보기' 표와 마우스 올림 값이 있다.
  같은 그래프를 결과 폴더의 그래프/ 에 발표용 PNG(가로 1200px, 제목·부제·설명·출처 포함)로도 저장한다
  (NN_<id>.png 는 보고서 순서, 한눈에 보기 그래프는 00a_·00b_). 모양이 틀린 항목은 그리지 않고 이유를 실행 로그에 쓴다.
  charts 가 없거나 비면 예전과 똑같이 만든다.

입력 약속(metrics.json): run_at · conditions · metrics[area·name·value·unit·target·detail·plain·source·kind·
measured_at·compare] · headline[title·values·chart·note] · missing[name·reason·expected]
· charts[id·area·title·subtitle·form·unit·units·categories·series·highlight·max·reference·errors·note·source·kind·
measured_at](없어도 됨).
"""
import argparse
import glob
import html
import json
import math
import os
import pathlib
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import zlib
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

try:  # 예전 PDF 렌더러의 Edge 경로 목록과 쉬운 설명 사전을 그대로 쓴다
    from eval.metrics_pdf import EDGE as _EDGE_PATHS, PLAIN as _OLD_PLAIN
except ImportError:  # pragma: no cover - metrics_pdf 가 없어져도 돈다
    _EDGE_PATHS = [r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
                   r'C:\Program Files\Microsoft\Edge\Application\msedge.exe']
    _OLD_PLAIN = {}

HTML_NAME = '지표_보고서.html'
DOCX_NAME = '지표_보고서.docx'
TITLE = '공고 파트 작업 지표'

AREAS = [('1', '수집·운영'), ('2', '전처리 품질'), ('3', 'AI 판정'), ('4', '추천 정확도'),
         ('5', '신청 불가 차단·자격 판정'), ('6', '가산점'), ('7', '조율 연결'), ('8', '서버 성능'),
         ('9', '비용'), ('10', '개발 품질'), ('11', '약점과 한계')]
ML_TITLE = '서비스 밖 실험(ML)'
ML_LEAD = ('이 부분의 실험은 서비스에 연결되어 있지 않다. 추천·자격 판정 결과에 쓰이지 않으며, '
           '모델을 따로 만들어 본 결과만 적는다.')
OTHER_TITLE = '그 밖의 지표'
KIND = {'live': '이번에 잼', 'record': '기록'}
NA = '측정 못 함'
COLUMNS = ['지표', '값', '기준·목표', '쉬운 설명', '근거', '측정']
PASS_WORDS = ('통과', '성공', '합격', 'pass')
KST = timezone(timedelta(hours=9))

RERUN_CMD = ['cd data-collection', r'.venv\Scripts\python.exe -X utf8 -m eval.metrics_report --all']
RERUN_POINTS = [
    '공용 DB는 읽기만 한다(쓰지 않는다).',
    '배치 서버의 배치 기록(로그)을 읽기 전용으로 복사해 온다.',
    '평가 질의 58개를 이 PC의 공고 서버와 실제 서버에 보내 응답을 잰다.',
    '단위 테스트를 돌린다.',
    '결과는 새 폴더 reports/metrics_<UTC 시각>Z/ 에 쓴다. 이전 결과 폴더는 덮지 않는다.',
    '서버 주소는 옵션으로 바꿀 수 있다(팀 계정으로 옮긴 뒤에는 새 주소를 넣는다).',
]
DOC_CMD = r'.venv\Scripts\python.exe -X utf8 -m eval.metrics_doc reports\metrics_<UTC 시각>Z'

# 색 — 흰 바탕, 짙은 회색 글자, 강조색은 짙은 청록 하나
INK, MUTED, FAINT, LINE, SOFT = '#1f2933', '#56616b', '#7b858e', '#dde2e5', '#f4f6f7'
ACCENT, ACCENT_SOFT, BEFORE = '#0f5c57', '#e8f1f0', '#aeb8bd'
FONT_STACK = '"Malgun Gothic", "맑은 고딕", "Apple SD Gothic Neo", "Noto Sans KR", sans-serif'
DOCX_FONT = '맑은 고딕'


# ── 글 정리: 주소·키·비밀번호 지우기 ─────────────────────────────

_SCRUB = [
    (re.compile(r'\b[a-zA-Z][a-zA-Z0-9+.\-]*://\S+'), '(주소 생략)'),
    (re.compile(r'\b(?:\d{1,3}\.){3}\d{1,3}(?::\d{1,5})?\b'), '(주소 생략)'),
    # 호스트 이름(도메인). search.app 같은 모듈 이름과 겹치는 끝말(app·dev 등)은 넣지 않는다.
    (re.compile(r'\b(?:[a-zA-Z0-9-]+\.)+(?:com|net|org|io|kr|xyz|link|cloud)(?::\d{1,5})?(?:/\S*)?(?!\w|\.\w)'),
     '(주소 생략)'),
    (re.compile(r'\bsk-[A-Za-z0-9_\-]{12,}'), '(키 생략)'),
    (re.compile(r'\bAKIA[0-9A-Z]{16}\b'), '(키 생략)'),
    (re.compile(r'(?i)\b(password|passwd|pwd|secret|api[_-]?key|access[_-]?key|token)\s*[=:]\s*\S+'), r'\1=(생략)'),
]


def scrub(text):
    """문서에 쓰기 전에 주소·키·비밀번호 모양을 지운다."""
    if text is None:
        return ''
    s = str(text)
    for pat, rep in _SCRUB:
        s = pat.sub(rep, s)
    return s


# ── 값 모양 ─────────────────────────────────────────────

def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _pct(x):
    s = '%.1f' % (x * 100)
    return (s[:-2] if s.endswith('.0') else s) + '%'


def fmt_number(v, unit):
    """숫자 하나를 단위에 맞춰 쓴다."""
    if unit == 'ratio':
        return '%.1f%%' % (v * 100)
    if unit == 'score':
        return '%.3f' % v
    if unit == 'count':
        return format(int(v), ',') if float(v).is_integer() else format(v, ',.2f')
    if unit == 'ms':
        return (format(v, ',.0f') if abs(v) >= 10 else '%.1f' % v) + 'ms'
    if unit == 'sec':
        return ('%.2f' % v if abs(v) < 10 else format(v, ',.1f')) + '초'
    if unit == 'min':
        return format(v, ',.1f') + '분'
    if unit == 'mb':
        return (format(v, ',.0f') if abs(v) >= 100 else '%.1f' % v) + 'MB'
    if unit == 'usd':
        if v != 0 and abs(v) < 0.01:
            return '$%.4f' % v
        if abs(v) < 1:
            # 1달러 미만은 셋째 자리까지(시간당 $0.052 × 730 = 한 달 $37.96 처럼 표와 그래프가 맞게), 끝의 0 하나는 뺀다
            s = '%.3f' % v
            return '$' + (s[:-1] if s.endswith('0') else s)
        return '$' + format(v, ',.2f')
    if float(v).is_integer():
        return format(int(v), ',')
    return '%.3g' % v


def pass_total(m):
    """'16회 중 16회 통과'처럼 목표가 분모인 개수면 그 분모를 돌려준다. 아니면 None."""
    if m.get('unit') != 'count':
        return None
    v = m.get('value')
    if not _is_num(v) or not float(v).is_integer():
        return None
    t = str(m.get('target') if m.get('target') is not None else '').replace(',', '').strip()
    if not t.isdigit():
        return None
    words = ('%s %s' % (m.get('name') or '', m.get('detail') or '')).lower()
    if not any(w in words for w in PASS_WORDS):
        return None
    total = int(t)
    if total <= 0 or v > total:
        return None
    return total


def fmt_value(m):
    """값 칸에 쓸 글. 값이 없으면 '측정 못 함'."""
    v = m.get('value')
    if v is None:
        return NA
    unit = m.get('unit')
    if unit == 'text' or not _is_num(v):
        return scrub(v)
    total = pass_total(m)
    if total is not None:
        return '%s/%s' % (format(int(v), ','), format(total, ','))
    return fmt_number(v, unit)


_CMP = re.compile(r'^\s*(≥|>=|≤|<=|>|<|=)?\s*(-?\d+(?:\.\d+)?)\s*$')


def fmt_target(m):
    """기준·목표 칸. 비율 단위의 '0'·'≥0.95' 같은 숫자 목표는 %로, 시간·비용은 단위를 붙인다."""
    t = m.get('target')
    if t is None or str(t).strip() == '':
        return '—'
    s = scrub(t).strip()
    unit = m.get('unit')
    hit = _CMP.match(s)
    if hit and unit in ('ratio', 'ms', 'sec', 'min', 'mb', 'usd'):
        op = {'>=': '≥', '<=': '≤'}.get(hit.group(1), hit.group(1) or '')
        x = float(hit.group(2))
        if unit == 'ratio':
            return op + _pct(x) if abs(x) <= 1 else s
        return op + fmt_number(x, unit)
    return s


def _fmt_when(s):
    if s is None or str(s).strip() == '':
        return ''
    s = str(s).strip()
    if 'T' in s:
        try:
            dt = datetime.fromisoformat(s.replace('Z', '+00:00'))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=KST)
            return dt.astimezone(KST).strftime('%Y-%m-%d %H:%M')
        except ValueError:
            pass
    return scrub(s)


def fmt_run_at(s):
    """측정 시각을 한국 시각으로."""
    try:
        dt = datetime.fromisoformat(str(s).replace('Z', '+00:00'))
    except (TypeError, ValueError):
        return scrub(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=KST)
    return dt.astimezone(KST).strftime('%Y-%m-%d %H:%M') + ' KST'


def fmt_measure(m):
    """(측정 방식, 측정 날짜)."""
    kind = m.get('kind')
    return KIND.get(kind, scrub(kind) if kind else '—'), _fmt_when(m.get('measured_at'))


def plain_of(m):
    return scrub(m.get('plain') or _OLD_PLAIN.get(m.get('name'), ''))


def null_reason(m, missing_by_name):
    reason = (m.get('detail') or '').strip()
    if not reason:
        reason = ((missing_by_name.get(m.get('name')) or {}).get('reason') or '').strip()
    return scrub(reason) or '이유가 기록되지 않았다.'


# ── 입력 정리 ───────────────────────────────────────────

def prepare(data):
    """metrics.json 을 그리기 좋게 묶는다. 원본 dict 는 바꾸지 않는다."""
    metrics = [m for m in (data.get('metrics') or []) if isinstance(m, dict)]
    missing = [x for x in (data.get('missing') or []) if isinstance(x, dict)]
    missing_by_name = {x.get('name'): x for x in missing}
    by_name = {m.get('name'): m for m in metrics}
    conds = data.get('conditions') or []
    if isinstance(conds, str):
        conds = [conds]
    known = {k for k, _ in AREAS}
    charts, chart_log = read_charts(data.get('charts'))

    def chart_key(ch):
        return ch['area'] if ch['area'] in known or ch['area'] == 'ML' else 'other'

    sections = []
    for key, title in AREAS:
        part = [m for m in metrics if str(m.get('area')).strip() == key]
        figs = [ch for ch in charts if chart_key(ch) == key]
        if part or figs:
            sections.append({'key': key, 'title': '%s. %s' % (key, title), 'metrics': part, 'ml': False,
                             'charts': figs})
    other = [m for m in metrics if str(m.get('area')).strip() not in known and str(m.get('area')).strip().upper() != 'ML']
    figs = [ch for ch in charts if chart_key(ch) == 'other']
    if other or figs:
        sections.append({'key': 'other', 'title': OTHER_TITLE, 'metrics': other, 'ml': False, 'charts': figs})
    ml = [m for m in metrics if str(m.get('area')).strip().upper() == 'ML']
    figs = [ch for ch in charts if chart_key(ch) == 'ML']
    if ml or figs:
        sections.append({'key': 'ML', 'title': ML_TITLE, 'metrics': ml, 'ml': True, 'charts': figs})
    ordered = [ch for sec in sections for ch in sec['charts']]       # 보고서에 나오는 순서로 번호를 매긴다
    for no, ch in enumerate(ordered, 1):
        ch['no'] = no
        ch['png'] = '%02d_%s.png' % (no, ch['id'])
    cards = []
    for c in data.get('headline') or []:
        if not isinstance(c, dict):
            continue
        values = []
        for v in c.get('values') or []:
            if not isinstance(v, dict):
                continue
            m = by_name.get(v.get('metric'))
            if m is None:
                reason = ((missing_by_name.get(v.get('metric')) or {}).get('reason')
                          or '이번 결과에 이 지표가 없다.')
                values.append({'label': scrub(v.get('label')), 'text': NA, 'na': True, 'reason': scrub(reason)})
            elif m.get('value') is None:
                values.append({'label': scrub(v.get('label')), 'text': NA, 'na': True,
                               'reason': null_reason(m, missing_by_name)})
            else:
                values.append({'label': scrub(v.get('label')), 'text': fmt_value(m), 'na': False, 'reason': ''})
        chart_metric = by_name.get(c.get('chart')) if c.get('chart') else None
        groups = chart_groups(chart_metric) if chart_metric else []
        cards.append({'title': scrub(c.get('title')), 'values': values, 'note': scrub(c.get('note')),
                      'chart': chart_metric if groups else None})
    return {'run_at': fmt_run_at(data.get('run_at')), 'conditions': [scrub(x) for x in conds],
            'sections': sections, 'cards': cards, 'missing': missing, 'missing_by_name': missing_by_name,
            'charts': ordered, 'chart_log': chart_log}


def read_charts(raw):
    """metrics.json 의 charts 목록을 검사한다. (그릴 그래프 목록, 실행 로그 줄 목록). 없거나 비면 ([], [])."""
    if raw is None:
        return [], []
    if not isinstance(raw, list):
        return [], ['charts 가 목록이 아니어서 그래프를 하나도 그리지 않았다']
    charts, log, seen = [], [], set()
    for c in raw:
        ch, why, notes = normalize_chart(c)
        name = scrub(c.get('id')) if isinstance(c, dict) and c.get('id') else '(id 없음)'
        if ch is None:
            log.append('그래프 뺌 %s: %s' % (name, why))
            continue
        if ch['id'] in seen:
            log.append('그래프 뺌 %s: 같은 id 가 앞에 있다' % name)
            continue
        seen.add(ch['id'])
        log.extend('그래프 %s: %s' % (name, x) for x in notes)
        charts.append(ch)
    return charts, log


# ── 그래프(SVG) ─────────────────────────────────────────

CHART_W = 380


def chart_groups(m):
    """compare → [(묶음 이름, [(이름, 값), (이름, 값)]), ...]. 그릴 숫자가 없으면 []."""
    c = (m or {}).get('compare')
    if not isinstance(c, dict):
        return []
    lb = scrub(c.get('label_before') or '이전')
    la = scrub(c.get('label_after') or '이후')
    groups = []
    if _is_num(c.get('before')) or _is_num(c.get('after')):
        groups.append(('이번 측정', [(lb, c.get('before')), (la, c.get('after'))]))
    ref = c.get('reference')
    if isinstance(ref, dict) and (_is_num(ref.get('before')) or _is_num(ref.get('after'))):
        # 참고 묶음은 자기 이름표가 있으면 그것을 쓴다(예: 9/28 평가의 "그때 서비스" — "지금 서비스"로 읽히지 않게)
        rb = scrub(ref.get('label_before') or c.get('label_before') or '이전')
        ra = scrub(ref.get('label_after') or c.get('label_after') or '이후')
        groups.append((scrub(ref.get('label') or '참고'), [(rb, ref.get('before')), (ra, ref.get('after'))]))
    return groups


def _chart_value(v, unit):
    if not _is_num(v):
        return NA
    return fmt_number(v, unit if unit != 'text' else None)


def _text_w(s, size):
    w = 0.0
    for ch in s:
        if ord(ch) >= 0x2E80 or ch in '→←↑↓—―':      # 한글·한자, 화살표·긴 줄표는 글자 한 칸 폭
            w += size * 0.98
        elif ch.isupper() or ch.isdigit() or ch in '%$@#&':
            w += size * 0.62
        elif ch in ' .,:;|!il()[]\'"·':
            w += size * 0.32
        else:
            w += size * 0.54
    return w


def _fit(s, size, limit):
    if _text_w(s, size) <= limit:
        return s
    while s and _text_w(s + '…', size) > limit:
        s = s[:-1]
    return s + '…'


def chart_svg(m, W=CHART_W):
    """metric 의 compare 로 가로 막대 그래프를 만든다. (svg, 폭, 높이). 그릴 것이 없으면 (None, 0, 0).

    영역 그래프와 같은 규칙: 나중 값(지금 서비스)은 청록 ONE, 앞의 값은 회색 OFF, 막대 끝은 4px 둥글게,
    값이 0이면 막대 없이 값 글자만(작은 막대가 작은 값처럼 보이지 않게)."""
    groups = chart_groups(m)
    if not groups:
        return None, 0, 0
    unit = m.get('unit')
    e = html.escape
    fs, pad, row_h = 13, 12, 26
    labels = [lbl for _, rows in groups for lbl, _ in rows]
    texts = [_chart_value(v, unit) for _, rows in groups for _, v in rows]
    label_w = min(max(_text_w(x, fs) for x in labels) + 12, 150)
    value_w = max(_text_w(x, fs) for x in texts) + 14
    bar_x = pad + label_w
    bar_max = max(W - bar_x - value_w - pad, 60)
    nums = [abs(v) for _, rows in groups for _, v in rows if _is_num(v)]
    vmax = max(nums) if nums and max(nums) > 0 else 1.0
    parts = []
    y = pad
    caption = _fit(scrub(m.get('name')), 12, W - 2 * pad)
    parts.append('<text x="%d" y="%d" font-size="12" fill="%s">%s</text>' % (pad, y + 12, MUTED, e(caption)))
    y += 26
    for gi, (glabel, rows) in enumerate(groups):
        if gi:
            y += 10
        parts.append('<text x="%d" y="%d" font-size="12" font-weight="700" fill="%s">%s</text>'
                     % (pad, y + 12, INK, e(_fit(glabel, 12, W - 2 * pad))))
        y += 20
        parts.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="1"/>'
                     % (bar_x, y + 3, bar_x, y + row_h * len(rows) - 3, AXIS))
        for i, (lbl, v) in enumerate(rows):
            strong = i == len(rows) - 1
            mid = y + row_h / 2
            parts.append('<text x="%.1f" y="%.1f" font-size="%d" text-anchor="end" fill="%s">%s</text>'
                         % (bar_x - 8, mid + 4.5, fs, MUTED, e(_fit(lbl, fs, label_w - 12))))
            if _is_num(v):
                bw = bar_max * abs(v) / vmax
                if v != 0:
                    parts.append(_hbar(bar_x, bar_x + bw, mid - 7, 14, ONE if strong else OFF))
                parts.append('<text x="%.1f" y="%.1f" font-size="%d" font-weight="%s" fill="%s">%s</text>'
                             % (bar_x + bw + 6, mid + 4.5, fs, '700' if strong else '400', INK,
                                e(_chart_value(v, unit))))
            else:
                parts.append('<text x="%.1f" y="%.1f" font-size="%d" fill="%s">%s</text>'
                             % (bar_x + 6, mid + 4.5, fs, FAINT, NA))
            y += row_h
    height = int(math.ceil(y + pad))
    summary = '; '.join('%s: %s' % (g, ', '.join('%s %s' % (lbl, _chart_value(v, unit)) for lbl, v in rows))
                        for g, rows in groups)
    svg = ('<svg role="img" aria-label="%s" width="%d" height="%d" viewBox="0 0 %d %d" '
           'font-family="\'Malgun Gothic\',\'Apple SD Gothic Neo\',sans-serif">'
           '<title>%s</title><rect width="%d" height="%d" fill="#ffffff"/>%s</svg>'
           % (e(scrub(m.get('name'))), W, height, W, height, e(summary), W, height,
              ''.join(parts)))
    return svg, W, height


# ── 영역 그래프(metrics.json 의 "charts") ─────────────────
#
# 값은 데이터 쪽(metrics_report.py)이 넣고, 여기서는 그리기만 한다. 항목마다:
#   id · area · title · subtitle · form · unit · units(pairs) · categories · series[{name, values}](1~3개)
#   · highlight · max · reference[{value, label}] · errors · note · source · kind · measured_at
# 모양이 틀린 항목은 그리지 않고 이유를 실행 로그에 남긴다(빈 값을 숫자로 채우지 않는다).

FORMS = ('line', 'column', 'bar', 'stacked', 'pairs')
UNITS = ('ratio', 'score', 'count', 'ms', 'sec', 'min', 'mb', 'usd')
SERIES_COLORS = ('#00897b', '#eb6834', '#4a3aa7')   # 색약 검사를 통과한 조합 — 순서를 바꾸지 않는다
ONE, OFF = SERIES_COLORS[0], '#aeb8bd'              # 계열 하나: 강조 막대 / 강조 안 된 막대
GRID, AXIS, REF = '#e6eaec', '#c3cbcf', MUTED
FIG_W = 360            # 그래프 그림 폭(화면에서는 카드 폭에 맞춰 늘고 준다 — 휴대폰에서도 0.9배 아래로 잘 안 줄게)
CARD_W = 600           # 발표 자료용 PNG 폭(2배로 저장 → 가로 1200px)
CARD_PAD = 28          # 발표용 그림 안쪽 여백
# 발표용 PNG 는 그래프를 이 배율로만 키운다(HTML 최대 폭 420/360 과 같게) — 그래프를 그만큼 넓게 그려
# 막대 두께(≤24px)·둥근 끝·틈이 HTML 과 같은 규격이 되고, 날짜 칸도 넓어진다.
PNG_SCALE = 420 / 360.0
PNG_FIG_W = int(round((CARD_W - 2 * CARD_PAD) / PNG_SCALE))
ML_TAG = '서비스 밖 실험(서비스에는 안 씀)'
CHART_DIR = '그래프'
EDGE_WORKERS = 4       # Edge 를 한꺼번에 몇 개까지 띄울지
LS, TS = 12.5, 11      # 이름표·눈금 글자 크기
SIDE = 6               # 그림 좌우 여백
SVG_FONT = "'Malgun Gothic','Apple SD Gothic Neo','Noto Sans KR',sans-serif"
_ID = re.compile(r'^[a-z0-9_]{1,64}$')
# 날짜 항목: 'YYYY-MM-DD' 와 데이터 쪽이 보내는 짧은 꼴 'M/D'(예: '9/8', '10/15') 둘 다
_DAY = re.compile(r'^(?:\d{4}-(\d{1,2})-(\d{1,2})|(\d{1,2})/(\d{1,2}))$')
_PNG_OUT = re.compile(r'^\d{2,3}[a-z]?_[a-z0-9_]+\.png$')


def _clean_values(values):
    out, bad = [], 0
    for v in values:
        if _is_num(v):
            out.append(v)
        else:
            bad += v is not None
            out.append(None)
    return out, bad


def normalize_chart(c):
    """charts 항목 하나를 검사해 그리기 좋게 다듬는다. 원본은 바꾸지 않는다.

    (그래프, None, 알림 목록) 또는 (None, 뺀 이유, [])."""
    if not isinstance(c, dict):
        return None, '항목 모양이 아니다', []
    cid = c.get('id')
    if not isinstance(cid, str) or not _ID.match(cid):
        return None, 'id 는 영문 소문자·숫자·밑줄이어야 한다', []
    form = c.get('form')
    if form not in FORMS:
        return None, '모르는 form(%s)' % scrub(form), []
    cats = c.get('categories')
    if not isinstance(cats, list) or not cats:
        return None, 'categories 가 비었다', []
    n = len(cats)
    raw_series = c.get('series')
    if not isinstance(raw_series, list) or not 1 <= len(raw_series) <= len(SERIES_COLORS):
        return None, 'series 는 1~3개여야 한다', []
    notes, series = [], []
    for s in raw_series:
        if not isinstance(s, dict) or not isinstance(s.get('values'), list) or len(s['values']) != n:
            return None, 'series 의 values 개수가 categories 개수와 다르다', []
        vals, bad = _clean_values(s['values'])
        if bad:
            notes.append('숫자가 아닌 값 %d개는 빈칸(측정 못 함)으로 뒀다' % bad)
        series.append({'name': scrub(s.get('name')).strip() or '값', 'values': vals})
    if not any(v is not None for s in series for v in s['values']):
        return None, '그릴 숫자가 하나도 없다', []
    if form == 'stacked' and any(v is not None and v < 0 for s in series for v in s['values']):
        return None, 'stacked 에 음수가 있다(부분-전체 그래프로 그릴 수 없다)', []
    source = scrub(c.get('source')).strip()
    if not source:
        return None, '출처(source)가 없다', []
    unit = c.get('unit') if c.get('unit') in UNITS else None
    units = None
    if form == 'pairs':
        raw_units = c.get('units')
        if isinstance(raw_units, list) and len(raw_units) == n:
            units = [u if u in UNITS else unit for u in raw_units]
        else:
            units = [unit] * n
    highlight = None
    if len(series) == 1 and form in ('bar', 'column') and isinstance(c.get('highlight'), list):
        highlight = {i for i in c['highlight'] if isinstance(i, int) and not isinstance(i, bool) and 0 <= i < n}
        highlight = highlight or None
    mx = c.get('max') if _is_num(c.get('max')) and c.get('max') > 0 else None
    refs = [{'value': r['value'], 'label': scrub(r.get('label')).strip()}
            for r in (c.get('reference') or []) if isinstance(r, dict) and _is_num(r.get('value'))]
    errors, raw_err = None, c.get('errors')
    if isinstance(raw_err, list) and len(raw_err) == n and form != 'stacked':
        errors = [(e[0], e[1]) if isinstance(e, (list, tuple)) and len(e) == 2 and _is_num(e[0]) and _is_num(e[1])
                  and e[0] <= e[1] else None for e in raw_err]
        errors = errors if any(errors) else None
    elif raw_err is not None:
        notes.append('errors 모양이 맞지 않아 구간을 그리지 않았다')
    if form == 'pairs' and (refs or mx is not None):
        notes.append('pairs 는 단위마다 축이 달라 기준선·max 를 쓰지 않았다')
        refs, mx = [], None
    area = str(c.get('area') if c.get('area') is not None else '').strip()
    return {'id': cid, 'area': 'ML' if area.upper() == 'ML' else area, 'form': form,
            'title': scrub(c.get('title')).strip() or cid, 'subtitle': scrub(c.get('subtitle')).strip(),
            'unit': unit, 'units': units, 'categories': [scrub(x).strip() or '—' for x in cats],
            'series': series, 'highlight': highlight, 'max': mx, 'reference': refs, 'errors': errors,
            'note': scrub(c.get('note')).strip(), 'source': source, 'kind': c.get('kind'),
            'measured_at': c.get('measured_at')}, None, notes


def _unit_at(ch, i):
    return ch['units'][i] if ch['form'] == 'pairs' else ch['unit']


def chart_table(ch):
    """그래프의 모든 값을 표로. (머리줄, 줄 목록, 기준선 글 목록)."""
    ser, errs = ch['series'], ch['errors']
    head = ['항목'] + [s['name'] for s in ser]
    if errs:
        head.append('구간' if len(ser) == 1 else '구간(%s)' % ser[0]['name'])
    rows = []
    for i, cat in enumerate(ch['categories']):
        u = _unit_at(ch, i)
        row = [cat] + [_chart_value(s['values'][i], u) for s in ser]
        if errs:
            e = errs[i]
            row.append('%s ~ %s' % (_chart_value(e[0], u), _chart_value(e[1], u)) if e else '—')
        rows.append(row)
    refs = [_ref_text(r, ch['unit']) for r in ch['reference']]
    return head, rows, refs


def chart_foot(ch):
    kind = KIND.get(ch.get('kind'), scrub(ch.get('kind')) if ch.get('kind') else '—')
    when = _fmt_when(ch.get('measured_at'))
    return '출처: %s · 측정: %s%s' % (scrub(ch.get('source')) or '—', kind, ' ' + when if when else '')


# 축 ─────────────────────────────

def _nice_step(span, n=4):
    raw = span / float(n)
    k = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5):
        if raw <= m * k * (1 + 1e-9):
            return m * k
    return 10 * k


def _domain(values, fixed_max=None, zero=True, n=4):
    """축 범위와 눈금. (아래 끝, 위 끝, 눈금 간격, 눈금 목록). fixed_max 가 값보다 크면 위 끝을 거기에 맞춘다."""
    vals = [v for v in values if _is_num(v)] or [0]
    lo, hi = min(vals), max(vals)
    if zero:
        lo, hi = min(lo, 0), max(hi, 0)
    keep = fixed_max is not None and fixed_max >= hi
    if keep:
        hi = fixed_max
    if hi - lo <= 0:
        if zero:
            hi = 1.0
        else:
            pad = abs(lo) * 0.1 or 1.0
            lo, hi = lo - pad, hi + pad
    step = _nice_step(hi - lo, n)
    lo2 = math.floor(lo / step + 1e-9) * step
    hi2 = hi if keep else math.ceil(hi / step - 1e-9) * step
    if hi2 <= lo2:
        hi2 = lo2 + step
    ticks, i = [], 0
    while lo2 + i * step <= hi2 + step * 1e-6:
        t = lo2 + i * step
        ticks.append(0.0 if abs(t) < step * 1e-9 else t)
        i += 1
    return lo2, hi2, step, ticks


def _decimals(x):
    for d in range(5):
        if abs(round(x, d) - x) < 1e-9 * max(1.0, abs(x)):
            return d
    return 4


def fmt_tick(v, unit, step):
    """눈금 글. 단위 규칙은 fmt_number 와 같고 자릿수만 눈금 간격에 맞춘다."""
    if abs(v) < abs(step) * 1e-9:
        v = 0.0
    if unit == 'ratio':
        return '%.*f%%' % (_decimals(step * 100), v * 100)
    d = _decimals(step)
    if unit == 'score':
        return '0' if v == 0 else '%.*f' % (max(d, 1), v)
    if unit == 'usd':
        return '$' + format(v, ',.%df' % (0 if v == 0 or d == 0 else max(d, 2)))
    return format(v, ',.%df' % d) + {'ms': 'ms', 'sec': '초', 'min': '분', 'mb': 'MB'}.get(unit, '')


def _ref_text(r, unit):
    """기준선 이름표. 이름에 값이 이미 들어 있으면('목표 70%') 값을 다시 붙이지 않는다."""
    v, label = r['value'], r['label']
    value = _chart_value(v, unit)
    if not label:
        return value
    same = {value, '%g' % v, value.replace(',', ''), '%.2f' % v}
    if unit == 'ratio':
        same.add(_pct(v))
    for s in same:
        if re.search(r'(?<![\d.,])%s(?![\d])' % re.escape(s), label):
            return label
    return '%s %s' % (label, value)


def _lin(d0, d1, r0, r1):
    span = (d1 - d0) or 1.0
    return lambda v: r0 + (v - d0) / span * (r1 - r0)


def _axis_cat(s):
    """날짜 항목을 축에 쓸 짧은 꼴 'M/D' 로. 날짜가 아니면 그대로."""
    m = _DAY.match(s)
    if not m:
        return s
    mo, d = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
    return '%d/%d' % (int(mo), int(d))


def _err_bounds(errs):
    return [b for e in (errs or []) if e for b in e]


# 그리기 조각 ─────────────────────

def _n(x):
    s = '%.1f' % x
    s = s[:-2] if s.endswith('.0') else s
    return '0' if s == '-0' else s


def _t(x, y, s, size=LS, fill=MUTED, anchor='start', bold=False, halo=False):
    return '<text x="%s" y="%s" font-size="%s" fill="%s"%s%s%s>%s</text>' % (
        _n(x), _n(y), _n(size), fill,
        ' text-anchor="%s"' % anchor if anchor != 'start' else '', ' font-weight="700"' if bold else '',
        ' stroke="#fff" stroke-width="3" stroke-linejoin="round" paint-order="stroke"' if halo else '',
        html.escape(s))


def _tbox(x, y, s, size, anchor='start'):
    """글자가 차지하는 상자(x0, y0, x1, y1). y 는 글자 기준선."""
    w = _text_w(s, size)
    x0 = x - w if anchor == 'end' else x - w / 2 if anchor == 'middle' else x
    return (x0, y - size * 0.82, x0 + w, y + size * 0.24)


def _hits(boxes, b, pad=2.0):
    return sum(1 for c in boxes if b[0] < c[2] + pad and c[0] < b[2] + pad and b[1] < c[3] + pad and c[1] < b[3] + pad)


def _line(x1, y1, x2, y2, color, width=1):
    return '<line x1="%s" y1="%s" x2="%s" y2="%s" stroke="%s" stroke-width="%s"/>' % (
        _n(x1), _n(y1), _n(x2), _n(y2), color, width)


def _hbar(x0, x1, y, h, fill, r=4.0):
    """가로 막대. x0 = 기준선 쪽(각짐), x1 = 데이터 쪽 끝(4px 둥글게)."""
    w = abs(x1 - x0)
    if w < 0.5:
        return ''
    r = min(r, h / 2.0, w)
    if r <= 0:
        return '<rect x="%s" y="%s" width="%s" height="%s" fill="%s"/>' % (_n(min(x0, x1)), _n(y), _n(w), _n(h), fill)
    sweep, sign = (1, -1) if x1 >= x0 else (0, 1)
    return ('<path d="M%s,%sH%sA%s,%s 0 0 %d %s,%sV%sA%s,%s 0 0 %d %s,%sH%sZ" fill="%s"/>'
            % (_n(x0), _n(y), _n(x1 + sign * r), _n(r), _n(r), sweep, _n(x1), _n(y + r), _n(y + h - r), _n(r), _n(r),
               sweep, _n(x1 + sign * r), _n(y + h), _n(x0), fill))


def _vbar(x, base, top, w, fill, r=4.0):
    """세로 막대. base = 기준선 쪽(각짐), top = 데이터 쪽 끝(4px 둥글게)."""
    h = abs(base - top)
    if h < 0.5:
        return ''
    r = min(r, w / 2.0, h)
    sweep, sign = (1, 1) if top <= base else (0, -1)
    return ('<path d="M%s,%sV%sA%s,%s 0 0 %d %s,%sH%sA%s,%s 0 0 %d %s,%sV%sZ" fill="%s"/>'
            % (_n(x), _n(base), _n(top + sign * r), _n(r), _n(r), sweep, _n(x + r), _n(top), _n(x + w - r), _n(r), _n(r),
               sweep, _n(x + w), _n(top + sign * r), _n(base), fill))


def _whisker(a, b, at, horizontal=True):
    """구간 수염(얇은 선 + 양 끝 짧은 막대)."""
    if horizontal:
        d = 'M%s,%sH%sM%s,%sV%sM%s,%sV%s' % (_n(a), _n(at), _n(b), _n(a), _n(at - 4), _n(at + 4), _n(b), _n(at - 4),
                                             _n(at + 4))
    else:
        d = 'M%s,%sV%sM%s,%sH%sM%s,%sH%s' % (_n(at), _n(a), _n(b), _n(at - 4), _n(a), _n(at + 4), _n(at - 4), _n(b),
                                             _n(at + 4))
    return '<path d="%s" stroke="%s" stroke-width="1.5" fill="none"/>' % (d, INK)


def _tip(text, inner, hit=None):
    """마우스를 올리면 '항목 · 계열 · 값' 이 뜨게 묶는다. hit = 눈에 안 보이는 넓은 과녁(x, y, w, h)."""
    target = ('<rect x="%s" y="%s" width="%s" height="%s" fill="#fff" fill-opacity="0"/>'
              % tuple(_n(v) for v in hit)) if hit else ''
    return '<g><title>%s</title>%s%s</g>' % (html.escape(text), target, inner)


def _wrap(s, size, limit, max_lines=2):
    """띄어쓰기 단위로 줄을 나눈다. 넘치는 마지막 줄은 … 로 줄인다."""
    words = str(s or '').split()
    lines, cur = [], ''
    for word in words:
        cand = cur + ' ' + word if cur else word
        if _text_w(cand, size) <= limit:
            cur = cand
            continue
        if cur:
            lines.append(cur)
        while _text_w(word, size) > limit and len(word) > 1:
            cut = len(word) - 1
            while cut > 1 and _text_w(word[:cut], size) > limit:
                cut -= 1
            lines.append(word[:cut])
            word = word[cut:]
        cur = word
    if cur:
        lines.append(cur)
    if len(lines) > max_lines:
        last = lines[max_lines - 1]
        while last and _text_w(last + '…', size) > limit:
            last = last[:-1]
        lines = lines[:max_lines - 1] + [last.rstrip() + '…']
    return lines


def _legend(series, W, y, kind='bar'):
    """계열 2개 이상일 때 늘 두는 범례. (조각, 높이)."""
    parts, x, row = [], SIDE, 0
    for j, s in enumerate(series):
        name = _fit(s['name'], 12, W - 2 * SIDE - 22)
        w = 18 + _text_w(name, 12)
        if x > SIDE and x + w > W - SIDE:
            x, row = SIDE, row + 1
        yy = y + row * 18
        if kind == 'line':
            parts.append('<line x1="%s" y1="%s" x2="%s" y2="%s" stroke="%s" stroke-width="2" stroke-linecap="round"/>'
                         % (_n(x), _n(yy + 8), _n(x + 12), _n(yy + 8), SERIES_COLORS[j]))
        else:
            parts.append('<rect x="%s" y="%s" width="10" height="10" rx="2" fill="%s"/>'
                         % (_n(x + 1), _n(yy + 3), SERIES_COLORS[j]))
        parts.append(_t(x + 18, yy + 12, name, 12, MUTED))
        x += w + 16
    return ''.join(parts), (row + 1) * 18


def _cat_layout(cats, W):
    """가로 막대의 항목 이름: 짧으면 막대 왼쪽('side'), 길면 막대 위 한 줄('above')."""
    widest = max(_text_w(c, LS) for c in cats)
    if widest <= W * 0.36:
        return 'side', widest + 10
    return 'above', 0.0


def _ref_band(refs, sx, unit, W, y):
    """가로 막대의 기준선 이름표를 그림 위쪽 띠에 겹치지 않게 놓는다. (조각, 띠 높이, 선 x 목록)."""
    rows, items = [], []
    for r in refs:
        x = sx(r['value'])
        s = _fit(_ref_text(r, unit), TS, W - 2 * SIDE)
        w = _text_w(s, TS)
        x0 = min(max(x - w / 2, SIDE), W - SIDE - w)
        for ri, row in enumerate(rows):
            if all(x0 > b + 8 or x0 + w < a - 8 for a, b in row):
                row.append((x0, x0 + w))
                break
        else:
            rows.append([(x0, x0 + w)])
            ri = len(rows) - 1
        items.append((ri, x0, s, x))
    parts = [_t(x0, y + ri * 15 + 11, s, TS, MUTED) for ri, x0, s, _ in items]
    return ''.join(parts), len(rows) * 15, [x for _, _, _, x in items]


def _hticks(ticks, sx, unit, step, y, W):
    """가로 축 눈금 글(겹치면 앞의 것을 뺀다 — 끝 눈금은 남긴다)."""
    placed = []
    for t in ticks:
        s = fmt_tick(t, unit, step)
        w = _text_w(s, TS)
        x0 = min(max(sx(t) - w / 2, SIDE), W - SIDE - w)
        while placed and x0 < placed[-1][0] + placed[-1][2] + 8:
            placed.pop()
        placed.append((x0, s, w))
    return ''.join(_t(x0, y, s, TS, FAINT) for x0, s, _ in placed)


def _vgrid(ticks, sy, unit, step, x_left, x_right):
    """세로 축 격자(가로 실선)와 눈금 글."""
    parts = []
    for t in ticks:
        yy = sy(t)
        parts.append(_line(x_left, yy, x_right, yy, GRID))
        parts.append(_t(x_left - 6, yy + 4, fmt_tick(t, unit, step), TS, FAINT, 'end'))
    return ''.join(parts)


def _xlabels(cats, xs, slot, y, W):
    """세로 그래프의 아래 항목 이름. 날짜면 칸이 좁을 때 간격을 두고 일부만(M/D, 첫날·마지막 날 포함),
    아니면 칸 폭에 맞춰 두 줄까지."""
    n = len(cats)
    if all(_DAY.match(c) for c in cats):
        texts = [_axis_cat(c) for c in cats]
        wmax = max(_text_w(s, TS) for s in texts)
        every = max(1, int(math.ceil((wmax + 8) / max(slot, 1e-6))))
        idx = list(range(0, n, every))
        if idx[-1] != n - 1:
            if (n - 1 - idx[-1]) * slot >= wmax + 8:
                idx.append(n - 1)
            else:
                idx[-1] = n - 1
        out = []
        for i in idx:
            w = _text_w(texts[i], TS)
            out.append(_t(min(max(xs[i] - w / 2, SIDE), W - SIDE - w), y, texts[i], TS, FAINT))
        return ''.join(out), 0
    items = []
    for i, c in enumerate(cats):
        lines = _wrap(c, TS, max(slot - 6, 12), 2)
        w = max(_text_w(ln, TS) for ln in lines)
        x0 = min(max(xs[i] - w / 2, SIDE), W - SIDE - w)
        items.append((x0, x0 + w, i, lines))
    kept = []                                        # 그림 끝에서 안으로 밀려 겹치면 앞의 것을 뺀다(마지막은 남긴다)
    for it in items:
        while kept and it[0] < kept[-1][1] + 6:
            if it[2] == n - 1:
                kept.pop()
            else:
                break
        if not kept or it[0] >= kept[-1][1] + 6:
            kept.append(it)
    out, extra = [], 0
    for _, _, i, lines in kept:
        for li, ln in enumerate(lines):
            w = _text_w(ln, TS)
            out.append(_t(min(max(xs[i] - w / 2, SIDE), W - SIDE - w), y + li * 13, ln, TS, FAINT))
        extra = max(extra, (len(lines) - 1) * 13)
    return ''.join(out), extra


def _ink_on(fill):
    """색 막대 안에 쓰는 글자색: 흰색과 진한 글자 중 대비가 큰 쪽."""
    def lum(h):
        def ch(c):
            return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
        r, g, b = (int(h[i:i + 2], 16) / 255.0 for i in (1, 3, 5))
        return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)
    L = lum(fill)
    return '#ffffff' if 1.05 / (L + 0.05) >= (L + 0.05) / (lum(INK) + 0.05) else INK


def _place_refs(refs, sy, unit, x_left, x_right, y_min, y_max, marks, band=()):
    """세로 그래프의 가로 기준선과 이름표. 이름표는 선 위·아래, 오른쪽·왼쪽 중 막대·점과 겹치지 않는 자리에 둔다.

    (조각, 이름표 상자 목록, 빈자리가 없어 그림 위 줄로 보낼 기준선 번호 목록). band 에 든 번호는 선만 긋는다."""
    parts, boxes, need = [], [], []
    for ri, r in enumerate(refs):
        yv = sy(r['value'])
        parts.append(_line(x_left, yv, x_right, yv, REF))
        if ri in band:
            continue
        s = _fit(_ref_text(r, unit), TS, x_right - x_left - 8)
        best = None
        for x, yy, anchor in ((x_right - 2, yv - 4, 'end'), (x_left + 4, yv - 4, 'start'),
                              (x_right - 2, yv + 12, 'end'), (x_left + 4, yv + 12, 'start')):
            b = _tbox(x, yy, s, TS, anchor)
            out = 1 if b[1] < y_min or b[3] > y_max else 0
            score = out + _hits(boxes, b) + _hits(marks, b, 0)
            if best is None or score < best[0]:
                best = (score, x, yy, anchor, b)
        score, x, yy, anchor, b = best
        if score:
            need.append(ri)
            continue
        parts.append(_t(x, yy, s, TS, MUTED, anchor, halo=True))
        boxes.append(b)
    return ''.join(parts), boxes, need


def _ref_keys(refs, unit, W, x_start, y):
    """그림 안에 빈자리가 없는 기준선 이름표를 그림 위 줄에 '— 이름 값' 꼴로 모은다. (조각, 높이)."""
    parts, x, row = [], x_start, 0
    for r in refs:
        s = _fit(_ref_text(r, unit), TS, W - x_start - SIDE - 22)
        w = 20 + _text_w(s, TS)
        if x > x_start and x + w > W - SIDE:
            x, row = x_start, row + 1
        yy = y + row * 16
        parts.append(_line(x, yy + 7, x + 14, yy + 7, REF, 1.5))
        parts.append(_t(x + 20, yy + 11, s, TS, MUTED))
        x += w + 14
    return ''.join(parts), (row + 1) * 16


REF_GAP = 10           # 기준선과 이만큼(px) 안쪽으로 가까운 격자선은 그리지 않는다(겹선 방지).
#                        아래 눈금 글자는 기준선이 끝난 아래에 있어 겹치지 않으므로 모두 남긴다(축 눈금이 빠져 보이지 않게).


def _off_refs(ticks, sx, ref_xs, keep_zero=True):
    """가로 막대의 격자선 중 기준선 가까이 있는 것을 뺀다(0 은 남긴다). 눈금 글자에는 쓰지 않는다."""
    return [t for t in ticks if (keep_zero and t == 0) or all(abs(sx(t) - rx) >= REF_GAP for rx in ref_xs)]


LABEL_REF_PAD = 4      # 값 글자 상자에서 이만큼(px) 안쪽까지 기준선이 오면 '글자를 가른다'고 본다


def _crosses(ref_xs, a, b, pad=LABEL_REF_PAD):
    """글자 상자 [a, b] 를 기준선이 지나가는지."""
    return any(a - pad <= rx <= b + pad for rx in ref_xs)


# 모양별 그림 ─────────────────────

def _body_bar(ch, W):
    cats, ser, unit = ch['categories'], ch['series'], ch['unit']
    k, hl, errs, refs = len(ser), ch['highlight'], ch['errors'], ch['reference']
    back, marks, vlabels, front = [], [], [], []
    y = 4.0
    if k > 1:
        leg, lh = _legend(ser, W, y)
        front.append(leg)
        y += lh + 8
    vals = [v for s in ser for v in s['values'] if v is not None]
    lo, hi, step, ticks = _domain(vals + [r['value'] for r in refs] + _err_bounds(errs), ch['max'])
    value_w = max(_text_w(t, LS) for t in [fmt_number(v, unit) for v in vals] + [NA]) + 10
    mode, label_w = _cat_layout(cats, W)
    x_left, x_right = SIDE + label_w, W - SIDE - value_w
    sx = _lin(lo, hi, x_left, x_right)
    x0 = sx(0)
    ref_xs = []
    if refs:
        band, band_h, ref_xs = _ref_band(refs, sx, unit, W, y)
        front.append(band)
        y += band_h + 4
    plot_top = y
    bar_h = (18, 14, 12)[k - 1]
    group_h = k * bar_h + (k - 1) * 2
    spans = []
    for i, cat in enumerate(cats):
        if mode == 'above':
            lines = _wrap(cat, LS, W - 2 * SIDE, 2)
            for li, ln in enumerate(lines):
                front.append(_t(SIDE, y + 12 + li * 16, ln, LS, MUTED, halo=True))
            y += 16 * len(lines) + 4
        else:
            front.append(_t(x_left - 8, y + group_h / 2.0 + 4.5, _fit(cat, LS, label_w - 10), LS, MUTED, 'end'))
        spans.append((y - 3, y + group_h + 3))
        for j, s in enumerate(ser):
            v = s['values'][i]
            by = y + j * (bar_h + 2)
            mid = by + bar_h / 2.0
            strong = k == 1 and hl is not None and i in hl
            color = SERIES_COLORS[j] if k > 1 else (ONE if hl is None or strong else OFF)
            label = _chart_value(v, unit)
            tip = '%s · %s · %s' % (cat, s['name'], label)
            inner, texts = [], []
            if v is None:
                inner.append(_t(x_left + 4, mid + 4.5, NA, 12, FAINT))
            else:
                end = sx(v)
                inner.append(_hbar(x0, end, by, bar_h, color))
                e = errs[i] if errs and j == 0 else None
                if e:
                    inner.append(_whisker(sx(e[0]), sx(e[1]), mid))
                    end = max(end, sx(e[1]))
                    tip += ' (구간 %s ~ %s)' % (_chart_value(e[0], unit), _chart_value(e[1], unit))
                # 값 글자는 늘 막대(구간) 끝에 붙인다 — 기준선을 넘겨 놓으면 기준을 넘은 값처럼 읽힌다.
                # 막대 밖 글자를 기준선이 가르면: 막대 안쪽 끝에 넣고(밖에는 아무것도 두지 않음),
                # 막대가 짧아 안에 못 넣으면 글자 뒤에 흰 바탕을 깔아 그 높이만큼 선을 끊는다.
                lw = _text_w(label, LS)
                lx = min(max(end, x0) + 6, W - SIDE - lw)
                ink = INK if (k > 1 or hl is None or strong) else MUTED
                in_right = (min(sx(v), sx(e[0])) if e else sx(v)) - 6       # 막대 안 글자의 오른쪽 끝(수염 앞)
                after = lx + lw                                               # 계열 이름을 붙일 자리
                if not _crosses(ref_xs, lx, lx + lw):
                    texts.append(_t(lx, mid + 4.5, label, LS, ink, bold=strong, halo=True))
                elif sx(v) > x0 and in_right - lw >= x0 + 4 and not _crosses(ref_xs, in_right - lw, in_right):
                    texts.append(_t(in_right, mid + 4.5, label, LS, _ink_on(color), 'end', bold=strong))
                    after = max(end, x0)
                else:
                    texts.append('<rect x="%s" y="%s" width="%s" height="14" fill="#ffffff"/>'
                                 % (_n(lx - 2), _n(mid - 6.5), _n(lw + 4)))
                    texts.append(_t(lx, mid + 4.5, label, LS, ink, bold=strong, halo=True))
                if k > 1 and i == 0:            # 첫 묶음에는 계열 이름도 막대 옆에 직접 단다
                    nx = after + 6
                    nw = _text_w(s['name'], 11.5)
                    if nx + nw <= W - SIDE and not _crosses(ref_xs, nx, nx + nw, 3):
                        texts.append(_t(nx, mid + 4.5, s['name'], 11.5, FAINT, halo=True))
            marks.append(_tip(tip, ''.join(inner), (SIDE, by - 1, W - 2 * SIDE, bar_h + 2)))
            if texts:
                vlabels.append(_tip(tip, ''.join(texts)))
        y += group_h + 12
    plot_bottom = y - 12
    grid = _off_refs(ticks, sx, ref_xs)
    back.append(_grid_spans(grid, sx, x0, spans if mode == 'above' else [(plot_top - 2, plot_bottom + 2)]))
    lines = ''.join(_line(rx, plot_top - 4, rx, plot_bottom + 2, REF) for rx in ref_xs)
    front.append(_hticks(ticks, sx, unit, step, plot_bottom + 17, W))
    return ''.join(back + marks + [lines] + vlabels + front), plot_bottom + 23


def _grid_spans(ticks, sx, x0, spans):
    """가로 막대의 세로 격자와 기준축. 막대 위에 이름을 쓰는 방식이면 막대 묶음 구간에만 긋는다(글자를 지나지 않게)."""
    parts = []
    for a, b in spans:
        parts.extend(_line(sx(t), a, sx(t), b, GRID) for t in ticks)
        parts.append(_line(x0, a, x0, b, AXIS))
    return ''.join(parts)


def _body_stacked(ch, W):
    cats, ser, unit, refs = ch['categories'], ch['series'], ch['unit'], ch['reference']
    k = len(ser)
    back, marks, front = [], [], []
    y = 4.0
    if k > 1:
        leg, lh = _legend(ser, W, y)
        front.append(leg)
        y += lh + 8
    totals = [sum(s['values'][i] for s in ser if s['values'][i] is not None) for i in range(len(cats))]
    lo, hi, step, ticks = _domain(totals + [r['value'] for r in refs], ch['max'])
    mode, label_w = _cat_layout(cats, W)
    x_left, x_right = SIDE + label_w, W - SIDE - 10
    sx = _lin(lo, hi, x_left, x_right)
    ref_xs = []
    if refs:
        band, band_h, ref_xs = _ref_band(refs, sx, unit, W, y)
        front.append(band)
        y += band_h + 4
    plot_top = y
    bar_h = 20
    spans = []
    for i, cat in enumerate(cats):
        if mode == 'above':
            lines = _wrap(cat, LS, W - 2 * SIDE, 2)
            for li, ln in enumerate(lines):
                front.append(_t(SIDE, y + 12 + li * 16, ln, LS, MUTED, halo=True))
            y += 16 * len(lines) + 4
        else:
            front.append(_t(x_left - 8, y + bar_h / 2.0 + 4.5, _fit(cat, LS, label_w - 10), LS, MUTED, 'end'))
        spans.append((y - 3, y + bar_h + 3))
        segs = [(j, s['values'][i]) for j, s in enumerate(ser) if s['values'][i] is not None and s['values'][i] > 0]
        if all(s['values'][i] is None for s in ser):
            marks.append(_t(x_left + 4, y + bar_h / 2.0 + 4.5, NA, 12, FAINT))
        cum, outside = 0.0, []
        for idx, (j, v) in enumerate(segs):
            last = idx == len(segs) - 1
            a, b = sx(cum), sx(cum + v)
            b_draw = b if last else b - 2            # 조각 사이 2px 흰 틈
            color = SERIES_COLORS[j] if k > 1 else ONE
            inner = [_hbar(a, b_draw, y, bar_h, color, 4.0 if last else 0.0)] if b_draw - a >= 0.5 else []
            text = fmt_number(v, unit)
            tip = '%s · %s · %s' % (cat, ser[j]['name'], text)
            if _text_w(text, 11.5) + 8 <= b_draw - a:
                inner.append(_t((a + b_draw) / 2.0, y + bar_h / 2.0 + 4, text, 11.5, _ink_on(color), 'middle'))
            else:                                    # 좁은 조각: 숫자가 그림에서 사라지지 않게 막대 밖에 '이름 값'으로
                outside.append(('%s %s' % (ser[j]['name'], text)) if k > 1 else text)
            marks.append(_tip(tip, ''.join(inner), (a, y - 2, max(b - a, 1.0), bar_h + 4)))
            cum += v
        extra = 0
        if outside:
            s_out = ' · '.join(outside)
            w_out = _text_w(s_out, 11.5)
            end = sx(cum)
            if end + 6 + w_out <= W - SIDE:          # 막대 오른쪽에 자리가 있으면 그 옆에
                marks.append(_t(end + 6, y + bar_h / 2.0 + 4, s_out, 11.5, INK, halo=True))
            else:                                    # 없으면 막대 바로 아래, 막대 끝에 맞춰
                s_out = _fit(s_out, 11.5, W - 2 * SIDE)
                w_out = _text_w(s_out, 11.5)
                marks.append(_t(max(min(end, W - SIDE), SIDE + w_out), y + bar_h + 13, s_out, 11.5, INK, 'end', halo=True))
                extra = 16
        y += bar_h + 12 + extra
    plot_bottom = y - 12
    grid = _off_refs(ticks, sx, ref_xs)
    back.append(_grid_spans(grid, sx, sx(0), spans if mode == 'above' else [(plot_top - 2, plot_bottom + 2)]))
    lines = ''.join(_line(rx, plot_top - 4, rx, plot_bottom + 2, REF) for rx in ref_xs)
    front.append(_hticks(ticks, sx, unit, step, plot_bottom + 17, W))
    return ''.join(back + marks + [lines] + front), plot_bottom + 23


def _body_column(ch, W, band=()):
    cats, ser, unit = ch['categories'], ch['series'], ch['unit']
    k, n, hl, errs, refs = len(ser), len(cats), ch['highlight'], ch['errors'], ch['reference']
    back, marks, front = [], [], []
    y = 4.0
    if k > 1:
        leg, lh = _legend(ser, W, y)
        front.append(leg)
        y += lh + 6
    vals = [v for s in ser for v in s['values'] if v is not None]
    lo, hi, step, ticks = _domain(vals + [r['value'] for r in refs] + _err_bounds(errs), ch['max'])
    tick_w = max(_text_w(fmt_tick(t, unit, step), TS) for t in ticks) + 8
    x_left, x_right = SIDE + tick_w, W - SIDE
    if band:
        keys, kh = _ref_keys([refs[i] for i in sorted(band)], unit, W, x_left, y)
        front.append(keys)
        y += kh + 2
    plot_top = y + 18
    plot_bottom = plot_top + 150
    sy = _lin(lo, hi, plot_bottom, plot_top)
    base = sy(0)
    slot = (x_right - x_left) / float(n)
    bw = min(20.0, slot * 0.72) if k == 1 else min(14.0, (slot * 0.8 - 2 * (k - 1)) / k)
    bw = max(bw, 1.0)
    group_w = k * bw + 2 * (k - 1)
    centers = [x_left + slot * (i + 0.5) for i in range(n)]
    back.append(_vgrid(ticks, sy, unit, step, x_left, x_right))
    back.append(_line(x_left, base, x_right, base, AXIS))
    bar_boxes, label_cands = [], []
    for i, cat in enumerate(cats):
        gx = centers[i] - group_w / 2.0
        for j, s in enumerate(ser):
            v = s['values'][i]
            bx = gx + j * (bw + 2)
            cx = bx + bw / 2.0
            strong = k == 1 and hl is not None and i in hl
            color = SERIES_COLORS[j] if k > 1 else (ONE if hl is None or strong else OFF)
            label = _chart_value(v, unit)
            tip = '%s · %s · %s' % (cat, s['name'], label)
            inner = []
            hit = (x_left + slot * i, plot_top - 18, slot, plot_bottom - plot_top + 18) if k == 1 else \
                (bx - 1, plot_top - 18, bw + 2, plot_bottom - plot_top + 18)
            if v is None:
                # 빈값: 글자 대신 작은 빈 동그라미(촘촘해도 겹치지 않게). 값은 마우스 올림·숫자 표에 '측정 못 함'
                inner.append('<circle cx="%s" cy="%s" r="2.5" fill="#fff" stroke="%s" stroke-width="1"/>'
                             % (_n(cx), _n(base - 5 if lo >= 0 or base > plot_top + 10 else base + 5), FAINT))
            else:
                top = sy(v)
                inner.append(_vbar(bx, base, top, bw, color))
                bar_boxes.append((bx, min(top, base), bx + bw, max(top, base)))
                tip_end = top
                e = errs[i] if errs and j == 0 else None
                if e:
                    inner.append(_whisker(sy(e[0]), sy(e[1]), cx, horizontal=False))
                    tip_end = min(top, sy(e[1]))
                    tip += ' (구간 %s ~ %s)' % (_chart_value(e[0], unit), _chart_value(e[1], unit))
                ly = tip_end - 5 if v >= 0 else top + 14
                label_cands.append((i, j, v, cx, ly, label, strong, k > 1 or hl is None or strong))
            marks.append(_tip(tip, ''.join(inner), hit))
    ref_parts, ref_boxes, need = _place_refs(refs, sy, unit, x_left, x_right, plot_top - 16, plot_bottom,
                                             bar_boxes, band)
    if need:                                         # 막대에 가리는 이름표는 그림 위 줄로 보내고 다시 그린다
        return _body_column(ch, W, set(band) | set(need))
    front.append(ref_parts)
    # 값 이름표: 적으면 모두, 많으면 강조한 막대·가장 큰 값·마지막 값만
    fits = all(_text_w(c[5], LS) <= (slot if k == 1 else bw + 2) - 2 for c in label_cands)
    if fits and n <= (8 if k == 1 else 6):
        chosen = label_cands
    else:
        chosen = [c for c in label_cands if c[6]]
        for j in range(k):
            mine = [c for c in label_cands if c[1] == j]
            if mine:
                chosen.append(max(mine, key=lambda c: c[2]))
                chosen.append(mine[-1])
    placed = list(ref_boxes)
    for c in sorted(set(chosen), key=lambda c: (not c[6], c[0] != n - 1, c[0])):
        b = _tbox(c[3], c[4], c[5], LS, 'middle')
        shift = max(x_left - b[0], 0) + min(W - SIDE - b[2], 0)   # 눈금 글자 칸·그림 끝에 걸리면 안쪽으로 민다
        b = (b[0] + shift, b[1], b[2] + shift, b[3])
        if b[1] < 0 or b[3] > plot_bottom + 1 or _hits(placed, b) or _hits(bar_boxes, b, 0):
            continue
        placed.append(b)
        front.append(_t(c[3] + shift, c[4], c[5], LS, INK if c[7] else MUTED, 'middle', bold=c[6], halo=True))
    xl, extra = _xlabels(cats, centers, slot, plot_bottom + 17, W)
    front.append(xl)
    return ''.join(back + marks + front), plot_bottom + 23 + extra


def _body_line(ch, W, band=()):
    cats, ser, unit = ch['categories'], ch['series'], ch['unit']
    k, n, errs, refs = len(ser), len(cats), ch['errors'], ch['reference']
    back, marks, front = [], [], []
    y = 4.0
    if k > 1:
        leg, lh = _legend(ser, W, y, kind='line')
        front.append(leg)
        y += lh + 6
    vals = [v for s in ser for v in s['values'] if v is not None]
    zero = ch['max'] is not None or min(vals) <= 0 or min(vals) < 0.5 * max(vals)
    lo, hi, step, ticks = _domain(vals + [r['value'] for r in refs] + _err_bounds(errs), ch['max'], zero=zero)
    tick_w = max(_text_w(fmt_tick(t, unit, step), TS) for t in ticks) + 8
    lasts = []
    for j, s in enumerate(ser):
        idx = [i for i, v in enumerate(s['values']) if v is not None]
        if idx:
            lasts.append((j, idx[-1], s['values'][idx[-1]]))
    end_w = max(_text_w(fmt_number(v, unit), LS) for _, _, v in lasts) + 14
    x_left, x_right = SIDE + tick_w, W - SIDE - end_w
    if band:
        keys, kh = _ref_keys([refs[i] for i in sorted(band)], unit, W, x_left, y)
        front.append(keys)
        y += kh + 2
    plot_top = y + 14
    plot_bottom = plot_top + 150
    sy = _lin(lo, hi, plot_bottom, plot_top)
    if n == 1:
        xs = [(x_left + x_right) / 2.0]
    else:
        xs = [x_left + 6 + i * (x_right - x_left - 12) / float(n - 1) for i in range(n)]
    slot = (x_right - x_left) / float(max(n - 1, 1))
    back.append(_vgrid(ticks, sy, unit, step, x_left, x_right))
    back.append(_line(x_left, plot_bottom, x_right, plot_bottom, AXIS))
    occupied = []
    for j, s in enumerate(ser):
        color = SERIES_COLORS[j] if k > 1 else ONE
        runs, cur = [], []
        for i, v in enumerate(s['values']):
            if v is None:
                if cur:
                    runs.append(cur)
                cur = []
            else:
                cur.append(i)
        if cur:
            runs.append(cur)
        for run in runs:
            if len(run) > 1:
                d = 'M' + ' L'.join('%s,%s' % (_n(xs[i]), _n(sy(s['values'][i]))) for i in run)
                marks.append('<path d="%s" fill="none" stroke="%s" stroke-width="2" stroke-linejoin="round" '
                             'stroke-linecap="round"/>' % (d, color))
            pts = [(xs[i], sy(s['values'][i])) for i in run]
            for a, b in zip(pts, pts[1:] + pts[-1:]):        # 점과 그 사이 선이 지나는 자리
                for f in (0.0, 0.25, 0.5, 0.75):
                    px, py = a[0] + (b[0] - a[0]) * f, a[1] + (b[1] - a[1]) * f
                    occupied.append((px - 3, py - 3, px + 3, py + 3))
        dots = sorted({i for run in runs for i in (run if n <= 12 else (run[0], run[-1]))})
        for i in dots:
            marks.append('<circle cx="%s" cy="%s" r="5" fill="%s" stroke="#fff" stroke-width="2"/>'
                         % (_n(xs[i]), _n(sy(s['values'][i])), color))
    if errs:
        for i, e in enumerate(errs):
            if e:
                marks.append(_whisker(sy(e[0]), sy(e[1]), xs[i], horizontal=False))
    ref_parts, ref_boxes, need = _place_refs(refs, sy, unit, x_left, x_right, plot_top - 14, plot_bottom,
                                             occupied, band)
    if need:
        return _body_line(ch, W, set(band) | set(need))
    front.append(ref_parts)
    placed = list(ref_boxes)
    for j, i, v in lasts:                            # 끝점에 값
        s = fmt_number(v, unit)
        b = _tbox(xs[i] + 9, sy(v) + 4.5, s, LS)
        if not _hits(placed, b):
            placed.append(b)
            front.append(_t(xs[i] + 9, sy(v) + 4.5, s, LS, INK, bold=k == 1, halo=True))
    xl, extra = _xlabels(cats, xs, slot, plot_bottom + 17, W)
    front.append(xl)
    for i, cat in enumerate(cats):                   # 날짜마다 세로 띠 과녁: 그 날짜의 모든 계열 값
        a = (xs[i - 1] + xs[i]) / 2.0 if i else x_left
        b = (xs[i] + xs[i + 1]) / 2.0 if i < n - 1 else x_right
        lines = ['%s · %s · %s' % (cat, s['name'], _chart_value(s['values'][i], unit)) for s in ser]
        if errs and errs[i]:
            lines[0] += ' (구간 %s ~ %s)' % (_chart_value(errs[i][0], unit), _chart_value(errs[i][1], unit))
        front.append(_tip('\n'.join(lines), '', (a, plot_top - 14, max(b - a, 1.0), plot_bottom - plot_top + 14)))
    return ''.join(back + marks + front), plot_bottom + 23 + extra


PAIR_UNIT_WORD = {'ms': 'ms', 'sec': '초', 'min': '분', 'mb': 'MB', 'usd': '달러', 'ratio': '%'}   # 칸 제목 끝 '(ms)'


def _pair_axes(cats, ser, units, errs):
    """작은 그래프 칸들의 축. 같은 단위 칸끼리는 축 하나(0 ~ 그 단위의 최댓값)와 값 글자 칸 폭을 함께 쓴다.

    눈금 글자를 그리지 않으므로 칸마다 따로 '보기 좋은 끝'으로 올리면 같은 ms 인데 작은 값의 막대가 더 길어진다.
    {단위: (아래 끝, 위 끝, 값 글자 칸 폭)}"""
    axes = {}
    for i in range(len(cats)):
        u = units[i]
        e = errs[i] if errs else None
        vals = [s['values'][i] for s in ser if s['values'][i] is not None] + (list(e) if e else [])
        vw = max(_text_w(_chart_value(s['values'][i], u), LS) for s in ser) + 10
        lo, hi, w0 = axes.get(u, (0.0, 0.0, 0.0))
        axes[u] = (min([lo] + vals), max([hi] + vals), max(w0, vw))
    return {u: (lo, hi if hi > lo else lo + 1.0, vw) for u, (lo, hi, vw) in axes.items()}


def _body_pairs(ch, W):
    cats, ser, units, errs = ch['categories'], ch['series'], ch['units'], ch['errors']
    k, n = len(ser), len(cats)
    axes = _pair_axes(cats, ser, units, errs)
    parts = []
    y = 4.0
    if k > 1:
        leg, lh = _legend(ser, W, y)
        parts.append(leg)
        y += lh + 10
    name_w = max(_text_w(s['name'], 11.5) for s in ser) + 8
    texts = [_chart_value(s['values'][i], units[i]) for i in range(n) for s in ser]
    value_w = max(_text_w(t, LS) for t in texts) + 10
    gap = 20
    cols = 2 if n > 1 and (W - 2 * SIDE - gap) / 2.0 >= min(name_w, 80) + value_w + 60 else 1
    pw = (W - 2 * SIDE - gap * (cols - 1)) / float(cols)
    name_w = min(name_w, pw * 0.4)
    bar_h, pitch = 12, 16
    panel_h = 20 + k * pitch
    for i, cat in enumerate(cats):
        col, row = i % cols, i // cols
        px, py = SIDE + col * (pw + gap), y + row * (panel_h + 14)
        u = units[i]
        e = errs[i] if errs else None
        lo, hi, vw = axes[u]
        bx0, bx1 = px + name_w, px + pw - vw
        sx = _lin(lo, hi, bx0, bx1)
        word = PAIR_UNIT_WORD.get(u)
        head = '%s (%s)' % (cat, word) if word and not cat.endswith('(%s)' % word) else cat
        parts.append(_t(px, py + 13, _fit(head, LS, pw), LS, INK, bold=True))
        parts.append(_line(sx(0), py + 19, sx(0), py + 19 + k * pitch - 1, AXIS))
        for j, s in enumerate(ser):
            v = s['values'][i]
            by = py + 20 + j * pitch
            mid = by + bar_h / 2.0
            label = _chart_value(v, u)
            parts.append(_t(bx0 - 6, mid + 4, _fit(s['name'], 11.5, name_w - 6), 11.5, MUTED, 'end'))
            tip = '%s · %s · %s' % (cat, s['name'], label)
            inner = []
            if v is None:
                inner.append(_t(sx(0) + 4, mid + 4, NA, 11.5, FAINT))
            else:
                end = sx(v)
                inner.append(_hbar(sx(0), end, by, bar_h, SERIES_COLORS[j] if k > 1 else ONE))
                if e and j == 0:
                    inner.append(_whisker(sx(e[0]), sx(e[1]), mid))
                    end = max(end, sx(e[1]))
                    tip += ' (구간 %s ~ %s)' % (_chart_value(e[0], u), _chart_value(e[1], u))
                inner.append(_t(max(end, sx(0)) + 6, mid + 4.5, label, LS, INK, halo=True))
            parts.append(_tip(tip, ''.join(inner), (px, by - 2, pw, bar_h + 4)))
    rows = int(math.ceil(n / float(cols)))
    return ''.join(parts), y + rows * (panel_h + 14) - 8


_BODIES = {'bar': _body_bar, 'stacked': _body_stacked, 'column': _body_column, 'line': _body_line,
           'pairs': _body_pairs}


def figure_svg(ch, W=FIG_W):
    """영역 그래프 하나의 SVG. (svg, 폭, 높이)."""
    body, h = _BODIES[ch['form']](ch, W)
    h = int(math.ceil(h))
    svg = ('<svg class="fig-svg" role="img" aria-label="%s" width="%d" height="%d" viewBox="0 0 %d %d" '
           'font-family="%s"><rect width="%d" height="%d" fill="#ffffff"/>%s</svg>'
           % (html.escape(ch['title'], quote=True), W, h, W, h, SVG_FONT, W, h, body))
    return svg, W, h


def card_svg(title, subtitle, svg, w, h, note, foot):
    """발표 자료(PPT)에 그대로 붙일 그림: 제목·부제·그래프·쉬운 설명·출처를 한 장에. (svg, 폭, 높이)."""
    W, pad = CARD_W, CARD_PAD
    inner = W - 2 * pad
    scale = inner / float(w)
    parts, y = [], float(pad)
    for ln in _wrap(title, 25, inner, 2):
        parts.append(_t(pad, y + 24, ln, 25, INK, bold=True))
        y += 34
    for ln in _wrap(subtitle, 16, inner, 4):
        parts.append(_t(pad, y + 17, ln, 16, MUTED))
        y += 23
    y += 14
    parts.append('<g transform="translate(%s,%s) scale(%s)">%s</g>' % (_n(pad), _n(y), '%.4f' % scale, svg))
    y += h * scale + 16
    for ln in _wrap(note, 17, inner, 7):
        parts.append(_t(pad, y + 18, ln, 17, INK))
        y += 26
    y += 4
    for ln in _wrap(foot, 13, inner, 3):
        parts.append(_t(pad, y + 14, ln, 13, FAINT))
        y += 19
    H = int(math.ceil(y + pad - 8))
    return ('<svg width="%d" height="%d" viewBox="0 0 %d %d" font-family="%s"><rect width="%d" height="%d" '
            'fill="#ffffff"/>%s</svg>' % (W, H, W, H, SVG_FONT, W, H, ''.join(parts))), W, H


# ── HTML ────────────────────────────────────────────────

CSS = '''
:root{--ink:%(ink)s;--muted:%(muted)s;--faint:%(faint)s;--line:%(line)s;--soft:%(soft)s;--accent:%(accent)s;--accent-soft:%(accent_soft)s}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%%;text-size-adjust:100%%}
body{margin:0;background:#fff;color:var(--ink);font-family:%(font)s;font-size:15px;line-height:1.6;word-break:keep-all;overflow-wrap:break-word}
main{max-width:1080px;margin:0 auto;padding:40px 32px 64px}
header{border-bottom:2px solid var(--accent);padding-bottom:16px}
.eyebrow{font-size:13px;color:var(--muted)}
h1{font-size:28px;line-height:1.3;margin:4px 0 6px;font-weight:700}
.runat{margin:0;font-size:14px;color:var(--muted)}
ul.cond{margin:6px 0 0;padding-left:18px;font-size:14px;color:var(--muted)}
h2{font-size:20px;line-height:1.35;margin:44px 0 14px;padding-bottom:6px;border-bottom:1px solid var(--line);break-after:avoid}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));grid-auto-flow:row dense;gap:16px}
.card{border:1px solid var(--line);border-radius:6px;padding:18px 20px;min-width:0;break-inside:avoid}
.card.wide{grid-column:1/-1}
@media screen and (min-width:760px){
  .card.wide{display:grid;grid-template-columns:minmax(220px,1fr) minmax(0,380px);column-gap:40px;align-items:start}
  .card.wide .chart{margin-top:0}
}
.card h3{font-size:15px;margin:0 0 10px;font-weight:600;color:var(--ink)}
.vals{display:flex;flex-wrap:wrap;gap:14px 36px}
.val{min-width:0}
.big{font-size:32px;line-height:1.15;font-weight:700;color:var(--ink);font-variant-numeric:tabular-nums}
.big.na{font-size:20px;line-height:1.4;color:var(--muted)}
.vlabel{font-size:13px;color:var(--muted);margin-top:2px}
.vreason{font-size:12.5px;color:var(--faint);margin-top:2px}
.chart{margin-top:16px}
.chart svg{display:block;width:100%%;max-width:380px;height:auto}
.note{font-size:13px;color:var(--faint);margin:12px 0 0}
.lead{margin:0 0 14px;padding:10px 14px;border-left:3px solid var(--accent);background:var(--accent-soft);font-weight:600}
.tbl{overflow-x:auto}
table{width:100%%;border-collapse:collapse;table-layout:fixed;font-size:14px}
th,td{text-align:left;vertical-align:top;padding:10px;border-bottom:1px solid var(--line)}
thead th{background:var(--soft);font-size:12.5px;font-weight:600;color:var(--muted)}
th.c-name{width:21%%} th.c-val{width:13%%;text-align:right} th.c-tgt{width:9%%} th.c-plain{width:30%%} th.c-src{width:15%%} th.c-when{width:12%%}
td.name{font-weight:600}
td.val{text-align:right;font-weight:700;font-variant-numeric:tabular-nums}
td.val.txt{text-align:left;font-weight:600}
td.val.na{text-align:left;font-weight:600;color:var(--muted)}
td.val .why{display:block;font-weight:400;font-size:12.5px;color:var(--faint);margin-top:2px}
td.tgt{font-size:13px;color:var(--muted)}
.det{display:block;font-size:12.5px;color:var(--muted);margin-top:4px}
td.src{font-size:12.5px;color:var(--faint);overflow-wrap:anywhere}
td.when{font-size:12.5px;color:var(--muted)}
td.when span{display:block}
pre{margin:8px 0 12px;padding:12px 14px;background:var(--soft);border:1px solid var(--line);border-radius:4px;white-space:pre-wrap;overflow-wrap:anywhere}
code{font-family:Consolas,"D2Coding",monospace;font-size:13.5px}
ul.plainlist{margin:6px 0 0;padding-left:20px} ul.plainlist li{margin-bottom:4px}
.tag{font-size:12px;color:var(--muted);border:1px solid var(--line);border-radius:3px;padding:0 6px;margin-left:6px;white-space:nowrap}
footer{margin-top:48px;padding-top:12px;border-top:1px solid var(--line);font-size:12.5px;color:var(--faint)}
.figs{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:16px;margin:0 0 20px}
.fig{margin:0;border:1px solid var(--line);border-radius:6px;padding:16px 18px 12px;min-width:0;break-inside:avoid}
.fig h3{font-size:15px;line-height:1.4;margin:0;font-weight:600}
.fig .sub{font-size:12.5px;line-height:1.5;color:var(--muted);margin:2px 0 10px}
.fig .plot svg{display:block;width:100%%;max-width:420px;height:auto}
.fig .fnote{font-size:13.5px;margin:10px 0 0}
.fig .fsrc{font-size:12px;color:var(--faint);margin:4px 0 0;overflow-wrap:anywhere}
.fig details{margin-top:8px}
.fig summary{cursor:pointer;font-size:13px;color:var(--muted)}
.ntbl{overflow-x:auto;margin-top:6px}
table.nums{width:auto;min-width:60%%;table-layout:auto;font-size:12.5px}
.nums th,.nums td{padding:4px 10px 4px 0;border-bottom:1px solid var(--line)}
.nums thead th{background:none;font-size:12px;text-align:right}
.nums thead th:first-child,.nums tbody th{text-align:left}
.nums tbody th{font-weight:400}
.nums td{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.nref{font-size:12px;color:var(--muted);margin:6px 0 0}
@media screen and (max-width:880px){.figs{grid-template-columns:minmax(0,1fr)}}
@media screen and (max-width:720px){
  main{padding:24px 16px 48px}
  h1{font-size:24px}
  h2{font-size:18px;margin-top:36px}
  .card{padding:16px}
  .fig{padding:12px}
  .big{font-size:28px}
  .tbl table,.tbl tbody,.tbl tr,.tbl td{display:block;width:auto}
  .tbl table{table-layout:auto}
  .tbl thead{display:none}
  .tbl tr{border:1px solid var(--line);border-radius:6px;padding:10px 12px;margin-bottom:10px}
  .tbl td{border:0;padding:3px 0;text-align:left}
  .tbl td.val,.tbl td.val.na,.tbl td.val.txt{text-align:left}
  .tbl td::before{content:attr(data-label);font-size:11.5px;font-weight:400;color:var(--faint)}
  .tbl td.val::before,.tbl td.tgt::before,.tbl td.when::before{display:inline-block;min-width:68px}
  .tbl td.plain::before,.tbl td.src::before{display:block}
  .tbl td.val.na .why{margin-left:68px}
  .tbl td.name::before{content:none}
  .tbl td.name{font-size:15px;padding-bottom:6px}
  .tbl td.when span{display:inline;margin-right:6px}
}
@media print{
  @page{size:A4;margin:14mm 12mm}
  body{font-size:10pt}
  main{max-width:none;padding:0}
  h2{margin-top:22px}
  .card,.fig,tr{break-inside:avoid}
  .figs{grid-template-columns:repeat(2,minmax(0,1fr))}
  .fig details{display:none}
  thead{display:table-header-group}
  .tbl{overflow:visible}
  .big{font-size:22pt}
}
''' % {'ink': INK, 'muted': MUTED, 'faint': FAINT, 'line': LINE, 'soft': SOFT, 'accent': ACCENT,
       'accent_soft': ACCENT_SOFT, 'font': FONT_STACK}


def _esc(s):
    return html.escape(scrub(s), quote=True)


def _html_card(card):
    vals = []
    for v in card['values']:
        if v['na']:
            vals.append('<div class="val"><div class="big na">%s</div><div class="vlabel">%s</div>'
                        '<div class="vreason">%s</div></div>' % (NA, _esc(v['label']), _esc(v['reason'])))
        else:
            vals.append('<div class="val"><div class="big">%s</div><div class="vlabel">%s</div></div>'
                        % (_esc(v['text']), _esc(v['label'])))
    chart = ''
    if card['chart'] is not None:
        svg, _, _ = chart_svg(card['chart'])
        if svg:
            chart = '<div class="chart">%s</div>' % svg
    note = '<p class="note">%s</p>' % _esc(card['note']) if card['note'] else ''
    return '<section class="card%s"><div class="ctext"><h3>%s</h3><div class="vals">%s</div>%s</div>%s</section>' % (
        ' wide' if chart else '', _esc(card['title']), ''.join(vals), note, chart)


def _html_row(m, missing_by_name):
    labels = COLUMNS
    if m.get('value') is None:
        val = '<td class="val na" data-label="%s">%s<span class="why">%s</span></td>' % (
            labels[1], NA, _esc(null_reason(m, missing_by_name)))
        detail = ''
    else:
        is_text = m.get('unit') == 'text' or not _is_num(m.get('value'))
        val = '<td class="val%s" data-label="%s">%s</td>' % (' txt' if is_text else '', labels[1], _esc(fmt_value(m)))
        detail = (m.get('detail') or '').strip()
    plain = plain_of(m)
    plain_html = _esc(plain) + ('<span class="det">%s</span>' % _esc(detail) if detail else '')
    kind, when = fmt_measure(m)
    when_html = '<span>%s</span>' % _esc(kind) + ('<span>%s</span>' % _esc(when) if when else '')
    return ('<tr><td class="name" data-label="%s">%s</td>%s<td class="tgt" data-label="%s">%s</td>'
            '<td class="plain" data-label="%s">%s</td><td class="src" data-label="%s">%s</td>'
            '<td class="when" data-label="%s">%s</td></tr>'
            % (labels[0], _esc(m.get('name')), val, labels[2], _esc(fmt_target(m)), labels[3],
               plain_html or '—', labels[4], _esc(m.get('source')) or '—', labels[5], when_html))


def _html_figure(ch):
    svg, _, _ = figure_svg(ch)
    head, rows, refs = chart_table(ch)
    thead = ''.join('<th scope="col">%s</th>' % html.escape(h) for h in head)
    body = ''.join('<tr><th scope="row">%s</th>%s</tr>' % (
        html.escape(r[0]), ''.join('<td>%s</td>' % html.escape(c) for c in r[1:])) for r in rows)
    ref = '<p class="nref">기준선: %s</p>' % html.escape(' · '.join(refs)) if refs else ''
    return ('<figure class="fig" id="chart-%s"><h3>%s</h3>%s<div class="plot">%s</div>%s<p class="fsrc">%s</p>'
            '<details><summary>숫자 보기</summary><div class="ntbl"><table class="nums"><thead><tr>%s</tr></thead>'
            '<tbody>%s</tbody></table>%s</div></details></figure>'
            % (ch['id'], html.escape(ch['title']),
               '<p class="sub">%s</p>' % html.escape(ch['subtitle']) if ch['subtitle'] else '', svg,
               '<p class="fnote">%s</p>' % html.escape(ch['note']) if ch['note'] else '',
               html.escape(chart_foot(ch)), thead, body, ref))


def build_html(model):
    out = ['<!doctype html><html lang="ko"><head><meta charset="utf-8">'
           '<meta name="viewport" content="width=device-width, initial-scale=1">'
           '<title>%s</title><style>%s</style></head><body><main>' % (TITLE, CSS)]
    out.append('<header><div class="eyebrow">S-Brain · 공고 수집·매칭·자격 판정</div><h1>%s</h1>'
               '<p class="runat">측정 %s</p>' % (TITLE, _esc(model['run_at'])))
    if model['conditions']:
        out.append('<ul class="cond">%s</ul>' % ''.join('<li>%s</li>' % _esc(c) for c in model['conditions']))
    out.append('</header>')
    if model['cards']:
        out.append('<h2>한눈에 보기</h2><div class="cards">%s</div>' % ''.join(_html_card(c) for c in model['cards']))
    head = ''.join('<th class="c-%s" scope="col">%s</th>' % (k, t) for k, t in
                   zip(('name', 'val', 'tgt', 'plain', 'src', 'when'), COLUMNS))
    for sec in model['sections']:
        out.append('<h2>%s</h2>' % _esc(sec['title']))
        if sec['ml']:
            out.append('<p class="lead">%s</p>' % ML_LEAD)
        if sec.get('charts'):
            out.append('<div class="figs">%s</div>' % ''.join(_html_figure(ch) for ch in sec['charts']))
        if sec['metrics']:
            rows = ''.join(_html_row(m, model['missing_by_name']) for m in sec['metrics'])
            out.append('<div class="tbl"><table><thead><tr>%s</tr></thead><tbody>%s</tbody></table></div>'
                       % (head, rows))
    out.append('<h2>다시 뽑는 방법</h2><p>PowerShell 에서 아래 두 줄을 실행한다.</p><pre><code>%s</code></pre>'
               '<ul class="plainlist">%s</ul><p>보고서 두 파일(HTML·Word)만 다시 만들 때: <code>%s</code></p>'
               % (html.escape('\n'.join(RERUN_CMD)), ''.join('<li>%s</li>' % html.escape(p) for p in RERUN_POINTS),
                  html.escape(DOC_CMD)))
    if model['missing']:
        items = []
        for x in model['missing']:
            tag = '<span class="tag">%s</span>' % ('예상한 빈칸' if x.get('expected') else '확인 필요')
            items.append('<li><b>%s</b> — %s%s</li>' % (_esc(x.get('name')), _esc(x.get('reason') or '이유가 기록되지 않았다.'), tag))
        out.append('<h2>이번에 재지 못한 것</h2><ul class="plainlist">%s</ul>' % ''.join(items))
    out.append('<footer>metrics.json 의 값을 그대로 옮긴 문서다. 만든 곳: eval/metrics_doc.py</footer>')
    out.append('</main></body></html>\n')
    return ''.join(out)


# ── Edge 로 SVG → PNG ───────────────────────────────────

def find_edge():
    for p in _EDGE_PATHS:
        if os.path.exists(p):
            return p
    for name in ('msedge', 'microsoft-edge', 'microsoft-edge-stable'):
        found = shutil.which(name)
        if found:
            return found
    return None


def _png_ready(path):
    try:
        size = os.path.getsize(path)
        if size < 64:
            return False
        with open(path, 'rb') as f:
            head = f.read(8)
            f.seek(-12, os.SEEK_END)
            tail = f.read()
        return head == b'\x89PNG\r\n\x1a\n' and tail.endswith(b'IEND\xaeB`\x82')
    except OSError:
        return False


MIN_SHOT_H = 100      # Edge headless 는 창 높이가 이보다 낮으면 화면 저장에 실패할 수 있다 — 크게 찍고 아래를 잘라 낸다
SHOT_SCALE = 2
_PNG_SIG = b'\x89PNG\r\n\x1a\n'


def _png_chunk(kind, body):
    return struct.pack('>I', len(body)) + kind + body + struct.pack('>I', zlib.crc32(kind + body) & 0xffffffff)


def crop_png_rows(path, rows):
    """PNG 를 위에서 rows 줄만 남기고 아래를 잘라 낸다. 성공하면 True(실패해도 원본은 그대로).

    PNG 의 줄 필터는 위 줄만 참조하므로 필터를 풀지 않고 압축만 풀어 앞의 줄들을 그대로 남기면 된다."""
    try:
        with open(path, 'rb') as f:
            data = f.read()
        if data[:8] != _PNG_SIG:
            return False
        pos, chunks = 8, []
        while pos + 8 <= len(data):
            n = struct.unpack('>I', data[pos:pos + 4])[0]
            kind = data[pos + 4:pos + 8]
            chunks.append((kind, data[pos + 8:pos + 8 + n]))
            pos += 12 + n
            if kind == b'IEND':
                break
        if not chunks or chunks[0][0] != b'IHDR':
            return False
        w, h, depth, ctype, comp, filt, interlace = struct.unpack('>IIBBBBB', chunks[0][1])
        if rows == h:
            return True
        if interlace or not 0 < rows < h or ctype not in (0, 2, 3, 4, 6):
            return False
        stride = (w * {0: 1, 2: 3, 3: 1, 4: 2, 6: 4}[ctype] * depth + 7) // 8
        raw = zlib.decompress(b''.join(body for kind, body in chunks if kind == b'IDAT'))
        if len(raw) != h * (stride + 1):
            return False
        packed = zlib.compress(raw[:rows * (stride + 1)], 6)
        out, done = [_PNG_SIG], False
        for kind, body in chunks:
            if kind == b'IHDR':
                body = struct.pack('>IIBBBBB', w, rows, depth, ctype, comp, filt, interlace)
            elif kind == b'IDAT':
                if done:
                    continue
                body, done = packed, True
            out.append(_png_chunk(kind, body))
        tmp = path + '.crop'
        with open(tmp, 'wb') as f:
            f.write(b''.join(out))
        os.replace(tmp, path)
        return True
    except (OSError, ValueError, struct.error, zlib.error):
        return False


def svg_to_png(svg, width, height, png_path, edge, work, timeout=40):
    """SVG 한 장을 Edge headless 화면 저장으로 PNG 로 만든다. 성공하면 True.

    Edge 실행 파일은 화면 저장을 다른 프로세스에 넘기고 바로 끝날 수 있어, PNG 가 다 써질 때까지 기다린다.
    그래프마다 따로 빈 프로필을 써서 앞 그래프의 Edge 와 섞이지 않게 한다.
    키가 MIN_SHOT_H 보다 작은 그림은 창을 그만큼 키워 찍은 뒤 아래 흰 여백을 잘라 낸다."""
    stem = os.path.splitext(os.path.basename(png_path))[0]
    page = os.path.join(work, stem + '.html')
    shot_h = max(int(height), MIN_SHOT_H)
    with open(page, 'w', encoding='utf-8') as f:
        f.write('<!doctype html><html><head><meta charset="utf-8"><style>html,body{margin:0;padding:0;'
                'background:#fff;overflow:hidden}svg{display:block}</style></head><body>%s</body></html>' % svg)
    cmd = [edge, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--no-first-run',
           '--no-default-browser-check', '--disable-extensions', '--disable-sync',
           '--user-data-dir=' + os.path.join(work, 'profile_' + stem), '--force-device-scale-factor=%d' % SHOT_SCALE,
           '--screenshot=' + png_path, '--window-size=%d,%d' % (width, shot_h), pathlib.Path(page).as_uri()]
    start = time.time()
    try:
        subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout)
    except (OSError, subprocess.SubprocessError):
        return False
    while time.time() - start < timeout:
        if _png_ready(png_path):
            time.sleep(0.2)          # 마지막 쓰기가 끝났는지 한 번 더 본다
            if _png_ready(png_path):
                if shot_h != height:
                    crop_png_rows(png_path, int(height) * SHOT_SCALE)   # 못 자르면 흰 여백이 남을 뿐 그림은 맞다
                return True
        time.sleep(0.15)
    return False


def _remove_tree(path, tries=20):
    for _ in range(tries):
        shutil.rmtree(path, ignore_errors=True)
        if not os.path.exists(path):
            return True
        time.sleep(0.25)             # 막 끝난 Edge 가 프로필 파일을 잠시 쥐고 있을 수 있다
    return False


def png_jobs(model, word=True, export=True):
    """Edge 로 만들 PNG 목록: [(열쇠, svg, 폭, 높이, 파일 이름)].

    ('card', i)·('chart', id) 는 Word 에 넣을 그래프 그림, ('out', 파일 이름) 은 그래프 폴더에 둘 발표용 그림
    (제목·부제·쉬운 설명·출처를 그림 안에 함께 넣는다). 그래프 폴더는 charts 가 있을 때만 만든다."""
    heads = []
    for ci, card in enumerate(model['cards']):
        if card['chart'] is not None:
            svg, w, h = chart_svg(card['chart'])
            if svg:
                heads.append((ci, card, svg, w, h))
    jobs = []
    if word:
        for ci, _, svg, w, h in heads:
            jobs.append((('card', ci), svg, w, h, 'w_card_%02d.png' % ci))
        for ch in model['charts']:
            svg, w, h = figure_svg(ch)
            jobs.append((('chart', ch['id']), svg, w, h, 'w_%s.png' % ch['id']))
    if export and model['charts']:
        for li, (_, card, _svg, _w, _h) in enumerate(heads[:26]):
            name = '00%s_headline.png' % chr(97 + li)
            m = card['chart']
            svg, w, h = chart_svg(m, W=PNG_FIG_W)
            foot = chart_foot({'source': m.get('source'), 'kind': m.get('kind'), 'measured_at': m.get('measured_at')})
            # 그림만 떼어 붙여도 뜻이 통하게 부제에 지표의 쉬운 설명을 넣는다
            jobs.append((('out', name),) + card_svg(card['title'], plain_of(m), svg, w, h, card['note'], foot) + (name,))
        for ch in model['charts']:
            svg, w, h = figure_svg(ch, W=PNG_FIG_W)
            jobs.append((('out', ch['png']),) + card_svg(ch['title'], png_subtitle(ch), svg, w, h, ch['note'],
                                                          chart_foot(ch)) + (ch['png'],))
    return jobs


def png_subtitle(ch):
    """발표용 PNG 부제. 서비스 밖 실험(ML) 그래프는 그 말이 부제에 없으면 앞에 붙인다(HTML·Word 의 절 머리말이 PNG 에는 없어서)."""
    sub = ch['subtitle']
    if ch.get('area') == 'ML' and '서비스 밖 실험' not in sub:
        sub = ML_TAG + (' · ' + sub if sub else '')
    return sub


def run_pngs(jobs, edge, work):
    """PNG 를 Edge 몇 개로 나눠 동시에 만든다(그래프마다 따로 빈 프로필). {열쇠: png 경로 또는 None}."""
    def one(job):
        key, svg, w, h, name = job
        path = os.path.join(work, name)
        ok = svg_to_png(svg, w, h, path, edge, work)
        if not ok:                                   # 한 번 더(Edge 가 가끔 첫 화면 저장을 놓친다) — 남은 반쪽 파일은 지우고
            try:
                os.remove(path)
            except OSError:
                pass
            ok = svg_to_png(svg, w, h, path, edge, work)
        return key, (path if ok else None)
    if not jobs:
        return {}
    with ThreadPoolExecutor(max_workers=min(EDGE_WORKERS, len(jobs))) as pool:
        return dict(pool.map(one, jobs))


def save_chart_pngs(folder, results, names, say):
    """그래프 폴더에 발표용 PNG 를 쓴다. 새로 하나라도 만들었으면 예전에 이 도구가 쓴 NN_*.png 는 먼저 지운다
    (다른 파일은 그대로 둔다). 쓴 파일 이름 목록을 돌려준다."""
    made = [(n, results.get(('out', n))) for n in names]
    ok = [(n, p) for n, p in made if p]
    for n, p in made:
        if not p:
            say('그래프 PNG 를 만들지 못했다: %s' % n)
    if not ok:
        return []
    if os.path.isdir(folder):
        for old in os.listdir(folder):
            if _PNG_OUT.match(old):
                try:
                    os.remove(os.path.join(folder, old))
                except OSError:
                    pass
    os.makedirs(folder, exist_ok=True)
    for n, p in ok:
        tmp = os.path.join(folder, n + '.tmp')
        shutil.copyfile(p, tmp)
        os.replace(tmp, os.path.join(folder, n))
    say('그래프 PNG %d장을 %s 폴더에 썼다.' % (len(ok), CHART_DIR))
    return [n for n, _ in ok]


# ── Word ────────────────────────────────────────────────

_RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline',
              'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing',
              'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText',
              'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']
_TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
                'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
                'tblCaption', 'tblDescription', 'tblPrChange']
_TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
               'textDirection', 'tcFitText', 'vAlign', 'hideMark']
_TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight',
               'tblHeader', 'tblCellSpacing', 'jc', 'hidden']


def _docx():
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def _put(parent, child, order):
    """XML 자식을 스키마 순서에 맞는 자리에 넣는다(같은 이름은 바꿔 끼운다)."""
    name = child.tag.rsplit('}', 1)[-1]
    for old in [c for c in parent if c.tag == child.tag]:
        parent.remove(old)
    rank = order.index(name)
    for c in parent:
        cname = c.tag.rsplit('}', 1)[-1]
        if cname in order and order.index(cname) > rank:
            c.addprevious(child)
            return child
    parent.append(child)
    return child


def build_docx(model, path, pngs=None, log=None):
    """Word 문서를 만든다. pngs = run_pngs 결과({('card', i)·('chart', id): png 경로}). 그림이 없으면 숫자 표로 대신한다.
    log: 그림을 만들려 했는데(pngs 가 있음) 못 만들어 표로 대신한 것을 한 줄씩 받는 함수."""
    pngs = pngs or {}
    tried = bool(pngs)

    def say(line):
        if tried and callable(log):
            log(line)
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt, RGBColor

    def rgb(hexstr):
        return RGBColor.from_string(hexstr.lstrip('#').upper())

    def el(tag, **attrs):
        node = OxmlElement(tag)
        for k, v in attrs.items():
            node.set(qn('w:' + k), str(v))
        return node

    def set_font(rpr):
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = _put(rpr, el('w:rFonts'), _RPR_ORDER)
        for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + a), None)
        for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + a), DOCX_FONT)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    usable = 17.0

    # 글꼴: 문서 기본값과 쓰는 스타일 모두 맑은 고딕(동아시아 글꼴 포함)
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        rpr_default = defaults.find(qn('w:rPrDefault'))
        if rpr_default is not None and rpr_default.find(qn('w:rPr')) is not None:
            set_font(rpr_default.find(qn('w:rPr')))
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.2
    nrpr = normal.element.get_or_add_rPr()
    set_font(nrpr)
    _put(nrpr, el('w:lang', val='ko-KR', eastAsia='ko-KR'), _RPR_ORDER)
    for name, size, color, before, after in (('Title', 22, INK, 0, 6), ('Heading 1', 15, ACCENT, 20, 8),
                                             ('Heading 2', 12, INK, 12, 4)):
        st = doc.styles[name]
        set_font(st.element.get_or_add_rPr())
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ('List Bullet', 'Table Grid'):
        try:
            set_font(doc.styles[name].element.get_or_add_rPr())
        except KeyError:
            pass
    title_ppr = doc.styles['Title'].element.find(qn('w:pPr'))
    if title_ppr is not None:
        for bottom in title_ppr.iter(qn('w:bottom')):
            for a in ('themeColor', 'themeShade', 'themeTint'):   # 테마 색이 남아 있으면 w:color 보다 앞선다
                bottom.attrib.pop(qn('w:' + a), None)
            bottom.set(qn('w:color'), ACCENT.lstrip('#').upper())

    props = doc.core_properties
    props.title = TITLE
    props.subject = 'metrics.json 지표 보고서'
    props.author = ''
    props.last_modified_by = ''
    props.comments = ''

    def para(text='', size=None, color=None, bold=False, align=None, after=None, style=None, mono=False):
        p = doc.add_paragraph(style=style)
        if text:
            add_run(p, text, size, color, bold, mono)
        if align:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def add_run(p, text, size=None, color=None, bold=False, mono=False):
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if bold:
            r.font.bold = True
        if mono:
            r.font.name = 'Consolas'
            rpr = r._element.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            fonts.set(qn('w:eastAsia'), DOCX_FONT)
        return r

    def shade(cell, fill):
        _put(cell._tc.get_or_add_tcPr(), el('w:shd', val='clear', color='auto', fill=fill.lstrip('#').upper()),
             _TCPR_ORDER)

    def style_table(table, widths):
        table.style = doc.styles['Table Grid']
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        tblpr = table._tbl.tblPr
        borders = el('w:tblBorders')
        for edge_name, val, sz, color in (('top', 'single', 8, ACCENT), ('left', 'nil', 0, ''),
                                          ('bottom', 'single', 4, '#c9d0d4'), ('right', 'nil', 0, ''),
                                          ('insideH', 'single', 4, LINE), ('insideV', 'nil', 0, '')):
            b = el('w:' + edge_name, val=val)
            if val != 'nil':
                b.set(qn('w:sz'), str(sz))
                b.set(qn('w:space'), '0')
                b.set(qn('w:color'), color.lstrip('#').upper())
            borders.append(b)
        _put(tblpr, borders, _TBLPR_ORDER)
        mar = el('w:tblCellMar')
        for side, w in (('top', 50), ('left', 90), ('bottom', 50), ('right', 90)):
            mar.append(el('w:' + side, w=w, type='dxa'))
        _put(tblpr, mar, _TBLPR_ORDER)
        twips = [int(round(w / 2.54 * 1440)) for w in widths]
        _put(tblpr, el('w:tblW', w=sum(twips), type='dxa'), _TBLPR_ORDER)
        for col, tw in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), twips):
            col.set(qn('w:w'), str(tw))
        for row in table.rows:
            _put(row._tr.get_or_add_trPr(), el('w:cantSplit', val='true'), _TRPR_ORDER)
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        head = table.rows[0]
        _put(head._tr.get_or_add_trPr(), el('w:tblHeader', val='true'), _TRPR_ORDER)
        for cell in head.cells:
            shade(cell, SOFT)

    def cell_text(cell, text, size=9, color=None, bold=False, align=None, first=True):
        p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.1
        if align:
            p.alignment = align
        add_run(p, text, size, color, bold)
        return p

    def header_row(table, names, right=()):
        for i, (cell, name) in enumerate(zip(table.rows[0].cells, names)):
            cell_text(cell, name, size=8.5, color=MUTED, bold=True,
                      align=WD_ALIGN_PARAGRAPH.RIGHT if i in right else None)

    def keep_together(table):
        """작은 표가 쪽 끝에서 머리줄만 남고 갈라지지 않게, 마지막 줄 말고는 '다음 줄과 함께'로 둔다."""
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    # 1. 제목·측정 시각·조건
    doc.add_heading(TITLE, level=0)
    para('측정 ' + model['run_at'], size=10, color=MUTED, after=2)
    for c in model['conditions']:
        para(c, size=9.5, color=MUTED, style='List Bullet', after=0)

    # 2. 한눈에 보기
    if model['cards']:
        doc.add_heading('한눈에 보기', level=1)
        for ci, card in enumerate(model['cards']):
            doc.add_heading(card['title'] or '(제목 없음)', level=2)
            for v in card['values']:
                p = para(after=0)
                if v['na']:
                    add_run(p, NA, 14, MUTED, True)
                else:
                    add_run(p, v['text'], 22, INK, True)      # 큰 숫자는 글자색(계열 색은 그래프에만)
                add_run(p, '   ' + v['label'], 10, MUTED)
                if v['na'] and v['reason']:
                    para(v['reason'], size=8.5, color=FAINT, after=2)
            if card['chart'] is not None:
                png = pngs.get(('card', ci))
                if png:
                    doc.add_picture(png, width=Cm(min(CHART_W / 96 * 2.54, usable)))
                else:
                    say('Word: 한눈에 보기 그래프 그림을 만들지 못해 숫자 표로 넣었다: %s' % scrub(card['chart'].get('name')))
                    groups = chart_groups(card['chart'])
                    c = card['chart'].get('compare') or {}
                    names = ['구분', scrub(c.get('label_before') or '이전'), scrub(c.get('label_after') or '이후')]
                    para('그래프 숫자(%s) — 그림을 만들지 못해 표로 대신한다.' % scrub(card['chart'].get('name')),
                         size=8.5, color=FAINT, after=2)
                    t = doc.add_table(rows=1 + len(groups), cols=3)
                    style_table(t, [4.5, 3.5, 3.5])
                    header_row(t, names, right=(1, 2))
                    for row, (glabel, pairs) in zip(t.rows[1:], groups):
                        cell_text(row.cells[0], glabel, bold=True)
                        for cell, (_, val) in zip(row.cells[1:], pairs):
                            cell_text(cell, _chart_value(val, card['chart'].get('unit')),
                                      align=WD_ALIGN_PARAGRAPH.RIGHT)
                    keep_together(t)
            if card['note']:
                para(card['note'], size=8.5, color=FAINT, after=6)

    def chart_block(ch):
        """영역 그래프 하나: 제목 · 부제 · 그림(없으면 숫자 표) · 쉬운 설명 · 출처."""
        p = para(after=0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        add_run(p, ch['title'], 11, INK, True)
        if ch['subtitle']:
            para(ch['subtitle'], size=8.5, color=MUTED, after=2).paragraph_format.keep_with_next = True
        png = pngs.get(('chart', ch['id']))
        if png:
            doc.add_picture(png, width=Cm(min(FIG_W / 96.0 * 2.54 * 1.15, usable)))
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        else:
            say('Word: 그래프 그림을 만들지 못해 숫자 표로 넣었다: %s' % ch['id'])
            para('그림을 만들지 못해 그래프 숫자를 표로 대신한다.', size=8.5, color=FAINT,
                 after=2).paragraph_format.keep_with_next = True
            head, rows, refs = chart_table(ch)
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            style_table(t, [4.6] + [2.8] * len(ch['series']) + ([3.6] if ch['errors'] else []))
            header_row(t, head, right=range(1, len(head)))
            for row, r in zip(t.rows[1:], rows):
                cell_text(row.cells[0], r[0])
                for cell, v in zip(row.cells[1:], r[1:]):
                    cell_text(cell, v, align=WD_ALIGN_PARAGRAPH.RIGHT)
            keep_together(t)
            if refs:
                para('기준선: ' + ' · '.join(refs), size=8.5, color=MUTED, after=2)
        if ch['note']:
            para(ch['note'], size=9.5, color=INK, after=1)
        para(chart_foot(ch), size=8, color=FAINT, after=8)

    # 3·4. 영역 표, ML
    widths = [3.7, 2.1, 1.8, 4.6, 2.8, 2.0]          # 합 17.0cm = A4 폭 - 좌우 여백
    for sec_ in model['sections']:
        doc.add_heading(sec_['title'], level=1)
        if sec_['ml']:
            lead = para(after=6)
            add_run(lead, ML_LEAD, 10, ACCENT, True)
        for ch in sec_.get('charts') or []:
            chart_block(ch)
        if not sec_['metrics']:
            continue
        t = doc.add_table(rows=1 + len(sec_['metrics']), cols=len(COLUMNS))
        style_table(t, widths)
        header_row(t, COLUMNS, right=(1,))
        for row, m in zip(t.rows[1:], sec_['metrics']):
            c = row.cells
            cell_text(c[0], scrub(m.get('name')), bold=True)
            if m.get('value') is None:
                cell_text(c[1], NA, bold=True, color=MUTED)
                cell_text(c[1], null_reason(m, model['missing_by_name']), size=8, color=FAINT, first=False)
                detail = ''
            else:
                is_text = m.get('unit') == 'text' or not _is_num(m.get('value'))
                cell_text(c[1], fmt_value(m), bold=True, align=None if is_text else WD_ALIGN_PARAGRAPH.RIGHT)
                detail = scrub((m.get('detail') or '').strip())
            cell_text(c[2], fmt_target(m), size=8.5, color=MUTED)
            cell_text(c[3], plain_of(m) or '—')
            if detail:
                cell_text(c[3], detail, size=8, color=MUTED, first=False)
            cell_text(c[4], scrub(m.get('source')) or '—', size=8, color=FAINT)
            kind, when = fmt_measure(m)
            cell_text(c[5], kind, size=8.5, color=MUTED)
            if when:
                cell_text(c[5], when, size=8, color=FAINT, first=False)
        keep_together(t)

    # 5. 다시 뽑는 방법
    doc.add_heading('다시 뽑는 방법', level=1)
    para('PowerShell 에서 아래 두 줄을 실행한다.', after=2)
    for line in RERUN_CMD:
        p = para(after=0)
        p.paragraph_format.left_indent = Cm(0.6)
        add_run(p, line, 9.5, INK, mono=True)
    para(after=2)
    for point in RERUN_POINTS:
        para(point, style='List Bullet', after=0)
    p = para(after=4)
    p.paragraph_format.space_before = Pt(6)
    add_run(p, '보고서 두 파일(HTML·Word)만 다시 만들 때: ')
    add_run(p, DOC_CMD, 9, INK, mono=True)

    # 6. 이번에 재지 못한 것
    if model['missing']:
        doc.add_heading('이번에 재지 못한 것', level=1)
        for x in model['missing']:
            p = para(style='List Bullet', after=0)
            add_run(p, scrub(x.get('name')), bold=True)
            add_run(p, ' — ' + (scrub(x.get('reason')) or '이유가 기록되지 않았다.'))
            add_run(p, '  (%s)' % ('예상한 빈칸' if x.get('expected') else '확인 필요'), 8.5, MUTED)

    para('metrics.json 의 값을 그대로 옮긴 문서다. 만든 곳: eval/metrics_doc.py', size=8, color=FAINT).paragraph_format.space_before = Pt(18)
    tmp = path + '.tmp'
    doc.save(tmp)
    os.replace(tmp, path)
    return path


# ── 진입점 ──────────────────────────────────────────────

def render(out_dir, log=print):
    """<out_dir>/metrics.json → 같은 폴더의 지표_보고서.html · 지표_보고서.docx (+ charts 가 있으면 그래프/ 폴더 PNG).

    (html 경로, docx 경로) 를 돌려준다. python-docx 가 없으면 docx 경로는 None 이다.
    log: 실행 로그를 한 줄씩 받는 함수(뺀 그래프와 그 이유, PNG 결과). None 이면 아무것도 쓰지 않는다."""
    say = log if callable(log) else (lambda _line: None)
    out_dir = os.path.abspath(out_dir)
    with open(os.path.join(out_dir, 'metrics.json'), encoding='utf-8') as f:
        data = json.load(f)
    model = prepare(data)
    for line in model['chart_log']:
        say(line)
    html_path = os.path.join(out_dir, HTML_NAME)
    tmp = html_path + '.tmp'
    with open(tmp, 'w', encoding='utf-8', newline='\n') as f:
        f.write(build_html(model))
    os.replace(tmp, html_path)
    want_docx, want_out = _docx(), bool(model['charts'])
    jobs = png_jobs(model, word=want_docx, export=want_out)
    edge = find_edge() if jobs else None
    results, work = {}, None
    try:
        if jobs and edge:
            work = tempfile.mkdtemp(prefix='metrics_doc_')
            results = run_pngs(jobs, edge, work)
            if want_out:
                save_chart_pngs(os.path.join(out_dir, CHART_DIR), results,
                                [j[4] for j in jobs if j[0][0] == 'out'], say)
        elif want_out:
            say('Edge 를 찾지 못해 그래프 PNG(%s 폴더)를 만들지 않았다. Word 에는 그래프 숫자를 표로 넣는다.' % CHART_DIR)
        docx_path = build_docx(model, os.path.join(out_dir, DOCX_NAME), results, log=say) if want_docx else None
    finally:
        if work:
            _remove_tree(work)
    return html_path, docx_path


def _latest():
    found = sorted(p for p in glob.glob(os.path.join(ROOT, 'reports', 'metrics_*'))
                   if os.path.isfile(os.path.join(p, 'metrics.json')))
    return found[-1] if found else None


def main(argv=None):
    ap = argparse.ArgumentParser(description='metrics.json → 지표_보고서.html · 지표_보고서.docx')
    ap.add_argument('out_dir', nargs='?', help='metrics.json 이 있는 폴더(빼면 가장 최근 reports/metrics_*)')
    args = ap.parse_args(argv)
    folder = args.out_dir or _latest()
    if not folder or not os.path.isfile(os.path.join(folder, 'metrics.json')):
        print('metrics.json 을 찾지 못했다: %s' % (folder or 'reports/metrics_*'), file=sys.stderr)
        return 2
    html_path, docx_path = render(folder)
    print('HTML:', html_path)
    if docx_path:
        print('Word:', docx_path)
    else:
        print('Word: python-docx 가 없어 만들지 않았다(pip install -r requirements-report.txt).')
    return 0


if __name__ == '__main__':
    sys.exit(main())
