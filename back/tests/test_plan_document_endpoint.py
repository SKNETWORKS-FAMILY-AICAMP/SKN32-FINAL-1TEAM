"""GET /projects/{id}/plan-document.docx 확인.
- 매칭/계획서가 아직 없는 프로젝트에서도 200 + 유효 docx(공식 양식 구조)가 나오는지
- generate 이후엔 실제 PlanSection 본문(문제인식/실현가능성/성장전략)이 채워지는지
검증한다."""
import io
import json
import os
import shutil

import pytest
from docx import Document

from app.models import Notice


def _payload(**overrides):
    base = {
        'biz_type': '개인', 'ceo_name': '박테스트',
        'founded_at': None, 'description': 'AI 기반 동네 헬스장 통합 예약 서비스',
        'team_members': [{'name': '박테스트', 'role': '대표', 'experience': '4년'}],
        'pricing_items': [{'service_name': '서버 호스팅', 'unit_price': 50000}],
    }
    base.update(overrides)
    return base


def _seed_notice(db_session):
    notice = Notice(
        notice_id='NOTICE-PLAN-DOC-TEST-1', source='k-startup',
        title='2026년도 초기창업패키지(일반형)', organizer='창업진흥원',
        recruitment_status='open', url='https://example.com/notice/1',
    )
    db_session.add(notice)
    db_session.commit()
    db_session.refresh(notice)
    return notice


def test_plan_document_before_generate_returns_valid_docx(authed_client, db_session):
    r = authed_client.post('/projects', data={'payload': json.dumps(_payload())})
    assert r.status_code == 201, r.text
    project_id = r.json()['project_id']

    r = authed_client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200
    assert r.headers['content-type'] == (
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    assert r.content[:2] == b'PK'

    doc = Document(io.BytesIO(r.content))
    full_text = '\n'.join(p.text for p in doc.paragraphs)
    table_text = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert '초기창업패키지' in full_text
    assert len(doc.tables) >= 5
    # placeholder notation for fields we don't collect yet (기업명 등은 아직 DB에 없음)
    assert '○' in table_text


def test_plan_document_after_generate_contains_real_sections(authed_client, db_session):
    notice = _seed_notice(db_session)
    r = authed_client.post('/projects', data={'payload': json.dumps(_payload())})
    project_id = r.json()['project_id']

    r = authed_client.post(f'/projects/{project_id}/generate', json={'notice_id': notice.notice_id})
    assert r.status_code == 200, r.text
    plan_sections = r.json()['plan']['sections']
    body_1_1 = next(s['body'] for s in plan_sections if s['tag'] == '1-1')

    r = authed_client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    full_text = '\n'.join(p.text for p in doc.paragraphs)
    assert '1. 문제 인식(Problem)_창업 아이템의 필요성' in full_text
    assert body_1_1[:15] in full_text  # 실제 생성된 계획서 본문이 들어갔는지


def test_plan_document_uses_intake_partners_not_empty_project_partners_table(authed_client):
    # [2026-09-22] IntakeForm.jsx "협력 기관" 입력은 project_plan_inputs.partners(JSON)에
    # 저장되는데, _build_plan_document_data는 계속 비어있는 project_partners 테이블을
    # 읽고 있어서 실제로 입력해도 사업계획서엔 항상 플레이스홀더만 나오던 버그 —
    # routers/projects.py partner_rows 수정 확인.
    r = authed_client.post('/projects', data={'payload': json.dumps(_payload(
        partners=[{'name': '○○대학 · 실증 지원', 'status': '협력 중'}],
    ))})
    assert r.status_code == 201, r.text
    project_id = r.json()['project_id']

    r = authed_client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    table_text = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert '○○대학 · 실증 지원' in table_text
    assert '협력 중' in table_text


def test_plan_document_uses_occupation_for_preliminary(authed_client):
    # [2026-09-22] 계획서 양식(예비창업자 전용 "직업" 항목)이 요구하는데 대응 입력칸이
    # 없어 계속 '○○○' 플레이스홀더로만 나가던 항목. "아이템 범주"는 함께 검토했지만
    # Agent가 짓는 항목으로 확인돼(재희님) 사용자 입력에 추가하지 않았다.
    r = authed_client.post('/projects', data={'payload': json.dumps(_payload(
        applicant_type='preliminary', occupation='대학생',
    ))})
    assert r.status_code == 201, r.text
    project_id = r.json()['project_id']

    r = authed_client.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    table_text = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert '대학생' in table_text


def test_plan_document_requires_ownership(login_as, db_session):
    owner = login_as('owner@example.com', '주인')
    r = owner.post('/projects', data={'payload': json.dumps(_payload())})
    project_id = r.json()['project_id']

    other = login_as('other@example.com', '남')
    r = other.get(f'/projects/{project_id}/plan-document.docx')
    assert r.status_code == 404


# ============================================================================
# GET /projects/{id}/plan-document.hwp (SB-59, 2026-09-18)
# ============================================================================
# 예비창업패키지·초기창업패키지(일반형) 둘 다 원본 .hwp + 표 좌표 매핑까지 끝났다
# (app/hwp_export.py의 _preliminary_steps/_early_general_steps). RHWP_BIN이 실제
# rhwp 실행 파일을 안 가리키는 환경(팀 공용 CI 등, 개인 다운로드 경로를 박아둘 수
# 없음)에서는 500(rhwp 실행 파일을 찾을 수 없음)이 정상이고, RHWP_BIN을 로컬에
# 실제로 맞춘 환경(.env)에서는 진짜 200 + .hwp가 나와야 한다 — 두 경우 다 검증한다.
from app.hwp_export import RHWP_BIN as _RHWP_BIN

_RHWP_AVAILABLE = shutil.which(_RHWP_BIN) is not None or os.path.isfile(_RHWP_BIN)


@pytest.mark.parametrize('applicant_type', [None, 'preliminary'])
def test_plan_document_hwp(authed_client, applicant_type):
    payload = _payload(applicant_type=applicant_type) if applicant_type else _payload()
    r = authed_client.post('/projects', data={'payload': json.dumps(payload)})
    project_id = r.json()['project_id']

    r = authed_client.get(f'/projects/{project_id}/plan-document.hwp')
    if _RHWP_AVAILABLE:
        assert r.status_code == 200, r.text
        assert r.headers['content-type'] == 'application/haansofthwp'
        assert len(r.content) > 1000
    else:
        assert r.status_code == 500
        assert 'rhwp 실행 파일을 찾을 수 없습니다' in r.json()['detail']
