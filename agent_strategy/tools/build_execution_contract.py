import json
import re
from pathlib import Path
import openpyxl

base=Path(__file__).resolve().parents[2]/'agent_strategy'
book=openpyxl.load_workbook(base/'전략_작성_검증1.xlsx',read_only=True,data_only=True)
functions={}
limits={'F01':1400,'F02':2200,'F03':2600,'F04':2000,'F05':1400,'F06':2800,'F07':2600,'F08':2200,'F09':1800,'F10':2000,'F11':1600,'F12':2400,'F13':1600,'F14':1000,'F15':1200,'F16':1800,'F17':0,'F18':1200,'F19':1400,'F20':0}
for row in book['함수모음'].iter_rows(values_only=True):
    if row and re.fullmatch('F[0-9]{2}',str(row[0])):
        fid,name,description,inputs,output,model=row[:6]
        tier=next((t for t in ['Sol','Terra','Luna'] if t in model),None)
        functions[fid]={'name':name.removesuffix('()'),'purpose':description,'inputs':inputs,'output':output,'modelLabel':model,'apiModel':f'gpt-5.6-{tier.lower()}' if tier else None,'maxOutputTokens':limits[fid]}
documents={k:[] for k in ['general','pre_startup','early_startup']}
for row in book['전략작성'].iter_rows(values_only=True):
    if len(row)<6 or row[3] not in ['F16','F17','F18']:continue
    match=re.match(r'(\d\.\d\.\d+)\.\s*(.*)',row[2])
    sid,title=match.groups()
    kind={'1':'general','2':'pre_startup','3':'early_startup'}[sid[0]]
    # Workbook's early-stage prose cross-references pre-stage 2.x; normalize to its own 3.x.
    if kind=='early_startup':title=re.sub(r'\b2\.(\d\.\d+)',r'3.\1',title)
    inputs=[s.strip() for s in row[4].split(',') if s.strip() not in ['section_spec','columns','rules','flow_type']]
    rule={}
    if '3가지' in title or '1 2 3' in title:rule['exactItems']=3
    if '5개 이상' in title:rule['minItems']=5
    if '1문장' in title:rule['maxSentences']=1
    if '1줄 소개' in title and '시장동향' not in title:rule['maxLines']=1
    if '3줄' in title:rule['maxLines']=3
    if row[3]=='F17':
        rule['requiredColumns']=(['비목','산출근거','정부지원사업비'] if '사업비' in title else ['구분','추진내용','추진기간','세부내용'] if '일정' in title or '기간을' in title else ['지표','단위','목표값','측정환경'] if '성능지표' in title else ['핵심기술','개발기능'])
    documents[kind].append({'sectionId':sid,'title':title,'functionId':row[3],'contentType':{'F16':'section','F17':'table','F18':'image'}[row[3]],'sourceKeys':inputs,'rules':rule,'enabled':row[3]!='F17'})
out={'version':2,'sourceWorkbook':'전략_작성_검증1.xlsx','tableGenerationEnabled':False,'functions':functions,'documents':documents}
(base/'runtime/execution_contract.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print({k:len(v) for k,v in documents.items()})
